from __future__ import annotations

import copy
import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\anju2\Desktop\Advance_java_research\23BCE0052-23BCE0050-23BCE0080_Strengthened.docx")
OUTPUT = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker\artifacts\manuscript_v2_expand\23BCE0052-23BCE0050-23BCE0080_Expanded_V2.docx")


def find_start(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise KeyError(f"Paragraph beginning with {prefix!r} was not found")


def copy_ppr(source: Paragraph, destination: Paragraph) -> None:
    if source._p.pPr is None:
        return
    destination_ppr = destination._p.get_or_add_pPr()
    for child in list(destination_ppr):
        destination_ppr.remove(child)
    for child in source._p.pPr:
        destination_ppr.append(copy.deepcopy(child))


def copy_rpr(source, destination) -> None:
    if source._r.rPr is None:
        return
    destination_rpr = destination._r.get_or_add_rPr()
    for child in list(destination_rpr):
        destination_rpr.remove(child)
    for child in source._r.rPr:
        destination_rpr.append(copy.deepcopy(child))


def insert_after(anchor, text: str, template: Paragraph, *, alignment=None, bold=None) -> Paragraph:
    new_element = OxmlElement("w:p")
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    paragraph.style = template.style
    copy_ppr(template, paragraph)
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    if template.runs:
        copy_rpr(template.runs[0], run)
    if bold is not None:
        run.bold = bold
    return paragraph


def insert_chain(anchor, entries, body_template, heading_template, caption_template=None):
    current = anchor
    for kind, text in entries:
        if kind == "heading":
            current = insert_after(current, text, heading_template, bold=True)
        elif kind == "main_heading":
            current = insert_after(current, text, heading_template, bold=True)
        elif kind == "caption":
            current = insert_after(current, text, caption_template or heading_template,
                                   alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        else:
            current = insert_after(current, text, body_template)
    return current


def set_text(paragraph: Paragraph, text: str, template: Paragraph, *, bold=None, alignment=None) -> None:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    paragraph.style = template.style
    copy_ppr(template, paragraph)
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    if template.runs:
        copy_rpr(template.runs[0], run)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def write_cell(cell, text: str, *, header=False, size=8.0) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = header
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        set_cell_shading(cell, "D9EAF7")


def make_objective_table_after(anchor: Paragraph, doc: Document):
    table = doc.add_table(rows=1, cols=4)
    # The retained manuscript uses a custom table style rather than Word's
    # default Table Grid style. Reuse it so the new table remains visually
    # consistent even when Table Grid is absent from the template.
    if doc.tables and doc.tables[0].style is not None:
        table.style = doc.tables[0].style
    table.autofit = False
    anchor._p.addnext(table._tbl)
    headers = ["Objective", "Implemented or documented basis", "Evidence in this study", "Validation still required"]
    widths = [0.75, 2.05, 1.75, 1.95]
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        set_cell_width(cell, width)
        write_cell(cell, text, header=True, size=7.5)
    rows = [
        ["O1", "Gateway admission, JWT/ownership checks, device-token verification, DTO validation, and ingress authorization.", "Implemented source behaviour; dated local authorization checks.", "Independent security assessment, credential lifecycle tests, and hardened live-route revalidation."],
        ["O2", "Kafka publication/consumption, InfluxDB persistence, MySQL state, configured retries/DLT, alert path, and observability components.", "Dated local event-path, alert, and single-host load observations; source-level tests.", "Correlation-based reconciliation, repeated performance studies, fault recovery, and multi-instance behaviour."],
        ["O3", "Fixed Random Forest, chronological split, validation thresholding, exact rerun, and model-artifact comparison.", "Historical CSV metrics reproduced exactly; interface smoke test and in-process checks are separately scoped.", "Documented external data source, causal feature provenance, and household-level evaluation."],
        ["O4", "Official REFIT checksum verification, cohort selection, actual-delta normalization, quality reports, and repeatability verification.", "Byte-identical H1/H2/H4 preprocessing outputs and auditable interval-energy estimates.", "Unit-contract approval, isolated credentials, pilot/full replay, label alignment, and REFIT anomaly evaluation."],
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row_values, widths):
            set_cell_width(cell, width)
            write_cell(cell, value, size=7.5)
    return table


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    # Capture style-bearing paragraphs before any XML insertions.
    intro_heading = find_start(doc, "1. Introduction")
    body = find_start(doc, "Residential energy monitoring is more than")
    subheading = find_start(doc, "1.1 Research Objectives")
    caption = find_start(doc, "Table I. Comparative summary")
    related_heading = find_start(doc, "2. Related Work")
    related_start = find_start(doc, "Rovetto et al.")
    related_arch_end = find_start(doc, "Veloso et al.")
    anomaly_start = find_start(doc, "At the anomaly-detection level")
    monitoring_start = find_start(doc, "IoT-monitoring studies")
    data_start = find_start(doc, "REFIT is treated separately")
    related_synthesis = find_start(doc, "Taken together, the reviewed work")
    gap_heading = find_start(doc, "2.1 Research Gaps Identified")
    gap1 = find_start(doc, "Research Gap 1:")
    gap2 = find_start(doc, "Research Gap 2:")
    gap3 = find_start(doc, "Research Gap 3:")
    gap_synthesis = find_start(doc, "Synthesis of Research Gaps")
    objectives_heading = find_start(doc, "1.1 Research Objectives")
    objective_alignment = find_start(doc, "Objective-to-evidence alignment:")
    method_heading = find_start(doc, "3.1 Overall System Methodology")
    method_first = find_start(doc, "The system accepts a Wokwi-compatible")
    method_second = find_start(doc, "The Usage Service consumes")
    method_boundary = find_start(doc, "Architecture implementation, historical-baseline")
    architecture_heading = find_start(doc, "3.2 Proposed System Architecture")
    architecture_explanation = find_start(doc, "The solid path runs")
    identity_heading = find_start(doc, "3.3 Identity, Admission, and Event Processing")
    identity_first = find_start(doc, "The gateway defines the API boundary")
    identity_second = find_start(doc, "After contract validation")
    model_heading = find_start(doc, "3.4 Baseline Model and Alerting Boundary")
    model_first = find_start(doc, "The retained Random Forest artefact")
    model_second = find_start(doc, "The model uses 300 trees")
    model_alert = find_start(doc, "The Alert Service has a configured")
    refit_heading = find_start(doc, "3.5 REFIT Preprocessing")
    refit_first = find_start(doc, "The five local cleaned REFIT files")
    refit_formula_intro = find_start(doc, "For each eligible pair")
    refit_condition_intro = find_start(doc, "Only intervals satisfying")
    refit_exclusions = find_start(doc, "Cross-gap intervals are excluded")
    algorithm_heading = find_start(doc, "3.6 Algorithmic Procedure")
    algorithm_last = find_start(doc, "Create the H1/H2/H4 primary cohort")
    workflow_heading = find_start(doc, "3.7 Evidence and Reproducibility Workflow")
    workflow_explanation = find_start(doc, "The upper workflow records")
    results_heading = find_start(doc, "4. Experimental Results")
    conditions_heading = find_start(doc, "4.1 Experimental Conditions")
    conditions_first = find_start(doc, "The reported observations")
    conditions_resource = find_start(doc, "The Windows laptop used")
    components_heading = find_start(doc, "4.2 Software Components")
    components_first = find_start(doc, "The local composition includes")
    historical_heading = find_start(doc, "4.3 Historical Baseline Protocol")
    historical_first = find_start(doc, "The historical protocol evaluates")
    reproduction_heading = find_start(doc, "4.4 Offline Reproduction")
    reproduction_first = find_start(doc, "The deterministic trainer was run")
    refit_results_heading = find_start(doc, "4.5 REFIT Data Configuration")
    refit_results_first = find_start(doc, "The H1, H2, and H4 primary-cohort")
    discussion_heading = find_start(doc, "4.6 Results and Discussion")
    discussion_rf = find_start(doc, "Table III shows that")
    discussion_smoke = find_start(doc, "The smoke test shows")
    discussion_refit = find_start(doc, "REFIT preprocessing produced")
    discussion_boundary = find_start(doc, "This result is a preparation")
    table_v_caption = find_start(doc, "Table V. Evidence boundaries")
    limitations_heading = find_start(doc, "5. Limitations and Future Work")
    limitation_historical = find_start(doc, "Historical CSV provenance and simulator boundary")
    limitation_refit = find_start(doc, "REFIT labels and runtime integration")
    limitation_resource = find_start(doc, "Resource, performance, and security boundary")
    limitation_future = find_start(doc, "Staged future validation")
    conclusion_heading = find_start(doc, "6. Conclusion")
    conclusion_first = find_start(doc, "This study documents a local")
    conclusion_second = find_start(doc, "The contribution is therefore")

    # Introduction: add a research narrative before the retained objectives.
    insert_chain(intro_heading, [
        ("heading", "1.1 Residential Telemetry as a Data-to-Decision Chain"),
        ("body", "A residential monitoring application receives more than a numerical consumption value. It must determine whether a message belongs to an authorised device, whether the payload is structurally usable, where the accepted record is retained, which downstream calculation sees it, and whether a later notification can be related to the same household context. A dashboard can display a recent value without answering these questions. In contrast, a monitoring system has to keep the telemetry path sufficiently explicit that an unexpected total, absent alert, or rejected reading can be located at the appropriate boundary."),
        ("body", "The distinction is especially important for simulator-compatible input. The present project accepts software-generated telemetry that follows the application contract, but that contract alone does not establish the physical origin, calibration, sampling behaviour, or billing meaning of a reading. The system is therefore presented as a local software prototype that demonstrates controlled data handling. Physical smart-meter and smart-plug studies remain relevant as measurement contexts, but they do not make a simulated request equivalent to a calibrated field instrument [12], [13], [15], [20]."),
        ("body", "The operational difficulty is not confined to the first HTTP request. Residential readings may be malformed, submitted with an unauthorised device identity, delayed while downstream consumers are unavailable, or interpreted under a unit convention that is not expressed by the payload itself. Each of those cases has a different corrective action. The study consequently treats admission, event publication, persistence, inference, and notification as separate responsibilities rather than using a single endpoint response as evidence that the whole data path behaved correctly."),
        ("heading", "1.2 System Challenges and Architectural Rationale"),
        ("body", "The implemented architecture separates browser-facing control activity from the telemetry data path. Keycloak and the API Gateway mediate user-facing access, whereas the ingestion route can admit a machine-originated record only after device-token checking. This separation reflects a practical difference between a user asking to view household data and a meter-like client sending a single measurement. It also makes a household identifier, device identifier, and caller identity distinct entities rather than assuming that one number proves the others."),
        ("body", "After admission, Kafka provides an asynchronous boundary between telemetry publication and downstream processing. In this design, decoupling is a design rationale: it allows the Ingestion Service to publish an accepted event without embedding time-series writes, alert decisions, and inference logic in the same request path. It does not, by itself, demonstrate an increase in throughput or fault tolerance. The project consequently reports code-level retry and dead-letter routing separately from runtime recovery measurements, and it does not claim exactly-once processing or complete record-level correlation across every service."),
        ("body", "The system also distinguishes time-series values from relational state. InfluxDB holds timestamped usage values, while MySQL is used for user, device, and alert records. This allocation allows a usage query to work over time-stamped values while ownership and alert state remain relational. The choice is aligned with the modular energy-management reasoning in [1]–[3] and [17], but the paper does not infer database scalability or billing accuracy from the choice of technologies."),
        ("heading", "1.3 Model Evidence and Data Provenance"),
        ("body", "An anomaly score cannot be interpreted independently of the data that produced it. Its meaning depends on the source of the measurements, the construction of the predictor variables, the definition of the abnormal class, the split between training and evaluation periods, and the decision rule used to turn a probability into a label. The legacy classifier in this repository can be reproduced under its documented local protocol, but the supplied CSV does not establish an external source, licence, household identity, label-generation process, or causal construction of Avg_Past_Consumption. Reproducibility of that experiment is therefore valuable but narrower than a claim of generalisation."),
        ("body", "The study deliberately separates four evidence classes: implemented source behaviour, dated local runtime observations, deterministic offline reconstruction of the legacy supervised experiment, and REFIT source/preprocessing evidence. These classes answer different questions. A local HTTP response can show that an interface executed; a repeated model artifact can show that a historical procedure is reproducible; and a verified checksum can show that a source file was retained correctly. None of them independently establishes anomaly accuracy on a new household."),
        ("heading", "1.4 Scope, Contribution, and Paper Organisation"),
        ("body", "The central research problem is to describe a residential energy-monitoring prototype without conflating its implemented software, historical model evidence, simulator-compatible checks, and recorded-household data preparation. The contribution is an evidence-aware account of a gateway-mediated, event-driven architecture; a reproducible reconstruction of a legacy Random Forest baseline; and a checksum-traceable REFIT preprocessing procedure for a future isolated evaluation. The paper does not claim that REFIT data have already passed through the live ingestion route, that the classifier is accurate on REFIT, or that alert behaviour has been proven under a REFIT replay."),
        ("body", "Section 2 situates the work within energy-management, monitoring, anomaly-detection, and dataset-provenance literature. Section 3 describes the implementation and the separated evidence streams. Section 4 interprets the historical baseline, dated local checks, and REFIT preprocessing results at their supported scope. Section 5 identifies the technical limitations and the acceptance-gated future evaluation, and Section 6 concludes the study."),
    ], body, subheading, caption)

    set_text(objectives_heading, "1.5 Research Objectives", subheading, bold=True)
    insert_chain(objective_alignment, [
        ("heading", "Objective Scope and Evaluation Boundary"),
        ("body", "Objective 1 addresses the admission boundary: a telemetry system cannot safely associate a reading with a household unless the request contract, caller category, and device association are checked before event publication. The implemented approach uses gateway-mediated user context for interactive requests and a device-bound token for machine telemetry. Available evidence includes the implementation and selected dated local authorisation checks; independent security evaluation and a credential lifecycle remain outside the study."),
        ("body", "Objective 2 addresses processing visibility. Kafka, time-series storage, relational state, alert handling, and observability components make the intended stages visible, while the source includes bounded retry and dead-letter routing for a failed time-series write. Dated local checks and component tests support selected parts of the path. They do not yet establish sustained consumer-lag behaviour, cross-service event reconciliation, or recovery under the complete REFIT workload."),
        ("body", "Objective 3 addresses reproducible interpretation of the legacy model. The study records the data split, features, configuration, threshold protocol, and exact reconstructed results so that a historical score is not treated as an opaque model claim. The evidence is sufficient to show reproduction of the available CSV experiment, but not to resolve the source and feature-provenance limitations or to support new-household accuracy."),
        ("body", "Objective 4 addresses preparation for a recorded-household experiment. REFIT files are verified, cohort selection is declared, interval energy is estimated with actual time differences, and derived artifacts are reproduced byte-for-byte. The objective intentionally stops before live ingestion, device provisioning, label alignment, and model evaluation. Those remaining activities are described as staged validation requirements rather than as completed work."),
    ], body, subheading, caption)

    # Themed literature review.
    insert_chain(related_heading, [
        ("heading", "2.1 Microservices and Service Separation for Energy Management"),
        ("body", "Recent energy-management research provides several reasons to isolate monitoring responsibilities. Rovetto et al. [1] present residential energy-efficiency monitoring through services, while Motta et al. [2] describe a cloud Home Energy Management System model that differentiates ingestion, operational, and analytical environments. These publications support a conceptual decomposition in which a reading should be handled differently from an operational request or an analytical result. They do not prescribe the exact services, storage systems, or security boundary used in the present local implementation."),
        ("body", "Veloso et al. [3] extend the modular argument into smart-grid settings by combining hybrid communication with a microservices solution. Zhu and Yue [17] likewise use technologies including Docker, Kafka, and InfluxDB in a smart-city energy-management setting. Their work demonstrates that the same categories of concerns can be separated in larger installations. However, published deployment outcomes from different locations, communication assumptions, and workloads cannot be transferred directly to this residential prototype. The present study therefore uses them to justify architectural roles, not to establish a numerical benchmark."),
        ("body", "Across these studies, a modular architecture is useful only if boundaries are inspectable. In the present system, the Gateway controls public admission; Ingestion publishes telemetry; Usage stores and aggregates consumption; the inference service produces a model response; Alert Service applies notification policy; and monitoring components observe services. That decomposition reduces coupling in the design, but a complete trace from input to final effect still requires runtime evidence. In particular, the core energy event does not by itself supply a durable end-to-end correlation identifier, so the current architecture should not be described as providing complete record-level tracing."),
        ("body", "The literature therefore establishes service separation as a viable architectural pattern, while leaving an application-specific question about evidence: which transitions have actually been observed and which remain design intent? This question motivates the following themes on measurement, anomaly interpretation, and source preparation."),
    ], body, subheading, caption)

    insert_chain(related_arch_end, [
        ("body", "A common implication of these architectures is that a high-level system diagram is not an evaluation protocol. A service may be implemented, started, or reachable without proving that a particular telemetry record is delivered exactly once, persisted under the intended timestamp, or carried into the alert decision that a user sees. This distinction guides the evidence labels used throughout the present paper."),
    ], body, subheading, caption)

    insert_chain(anomaly_start, [
        ("heading", "2.2 Smart-Meter Anomaly Detection and Evaluation Context"),
        ("body", "Hernández et al. [4] relate smart-meter observations to anomalies in daily activity. The important lesson for this work is methodological rather than numerical: the meaning of an anomaly depends on the measurement setting and the target definition. A classifier score can appear strong under a documented experiment yet be inappropriate for a different household signal, label ontology, or feature set. This is directly relevant to the present legacy CSV, whose labels and source history are not fully documented."),
        ("body", "Kulkarni et al. [8] discuss faults, billing errors, and misuse in a smart-meter anomaly setting. That application context supports the need for careful admission and alerting, but the publication's data and labels are not a substitute for the repository's historical CSV or for REFIT. Maitra et al. [10] propose a convolutional autoencoder with a dynamic threshold; its design illustrates a different way to treat temporal variation, but it is an arXiv preprint and is not used as an empirical comparator for the fixed Random Forest retained here."),
        ("body", "Liu and Hsiao [12] combine smart-plug monitoring with regression-based anomaly detection. Shuaibu et al. [19] describe appliance monitoring and anomaly inference over DDS middleware. Both extend beyond the evidence in this paper: smart-plug systems make particular physical measurement assumptions, and appliance-level pipelines define a different target from a whole-house aggregate energy record. Their results cannot be ranked against the historical Random Forest score without a common dataset, labels, split, and feature contract."),
        ("body", "The existing classifier uses a supervised Random Forest rather than an autoencoder, a regression model, or a multi-task network. Its contribution in this study is not a new learning algorithm. Instead, the work makes the legacy configuration and its reproducibility boundary explicit: a 70/15/15 chronological split, validation-based threshold selection, a fixed feature set, and a separated held-out test. This allows the historical result to be understood at the correct level of evidence."),
        ("body", "The studies in this theme collectively show that model type alone does not define validity. Measurement source, feature availability, label semantics, and evaluation protocol are equally decisive. The present paper uses that conclusion to avoid treating a Wokwi-compatible inference response or a REFIT checksum as an anomaly-detection result."),
    ], body, subheading, caption)

    insert_chain(monitoring_start, [
        ("heading", "2.3 IoT Monitoring, Physical Measurement, and Notification Context"),
        ("body", "Karthika et al. [11] provide IoT energy-monitoring context, while Hidayat and Al-Faris [13] use an ESP32/PZEM-004T prototype to obtain physical electrical parameters. Hossain and Hassan [15] similarly report a real-time meter configuration. De Sousa and Rabelo [20] include residential consumption and power-quality considerations. These studies are valuable because they make the measurement layer visible; they also highlight why an application payload generated by a simulator should not be presented as a physical electrical measurement."),
        ("body", "Notification is another boundary that should be separated from the predictive score. A threshold crossing, model anomaly, stored alert, e-mail send attempt, and final recipient delivery are related but not identical events. The present system includes threshold and anomaly pathways plus a 15-minute duplicate-suppression policy. It is a local design and a selectively tested implementation, not a replacement for the operational guarantees that may be needed in a deployed utility or clinical-like notification system."),
    ], body, subheading, caption)

    insert_chain(data_start, [
        ("heading", "2.4 Household Data, Preprocessing, and Reproducibility"),
        ("body", "REFIT is different from the architecture and model studies because it is a documented household-data resource. Murray et al. [5] describe the measurement study, and the University of Strathclyde and Zenodo records [6], [7] provide the cleaned-data provenance used in this project. The dataset is valuable because it exposes household aggregate active power, paired human-readable and Unix timestamps, and a documented distribution. It should not be represented as a ready-made aggregate anomaly benchmark merely because it contains rich time-series data."),
        ("body", "For a later event-stream experiment, source integrity, timestamp ordering, the treatment of gaps, and the power-to-energy convention all affect the values presented to the application. The current work therefore verifies the local files against published MD5 values and records a deterministic conversion rule. This approach establishes a repeatable derived input, but it cannot resolve the semantic status of every source annotation or provide aggregate anomaly labels without a separate alignment protocol."),
    ], body, subheading, caption)

    insert_chain(related_synthesis, [
        ("heading", "2.5 Critical Synthesis"),
        ("body", "The literature supports three linked observations. First, modular architectures can make energy-monitoring responsibilities separable, but their production or city-scale outcomes do not transfer automatically to a local composition. Second, anomaly-detection outcomes are conditional on a dataset, labels, and evaluation design; different algorithms cannot be compared by a score alone. Third, physical monitoring systems and household datasets impose measurement and provenance assumptions that simulator-compatible telemetry does not remove. The present study uses these observations to frame an implementation and reproducibility contribution rather than an unqualified performance claim."),
        ("body", "Table I is intentionally compact because the single-column layout cannot accommodate a reliable seven-column literature matrix without reducing readability. Its existing columns jointly capture the research problem, method or setting, reported contribution, and evaluation boundary. The surrounding themed discussion provides the context that a table cannot carry safely."),
    ], body, subheading, caption)
    set_text(caption, "Table I. Thematic comparison of architecture, monitoring, anomaly-detection, and data-provenance studies", caption, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Add the official REFIT provenance row to Table I without redesigning a readable table.
    literature_table = doc.tables[0]
    values = [
        "[5]–[7], 2016–2017",
        "Household energy-data provenance",
        "REFIT descriptor and cleaned-data records",
        "Documents recorded household aggregate and appliance measurements and the cleaned-data distribution.",
        "Dataset/provenance resource, not an anomaly-model benchmark; supports source verification and preprocessing rationale.",
    ]
    new_row = literature_table.add_row()
    for cell, value in zip(new_row.cells, values):
        write_cell(cell, value, size=7.5)

    set_text(gap_heading, "2.6 Research Gaps Identified", subheading, bold=True)
    insert_chain(gap1, [
        ("body", "The literature supports separation of monitoring, storage, analytics, and control concerns, but the evidence required to diagnose a single record remains less explicit in many architectural descriptions. A failure can occur before admission, during broker delivery, at persistence, during model handling, or when an alert policy applies. Without a declared boundary for each stage, it is difficult to distinguish a rejected request from a lost, duplicated, or policy-suppressed event."),
        ("body", "The present project addresses part of this problem by exposing a gateway, ingestion service, Kafka topic, usage consumer, time-series store, model service, and alert service as separate components. Its source also includes code-level retry and dead-letter handling. However, the present study does not establish complete correlation-based reconciliation across those services, exactly-once semantics, or multi-instance alert correctness. The gap is therefore addressed at the level of inspectable architecture and selected local evidence, not closed by a broad reliability claim."),
    ], body, subheading, caption)
    insert_chain(gap2, [
        ("body", "The reviewed anomaly studies confirm that model metrics are protocol-specific. A repeated training run, a thresholded response, and a hold-out score are not equivalent measures. They answer, respectively, whether a procedure can be recreated, whether a service can execute, and how a fixed model behaved on a particular evaluation slice. Confusing these evidence types is especially risky when a historical CSV lacks transparent label-generation and feature-provenance documentation."),
        ("body", "This project resolves the reporting problem by separating the legacy supervised result, its exact offline reconstruction, the Wokwi-compatible interface check, and REFIT preprocessing. The separation improves evidence clarity; it does not make the legacy model more accurate or create an external validation set. Objective 3 therefore remains an objective about reproducibility and bounded interpretation rather than a claim of generalisation."),
    ], body, subheading, caption)
    insert_chain(gap3, [
        ("body", "A household-data experiment must establish more than a file download. The experiment needs a source-verifiable cohort, a stated timestamp policy, a physically coherent transformation from power to energy when the target contract expects energy, an identity-to-device mapping, and a label protocol that belongs to the evaluated signal. These requirements are coupled: a correct interval calculation does not prove that the payload is authorised, and an authorised payload does not turn appliance-level annotations into aggregate household labels."),
        ("body", "The current study addresses the preparation part of this gap through checksum-verified REFIT inputs, a H1/H2/H4 primary cohort, a separate H3/H11 solar-sensitivity cohort, actual-delta interval estimates, and byte-identical outputs. The remaining runtime and semantic stages are deliberately open. Any future claim about REFIT accuracy must first show a verified derived_interval_energy_kwh-to-energyConsumed contract, isolated devices and credentials, request/event/persistence reconciliation, and a defensible aggregate-label alignment procedure."),
    ], body, subheading, caption)
    insert_chain(gap_synthesis, [
        ("body", "Together, the three gaps map directly onto the four objectives. Objectives 1 and 2 establish inspectable admission and processing boundaries; Objective 3 provides a reproducible historical baseline while preserving its external-validity limitation; and Objective 4 provides a controlled recorded-household input for the later validation sequence. The connection is methodological: it specifies what evidence would be needed at the next stage rather than claiming that all stages have already been completed."),
    ], body, subheading, caption)

    # Methods and implementation detail.
    insert_chain(method_heading, [
        ("heading", "3.1.1 Evidence Streams and Study Boundary"),
        ("body", "The methodology uses deliberately separated evidence streams. The first stream is implementation evidence drawn from service code, configuration, schemas, and architecture documentation. The second is dated local runtime evidence from controlled requests, authorization checks, storage observations, alert cases, and single-host tests. The third is an offline reconstruction of the supplied legacy supervised experiment. The fourth is checksum and preprocessing evidence for the REFIT cohort. The results section keeps these streams distinct because their inputs, controlled conditions, and permissible conclusions differ."),
        ("body", "This separation prevents a common interpretation error. An implemented component can be inspected statically without being exercised under a workload; a dated runtime check can establish a specific observed path without establishing its current post-hardening behaviour; and a deterministic preprocessing run can prove that an artifact is reproducible without proving that the application accepted it. The methodology therefore records both the positive evidence and the boundary that remains open."),
    ], body, subheading, caption)

    insert_chain(method_first, [
        ("body", "The minimal telemetry contract is intentionally small: deviceId, non-negative energyConsumed, and timestamp. The contract does not contain an event identifier or an explicit unit field. Consequently, the ingestion layer can reject a structurally invalid reading, but it cannot infer whether a client has supplied watts, kWh, or a double-converted value merely from the JSON schema. This is a decisive limitation for REFIT: the derived interval-energy field is a documented future compatibility candidate, not an accepted live mapping."),
        ("body", "For a user-originated call, the security design uses a JWT-derived principal, role checks, and ownership context. For telemetry that represents a device client, a device-bound token is presented to the gateway. The gateway removes externally supplied privileged headers before a valid machine-token request is forwarded. At the ingestion layer, authorization is checked again: a user path is verified against device ownership, while a machine path is verified against the exact device identifier in the payload. This defence-in-depth arrangement is implemented locally, but it is not a demonstration of mTLS, automatic credential rotation, or independent security assurance."),
    ], body, subheading, caption)

    insert_chain(method_second, [
        ("body", "After authorization, Ingestion constructs an EnergyUsageEvent and publishes it to the energy-usage topic. Usage consumes this event as part of the usage-service consumer group and writes a time-series point with deviceId as a tag, energyConsumed as a field, and the event timestamp as the write time. A scheduled task aggregates a preceding one-hour window every ten seconds, resolves household context, compares the sum with configured thresholds, and can publish an energy-alerts event. This describes the implemented route and its intended aggregation semantics; it should not be read as a metering or billing algorithm."),
        ("body", "The usage consumer has a code-level failure policy for time-series persistence. A failed handler invocation receives the initial attempt plus two configurable retries, followed by publication to an energy-usage.DLT record if recovery is required. This is a bounded recovery design, not a completed recovery experiment. A future fault study would need to verify the retry count, dead-letter publication, partition behaviour, operator visibility, and absence of unintended duplicate accounting in the live configuration."),
    ], body, subheading, caption)

    insert_chain(method_boundary, [
        ("body", "The architecture evidence is complemented by dated local observations rather than replaced by them. For example, a 2026-10-11 check recorded unauthorized telemetry rejection, device-token mismatch rejection, accepted owner and device submissions, and the later appearance of 0.37 kWh and 0.19 kWh as 0.56 kWh in the Usage API. These observations support a particular local Gateway-to-Ingestion-to-Kafka-to-Usage-to-InfluxDB route. They do not prove long-duration ordering, sustained consumer behavior, universal fault recovery, or the REFIT path."),
        ("body", "A separate 2026-10-05 single-host run provides historical evidence for controlled alert and load scenarios. It predates later private-caller and device-token hardening, so it is reported only as historical local evidence and not as proof of the final hardened route. Treating the chronology explicitly prevents older runtime observations and later source-level checks from being pooled into a stronger unsupported claim."),
    ], body, subheading, caption)

    insert_chain(architecture_explanation, [
        ("body", "At the public boundary, the frontend interacts with the API Gateway, which routes to the User, Device, Ingestion, Usage, Insight, and Alert services. Keycloak provides the OpenID Connect and JWT identity context for interactive use. The explicit gateway role is significant: service routes are not intended to decide public authentication independently of the boundary context, and internal headers used for trusted downstream calls are not meant to be supplied by an external browser or device."),
        ("body", "The data plane begins at Ingestion, where a validated request becomes an energy-usage event. Kafka isolates that publication from persistence and subsequent consumers. Usage owns time-series writes and rolling sums; InfluxDB stores the resulting time-stamped usage values. MySQL holds the relational facts needed for profiles, devices, and alerts. Alert Service consumes threshold and anomaly events, persists a non-suppressed alert, and may attempt e-mail delivery. Mailpit is a local delivery sink used in dated validation records rather than evidence of a production e-mail service."),
        ("body", "The baseline FastAPI service loads a persisted model and can expose prediction information, including imputed feature names when optional contextual fields are absent. Prometheus and Grafana are configured for service observation. Their presence supports inspectability of selected operational signals, but configured dashboards and metrics are not equivalent to a completed observability or incident-response evaluation. The figure therefore distinguishes the implemented solid path from the dashed REFIT branch, where provisioning, replay, reconciliation, and household-level evaluation remain pending."),
        ("body", "The architectural rationale is separation of responsibilities, not a claim that every service must run independently in a production deployment. The local Docker composition contains multiple supporting services and is resource constrained. Its deployment analysis has shown that the main stack and a future isolated REFIT provisioning stack should not run concurrently on the documented 8 GB workstation."),
    ], body, subheading, caption)

    insert_chain(identity_first, [
        ("body", "The user-facing and device-facing security paths have different inputs and decisions. A user request has an authenticated principal and an ownership rule: a USER may operate only on the household context carried through the trusted gateway boundary, while an ADMIN path is separately authorised. A device request carries a device token in the designated header and is checked against the payload deviceId. The configuration-based mapping is deliberately not equated with a database-generated token lifecycle; device creation persists device metadata and ownership, but does not by itself generate, rotate, or store a device credential."),
        ("body", "The gateway treats caller-supplied internal headers as untrusted. For machine-token traffic, it strips public attempts to set privileged user or gateway-proof headers before forwarding the request. For interactive traffic, the trusted downstream context is derived from the authenticated principal. Ingestion then fails closed if its ownership lookup or token mapping check does not authorize the input. Constant-time comparison is used for token matching. These implementation details make the boundary more explicit, but they do not establish resistance to every threat model or prove production secret-management practice."),
        ("body", "The minimum DTO requires a non-null device identifier, a non-negative energy value, and a timestamp. The validation prevents missing or negative contract values from proceeding through the normal route. It does not validate semantic plausibility, expected household power, or source calibration. In particular, a device identifier cannot be assigned from a REFIT house number merely because the two numbers happen to match; any replay requires separately provisioned experiment profiles, devices, and unique credentials."),
    ], body, subheading, caption)

    insert_chain(identity_second, [
        ("body", "The EnergyUsageEvent carries the device identifier, energyConsumed value, and timestamp. It is not an enriched analytical record and does not contain a durable correlation identifier for every downstream boundary. This fact limits the precision of any current claim about per-record auditability. Future replay work must create an external request log and reconciliation ledger rather than assume that ordinary broker offsets or database rows alone solve end-to-end traceability."),
        ("body", "Source-level tests exercise several expected security and processing behaviours, including rejection before Kafka publication when authorization fails, trusted-header handling, and selected usage-access rules. These tests strengthen confidence in component logic but do not replace a live Keycloak, gateway, broker, database, and e-mail integration study. The paper accordingly labels them as component evidence rather than as a full system-security result."),
    ], body, subheading, caption)

    insert_chain(model_first, [
        ("body", "The legacy training procedure begins with five supplied row-level predictors: Electricity_Consumed, Temperature, Humidity, Wind_Speed, and Avg_Past_Consumption. It derives Hour and DayOfWeek from the timestamp and computes ConsumptionDeviation as Electricity_Consumed minus Avg_Past_Consumption. Records are parsed, invalid numeric fields are removed, negative electricity values and complete duplicates are excluded, and the remaining rows are stably ordered chronologically. The procedure does not construct lagged or rolling features from later observations, but the causal provenance of the supplied Avg_Past_Consumption field remains undocumented."),
        ("body", "The data are divided into 3,500 training rows, 750 validation rows, and 750 held-out test rows. Validation F1 selects the decision threshold over the documented candidate range, yielding 0.31 for the retained Random Forest. The held-out test is not used to select that threshold. The wider trainer contains a deterministic test PR-AUC tie-breaker in its model-selection logic, so the paper avoids the broader unsupported statement that test data played no role in every model-selection decision."),
        ("body", "At inference time, a Wokwi-compatible payload may omit temperature, humidity, wind speed, and Avg_Past_Consumption. The service fills absent contextual values with training medians, derives temporal features from the supplied timestamp, obtains a probability from the fixed artifact, and applies the stored threshold. It reports the imputed feature names so that a three-field payload is not silently mistaken for a fully observed feature vector. This interface decision is transparent, but median imputation is not an evaluation-quality substitute for a real REFIT feature contract."),
    ], body, subheading, caption)

    insert_chain(model_second, [
        ("body", "The retained model configuration is 300 trees, min_samples_leaf = 2, class_weight = balanced_subsample, and random_state = 42. The saved training artifact can use broad parallelism while the deployed inference process limits the effective worker count to reduce oversubscription on the local host. This operational setting is a resource-management decision, not an optimisation result. The paper reports it only to explain why a local in-process benchmark cannot be extrapolated to the total request path."),
    ], body, subheading, caption)

    insert_chain(model_alert, [
        ("body", "Alert Service receives both threshold-generated energy-alerts and ML-generated anomaly-results. A normal model decision creates no anomaly alert. For an alert candidate, the service forms a key from the user context, a device or household scope, and alert type; it searches for a recent matching alert and suppresses a duplicate inside the configured 15-minute window. A non-suppressed alert is persisted, and the e-mail attempt/sent result is recorded separately. Thus, persistence and e-mail delivery are not conflated in the implementation."),
        ("body", "A focused mock-based test on 2026-10-10 executed first-alert, in-window duplicate, post-window, and user/device separation cases with four passing tests. A dated 2026-10-05 local duplicate scenario also recorded two accepted synthetic submissions, one stored ML_ANOMALY alert, and a suppression log. These are useful bounded checks, but they do not prove cluster-safe atomic deduplication, arbitrary replay behaviour, full database failure recovery, or REFIT-specific notification correctness."),
    ], body, subheading, caption)

    insert_chain(refit_first, [
        ("body", "REFIT is used as a recorded-household source for a future experiment, not as a direct replacement for the legacy labelled CSV. The five local cleaned files were checked against official Zenodo MD5 values before their headers and quality characteristics were inspected. The primary cohort consists of Houses 1, 2, and 4 because each has a selected 14-day window satisfying the pre-registered quality gates and no documented solar-interference caveat applied to its primary record. Houses 3 and 11 are retained as a solar-sensitivity cohort because the source documents solar-generation interference and no net-export correction has been performed."),
        ("body", "The cleaned files contain a human-readable Time field, a numeric Unix timestamp, Aggregate, appliance channels, and Issues context. Unix is used to establish chronological order because the human-readable timezone convention has not been independently verified for this analysis; the Time string is retained verbatim. Aggregate is treated as whole-house active power in watts. It is not silently relabelled as energy, and the Issues column is not converted into an anomaly label."),
        ("body", "Each selected pair of adjacent records yields a candidate interval. The method uses the aggregate power at the start of the interval and the observed elapsed duration, so it does not blindly multiply every row by the nominal eight-second cadence. Only intervals with 0 < delta_t <= 16 s and a non-negative start power are included. A longer cross-gap interval is excluded rather than filled, capped, interpolated, or synthesised. This policy avoids hiding the energy contribution of an unobserved period behind an unsupported assumption."),
    ], body, subheading, caption)

    insert_chain(refit_formula_intro, [
        ("body", "In dimensional terms, watts multiplied by seconds produces joules. Dividing by 3,600,000 converts watt-seconds to kilowatt-hours. The left-rectangle/sample-hold interpretation assumes the first aggregate sample represents the following observed interval; it is a documented numerical estimate, not proof that REFIT readings are point values rather than interval averages. This semantic limitation is preserved for the later compatibility review."),
    ], body, subheading, caption)

    insert_chain(refit_condition_intro, [
        ("body", "The maximum 16-second rule permits a modest tolerance around the nominal cadence while preventing the integration of an unobserved long gap as though it were a continuous measurement. For each excluded candidate, the source record and reason remain in the audit output. At the end of a selected window, the final record has no successor and is likewise retained in the source audit rather than contributing an invented interval."),
    ], body, subheading, caption)

    insert_chain(refit_exclusions, [
        ("body", "The generated artifacts include cleaned primary and sensitivity cohorts, a manifest, per-household quality report, timestamp-interval report, excluded-record log, and preprocessing summary. IQR fences flag unusual values diagnostically but do not remove a valid aggregate value. This distinction matters: an unusual power level may be an appliance or household event, a sensor condition, or simply a legitimate observation. It cannot be called an anomaly label without a verified task-specific annotation policy."),
        ("body", "The primary output contains 468,847 retained intervals across the three windows. Its SHA-256 hash is recorded separately from the five source-file MD5 values. Repeating the preprocessing against the same checksum-verified files and configuration produced byte-identical output artifacts. Source integrity, derived-artifact identity, and live system acceptance are therefore three different claims, and only the first two are completed here."),
    ], body, subheading, caption)

    insert_chain(algorithm_last, [
        ("heading", "Algorithm 2. Telemetry Admission and Event Publication"),
        ("body", "Input: an HTTP telemetry request containing deviceId, energyConsumed, timestamp, and either an authenticated user context or a device-bound token. Output: a rejected request or an authorised energy-usage event published to Kafka."),
        ("body", "1. Reject a request with missing required fields or a negative energyConsumed value. 2. If a machine token is presented, remove untrusted caller-supplied privileged headers and verify that the token matches the payload deviceId. 3. Otherwise, require trusted gateway context and verify that the authenticated user owns the payload device. 4. Reject the request if either authorization path fails or if the required ownership lookup cannot be completed. 5. Construct an EnergyUsageEvent with deviceId, energyConsumed, and timestamp. 6. Publish the event to the energy-usage topic. 7. Return the ingestion result without claiming that downstream persistence, inference, or alerting has already completed."),
        ("body", "This algorithm expresses the source-level service contract. It is not a replay algorithm and does not introduce an idempotency key, unit conversion, or a cross-service correlation identifier that the current contract does not contain."),
    ], body, subheading, caption)

    set_text(workflow_heading, "3.7 Evidence and Reproducibility Workflow", subheading, bold=True)
    insert_chain(workflow_explanation, [
        ("heading", "3.8 Reproducibility Protocol"),
        ("body", "The historical baseline protocol records a copied 5,000-row input, the stable chronological split, fixed preprocessing, model configuration, validation threshold rule, and artifact-hash comparison. A rerun reproduces the reported metrics exactly and matches the stored Random Forest artifact hash. This establishes a repeatable local procedure under the available environment. It does not verify the external source, physical units, labels, or causal availability of all predictor values."),
        ("body", "The REFIT protocol records a different reproducibility chain: official local-file MD5 verification, declared source columns, a cohort-selection file, a preprocessing configuration, an excluded-record log, quality reports, derived SHA-256 values, and an independent rerun. The byte-identical result demonstrates deterministic transformation of the selected data. It does not demonstrate a Kafka replay, database write, live alert, or anomaly evaluation."),
        ("body", "For this reason, the workflow figure should be read as an evidence map rather than a claim that every arrow has been executed with REFIT. The solid and completed branches have different evidence sources from the dashed future branch. This presentation permits the next experiment to be audited against explicit inputs and acceptance checks."),
    ], body, subheading, caption)

    # Results and discussion.
    insert_chain(conditions_first, [
        ("body", "The evidence was collected at different times and with different scopes. Dated local Docker checks on 2026-10-05 covered controlled synthetic integrations, a single-host load scenario, alert cases, and selected recovery observations. Later source hardening and component tests were recorded on 2026-10-10 and 2026-10-11. Offline supervised reconstruction and input-contract checks were recorded on 2026-10-15. REFIT preprocessing was performed separately. The study does not combine these dates into a single homogeneous benchmark because the deployed configuration and tested boundary differed."),
        ("body", "This chronology changes the interpretation of a pass. For example, a pre-hardening local load result gives bounded evidence about the older local endpoint, whereas later component tests show expected filter and authorization behaviour without proving a live Docker route. Both are useful, but neither can be relabelled as a current production-capacity or end-to-end security result."),
    ], body, subheading, caption)

    insert_chain(components_first, [
        ("body", "The components are grouped by responsibility rather than treated as interchangeable infrastructure. The gateway and Keycloak form the public identity and routing boundary. User and Device services hold household/device context; Ingestion accepts authorised telemetry; Kafka transports energy-usage events; Usage writes and aggregates time-series values; InfluxDB stores those values; the ML service exposes the retained inference contract; Alert Service persists and sends alert outcomes; Insight Service supplies deterministic user-facing calculations; and Prometheus/Grafana provide monitoring. MySQL supports transactional state including users, devices, and alerts."),
        ("body", "A dated runtime validation found the gateway health endpoint, Keycloak discovery, frontend, Prometheus readiness, and Grafana login route reachable in the local composition. It also confirmed selected authorization outcomes and the 0.56 kWh controlled usage result. An optional ML container was intentionally not included in that particular 2026-10-11 validation because the model directory was inaccessible to Docker on the workstation. The paper therefore treats ML endpoint evidence separately from broader service availability."),
    ], body, subheading, caption)

    insert_chain(historical_first, [
        ("body", "The copied legacy CSV contains 5,000 records with 4,750 Normal and 250 Abnormal labels. The chronological partitions contain both classes: 3,335 Normal and 165 Abnormal rows in training, 705 and 45 in validation, and 710 and 40 in testing. The reconstruction also found zero timestamp, complete raw-record, and engineered-feature fingerprint overlap across the partitions. These checks reduce obvious row leakage under the stored protocol, but they do not resolve how Avg_Past_Consumption was constructed before the CSV entered the repository."),
        ("body", "The logistic-regression comparator used the same held-out split and its own validation-selected threshold of 0.28. Table III is thus an internal comparison under one supplied dataset rather than a contest against the external methods in Section 2. The Random Forest's higher F1, ROC-AUC, and PR-AUC are meaningful only within this legacy protocol and should not be interpreted as an estimate of utility fraud detection, appliance anomaly detection, or performance on recorded REFIT households."),
    ], body, subheading, caption)

    insert_chain(reproduction_first, [
        ("body", "The offline reconstruction verifies more than a single displayed score. All reported metric differences between the stored and rerun comparison files were zero, and the reproduced Random Forest artifact hash matched the legacy artifact. On the fixed copied held-out CSV, a stratified bootstrap calculation also yielded a conditional 95% F1 interval of 0.838300 to 0.957444. This interval describes resampling variability of that particular held-out dataset and fixed model; it is not a population confidence interval for unseen households."),
        ("body", "Five in-process validation cases passed for the ML input contract. A separate six-case synthetic Wokwi-scale check found that five values lay above the historical Electricity_Consumed maximum of 1. These results show that syntactically acceptable interface values can be outside the observed historical range. They do not tell whether the anomaly decision is correct for those values and motivate a future out-of-distribution or contract-range policy."),
        ("body", "For the actual three-field smoke request, deviceId 1, energyConsumed 0.55, and timestamp 2026-10-11T00:00:00Z produced anomaly = false, probability = 0.0, threshold = 0.31, and four explicit imputed contextual fields. The observation establishes that the fixed model loaded and processed the input according to the service's imputation rule. It does not measure the distribution of response times, broker-to-alert operation, or classifier accuracy."),
    ], body, subheading, caption)

    insert_chain(refit_results_first, [
        ("body", "The primary-cohort summaries must be read as interval counts rather than independent household examples. H1 contains 153,578 retained intervals and an estimated 104.820956111 kWh after 184 excluded candidates; H2 contains 154,985 retained intervals and an estimated 170.354373889 kWh after 112 exclusions; and H4 contains 160,284 retained intervals and an estimated 111.223081389 kWh after 40 exclusions. The excluded counts are dominated by cross-gap candidates plus the terminal record without a successor, not by a model decision."),
        ("body", "The selected H1, H2, and H4 windows passed the stated reconciliation and timestamp-quality checks. Native-cadence shares were 95.97%, 95.64%, and 99.30%, respectively. Their IQR diagnostics flag many unusual aggregate values but do not remove them; this preserves the distinction between a numerical distributional flag and a verified anomaly label. H3 and H11 remain outside the primary summary because documented solar-generation interference has not been corrected to a consumption-only or net-import signal."),
        ("body", "The byte-identical rerun and primary SHA-256 establish that another execution using the same verified inputs and configuration reproduced the same derived files. This is strong evidence for preprocessing repeatability, not a claim that the interval count represents 468,847 independent evaluations or that the estimated 386.398411389 kWh has passed the live energyConsumed contract."),
    ], body, subheading, caption)

    insert_chain(discussion_rf, [
        ("body", "The historical comparison supports a narrow conclusion: under the available CSV, feature preparation, chronological split, and threshold protocol, the retained Random Forest outperformed the retained logistic-regression comparator. It does not show that Random Forest is universally superior to regression, autoencoding, dynamic-thresholding, or other approaches reviewed in Section 2. The model choice is a legacy artifact result whose transparency is improved by exact reconstruction and explicit limitations."),
        ("body", "The reproducibility result is still useful. An exact repeated score and model hash make the historical baseline inspectable and reduce ambiguity about configuration drift. They cannot repair missing documentation outside the repository. The unresolved source, licence, label-generation method, household identity, physical units, and Avg_Past_Consumption construction remain material threats to any broader scientific interpretation."),
    ], body, subheading, caption)

    insert_chain(discussion_smoke, [
        ("body", "Dated local runtime observations provide modest but genuine evidence beyond static code. The 2026-10-11 controlled inputs demonstrate selected authorization decisions and a Gateway-to-Usage persistence path. Earlier 2026-10-05 local evidence records a 450-request ingestion run with 450 accepted checks, a mean HTTP duration of 39.44 ms, and p95 of 109.77 ms under a ramping single-host scenario. These values are historical local endpoint observations, not an inference benchmark, broker throughput measure, or deployment-scale capacity result."),
        ("body", "The dated local alert evidence is also specific. A controlled threshold case observed an alert with sent = true and a matching Mailpit message after 13.64 s, while a duplicate scenario recorded one stored alert and a suppression message inside the configured window. The later focused mock test covered four logical deduplication cases. Together they support the presence of a duplicate-aware policy and selected local behaviour, but they do not establish exactly-once handling, multi-instance atomicity, arbitrary replay correctness, or a REFIT-driven alert evaluation."),
        ("body", "A controlled fault record observed ingestion acceptance while Mailpit was unavailable and later observed container-state recovery. A separate ML-container fault record did not establish post-restart Kafka reconnection or unique downstream event processing. The results are therefore reported as limited fault observations. They are not used to claim resilience of every service combination or end-to-end recovery of an alert lifecycle."),
    ], body, subheading, caption)

    insert_chain(discussion_refit, [
        ("body", "The REFIT result makes a different contribution. It turns a documented historical household source into a deterministic, auditable interval-energy artifact while retaining the evidence needed to inspect exclusions and quality. The chosen H1/H2/H4 cohort provides a realistic basis for a future chronological stream, whereas H3/H11 make the solar-interference limitation visible instead of hiding it in a pooled total."),
    ], body, subheading, caption)

    insert_chain(discussion_boundary, [
        ("body", "No table in this paper should be read as a substitute for a pre-registered end-to-end study. Before a real REFIT evaluation, the project must verify that derived_interval_energy_kwh maps to the semantic unit expected by energyConsumed without a second conversion; create isolated H1/H2/H4 profiles and device-bound credentials; run a bounded pilot; reconcile submitted, accepted, persisted, and rejected events; and only then evaluate an aligned aggregate label protocol. Without those steps, a REFIT F1-score would not be defensible."),
    ], body, subheading, caption)

    # Insert an objective-to-evidence table after existing Table V, followed by a short discussion.
    table_v = doc.tables[4]
    table_v_anchor_element = table_v._tbl
    new_p_element = OxmlElement("w:p")
    table_v_anchor_element.addnext(new_p_element)
    section_47 = Paragraph(new_p_element, doc._body)
    section_47.style = subheading.style
    copy_ppr(subheading, section_47)
    r = section_47.add_run("4.7 Objective-to-Evidence Interpretation")
    copy_rpr(subheading.runs[0], r)
    r.bold = True
    table_vi_caption = insert_after(section_47, "Table VI. Objective-to-evidence status and remaining validation gates", caption,
                                    alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    table_vi = make_objective_table_after(table_vi_caption, doc)
    p_after_table_element = OxmlElement("w:p")
    table_vi._tbl.addnext(p_after_table_element)
    p_after_table = Paragraph(p_after_table_element, doc._body)
    p_after_table.style = body.style
    copy_ppr(body, p_after_table)
    p_after_table.add_run("Table VI makes the evidence boundary operational. The four objectives are not presented as equally complete: Objectives 1 and 2 have implementation and selected local functional support, Objective 3 has the strongest completed offline evaluation, and Objective 4 has completed source and preprocessing work but no live replay. This mapping prevents the paper from converting a partially completed pipeline into a single undifferentiated success claim.")

    # Replace the compressed limitation lead-ins with proper subsections and expand them.
    set_text(limitation_historical, "5.1 Historical Data Provenance and Generalisation", subheading, bold=True)
    insert_chain(limitation_historical, [
        ("body", "The legacy supervised CSV permits an exact local reconstruction, but the repository does not establish its external source, licence, acquisition conditions, physical units, household identity, or label-generation process. These absences limit scientific interpretation even when the model code and metrics are deterministic. A score obtained under an opaque data-generating process cannot be treated as a population estimate for household anomalies."),
        ("body", "Avg_Past_Consumption is a second provenance concern. The code correctly computes ConsumptionDeviation from the supplied history field and avoids creating lag features from future rows, yet it cannot demonstrate whether the supplied history value was causal at the original point of prediction. This uncertainty matters because the model may benefit from a feature that would not be available under a live household deployment. Future work requires a documented causal feature construction before model transfer or retraining is claimed."),
        ("heading", "5.2 Simulator-Compatible Telemetry and Physical Measurement"),
        ("body", "Wokwi-compatible requests are useful for software integration demonstration because they exercise the application contract without requiring physical hardware in every test. They do not calibrate voltage, current, power factor, electrical noise, sampling error, clock drift, or meter installation. Consequently, the project does not make a utility-billing or physical-meter accuracy claim. Any physical validation would require a documented meter, reference instrument, sampling protocol, and comparison procedure."),
    ], body, subheading, caption)

    set_text(limitation_refit, "5.3 REFIT Labels, Source Semantics, and Aggregate Evaluation", subheading, bold=True)
    insert_chain(limitation_refit, [
        ("body", "REFIT provides aggregate power and appliance-level channels, but the current paper does not transform its available annotation evidence into whole-house anomaly ground truth. The official anomaly resource is appliance-level and only House 1 overlaps with the H1/H2/H4 primary cohort; its detailed schema, timestamp semantics, count, and aggregate-alignment method were not verified in the present work. The REFIT Issues field is retained as context, not used as a label."),
        ("body", "The interval-energy value is an actual-delta sample-hold estimate from Aggregate power in watts. It is deliberately not called a direct REFIT kWh measurement. The source reading semantics, possible treatment of a point sample versus interval average, and the mapping into the application's energyConsumed field still require review. A future evaluation must prevent both unit substitution (watts sent as kWh) and a second conversion of an already estimated interval energy."),
        ("heading", "5.4 Runtime Integration and Alert Correctness"),
        ("body", "The prepared REFIT cohort has not been provisioned as three isolated device identities. No H1/H2/H4 device-token mapping, live ingestion run, event reconciliation ledger, or full chronological replay has been completed. Existing synthetic local alert and deduplication checks provide bounded evidence for the application, but they do not establish what will happen to a recorded household stream or to events whose timestamps span days rather than the local test window."),
        ("body", "The current deduplication mechanism is also a policy rather than an exactly-once guarantee. It keys an alert by user, device/household scope, and type within a configured temporal window. It should therefore be evaluated for false suppression, repeated delivery, concurrent consumers, restarts, and database timing before being claimed as a reliable operational guarantee. A future replay must log both suppressed and delivered decisions alongside the corresponding source events."),
    ], body, subheading, caption)

    set_text(limitation_resource, "5.5 Resource and Deployment Constraints", subheading, bold=True)
    insert_chain(limitation_resource, [
        ("body", "The documented machine has 8 GB RAM and a 3.696 GiB Docker allocation, while the main stack was observed at approximately 2.89 GiB. A separately isolated REFIT provisioning stack must therefore not run concurrently with the main stack on this workstation. This is a local experimental-control and safety constraint, not evidence that the architecture cannot run on hardware with appropriate resources."),
        ("heading", "5.6 Performance, Resilience, and Security Evaluation"),
        ("body", "The existing single-host k6 result, focused component tests, local alert checks, and limited fault observations are informative but insufficient for a broad operational claim. The k6 run did not measure distributed load, consumer lag, database saturation, or replay performance. The in-process model benchmark excludes HTTP, Kafka, storage, and alert overhead. Fault records do not prove post-restart reconnection or unique message processing. Security tests do not constitute an independent penetration test, mTLS deployment, secret rotation validation, or a production workload-identity review."),
    ], body, subheading, caption)

    set_text(limitation_future, "5.7 Prioritised Future Validation Plan", subheading, bold=True)
    insert_chain(limitation_future, [
        ("body", "Stage 1 - Contract review: verify the semantic unit and aggregation meaning of energyConsumed against the derived_interval_energy_kwh field. Acceptance requires a documented mapping, payload validation, and no implicit or duplicate conversion. Stage 2 - Isolated provisioning: create a separate non-production identity environment with one profile, one aggregate-meter device, and one unique device credential for each of H1, H2, and H4. Acceptance requires redacted evidence that the mapping is one-to-one and does not reuse an existing household device."),
        ("body", "Stage 3 - Pilot replay: transmit a bounded chronological sample with a dry-run log and conservative retry handling. Acceptance requires reconciliation of submitted, accepted, rejected, ambiguous, broker-consumed, and persisted records. The pilot should stop on authentication, schema, timestamp, or unit disagreement. Stage 4 - Label protocol: obtain and inspect the official annotation schema, define the target aggregate label, verify timestamp alignment, and pre-register the temporal or household split before reporting any accuracy metric."),
        ("body", "Stage 5 - Controlled full replay: replay the H1/H2/H4 primary cohort only after pilot approval, preserve original source time separately from replay time, and report all reconciliations. Stage 6 - Operational evaluation: use controlled duplicate, notification, failure, restart, and input-range scenarios to test alert behaviour; repeat latency and throughput measurements with stated hardware conditions. Each stage is an acceptance gate rather than an automatic continuation of the previous one."),
    ], body, subheading, caption)

    insert_chain(conclusion_second, [
        ("body", "The four objectives are satisfied at different levels. The paper documents an implemented admission and event-processing architecture, including explicit user and device telemetry boundaries. It reconstructs the legacy supervised baseline exactly under the available historical CSV protocol. It also establishes checksum-verified and byte-reproducible REFIT preparation for a defined primary cohort. These outcomes are useful because they make the present state inspectable and reduce ambiguity about what the project has and has not shown."),
        ("body", "The strongest empirical claims remain limited. The reported Random Forest metrics are historical CSV-specific results, the Wokwi-compatible response is an interface smoke test, the dated local system observations are single-host and configuration-specific, and REFIT has not been sent through the application or used for aggregate anomaly accuracy. The system should therefore be understood as a local implementation and reproducibility study rather than a completed real-household deployment."),
        ("body", "The next credible advance is not another generic demonstration. It is a controlled, isolated REFIT evaluation with a verified energy contract, device-bound credentials, pilot reconciliation, a defensible label-alignment protocol, and measured alert, performance, and recovery outcomes. Completing those gates would allow future work to move from transparent preparation and historical reproducibility toward a substantiated household-level anomaly-detection claim."),
    ], body, subheading, caption)

    doc.save(OUTPUT)

    # Minimal structural evidence for the build log.
    built = Document(OUTPUT)
    text = "\n".join(p.text for p in built.paragraphs)
    word_count = sum(len(p.text.split()) for p in built.paragraphs) + sum(
        len(c.text.split()) for t in built.tables for row in t.rows for c in row.cells
    )
    report = OUTPUT.with_suffix(".build.txt")
    report.write_text(
        "Output: " + str(OUTPUT) + "\n"
        + "SHA-256: " + hashlib.sha256(OUTPUT.read_bytes()).hexdigest() + "\n"
        + "Approximate words including table cells: " + str(word_count) + "\n"
        + "Objectives: " + str(text.count("Objective 1:") + text.count("Objective 2:") + text.count("Objective 3:") + text.count("Objective 4:")) + "\n"
        + "Research gaps: " + str(text.count("Research Gap 1:") + text.count("Research Gap 2:") + text.count("Research Gap 3:")) + "\n"
        + "Tables: " + str(len(built.tables)) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
