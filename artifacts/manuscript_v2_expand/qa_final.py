import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document


MANUSCRIPT = Path(
    r"C:\Users\anju2\Desktop\Home-Energy-Tracker\artifacts\manuscript_v2_expand\23BCE0052-23BCE0050-23BCE0080_Expanded_V2.docx"
)
SOURCE_MANUSCRIPT = Path(
    r"C:\Users\anju2\Desktop\Advance_java_research\23BCE0052-23BCE0050-23BCE0080_Strengthened.docx"
)


def text_from_element(element):
    return "".join(element.xpath(".//w:t/text()"))


def count_words(parts):
    return len(
        re.findall(r"[A-Za-z0-9]+(?:[-_'][A-Za-z0-9]+)*", " ".join(parts))
    )


document = Document(MANUSCRIPT)
source_document = Document(SOURCE_MANUSCRIPT)
before_references = []
references = []
found_references = False
for child in document.element.body.iterchildren():
    text = text_from_element(child)
    if child.tag.endswith("}p") and text.strip() == "References":
        found_references = True
        continue
    (references if found_references else before_references).append(text)

with zipfile.ZipFile(MANUSCRIPT) as archive:
    comment_part_exists = "word/comments.xml" in archive.namelist()

print(f"content_words_excluding_references={count_words(before_references)}")
print(f"reference_words={count_words(references)}")
print(f"total_words={count_words(before_references) + count_words(references)}")
print(f"tables={len(document.tables)}")
print(
    "reference_entries="
    + str(sum(1 for text in references if re.match(r"^\[\d+\]", text.strip())))
)
print(
    "to_be_provided_occurrences="
    + str(sum(text.lower().count("to be provided") for text in before_references + references))
)
print(f"comments_xml={comment_part_exists}")
print(
    "title_author_affiliation_block_unchanged="
    + str(
        [paragraph.text for paragraph in document.paragraphs[:7]]
        == [paragraph.text for paragraph in source_document.paragraphs[:7]]
    )
)
print(f"sha256={hashlib.sha256(MANUSCRIPT.read_bytes()).hexdigest()}")
