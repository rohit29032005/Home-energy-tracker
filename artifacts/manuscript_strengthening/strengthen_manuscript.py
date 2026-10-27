from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\anju2\Desktop\Advance_java_research\23BCE0052-23BCE0050-23BCE0080.docx")
OUTPUT = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker\artifacts\manuscript_strengthening\23BCE0052-23BCE0050-23BCE0080_Strengthened.docx")


def copy_run_properties(source_run, target_run):
    if source_run is None or source_run._r.rPr is None:
        return
    target_run._r.insert(0, deepcopy(source_run._r.rPr))


def replace_paragraph(paragraph, text: str):
    """Replace text while retaining the paragraph and first-run formatting."""
    source_run = paragraph.runs[0] if paragraph.runs else None
    paragraph._p.clear_content()
    run = paragraph.add_run(text)
    copy_run_properties(source_run, run)
    return paragraph


def insert_paragraph_after(paragraph, text: str, template_paragraph=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    template = template_paragraph if template_paragraph is not None else paragraph
    if template._p.pPr is not None:
        new_paragraph._p.insert(0, deepcopy(template._p.pPr))
    source_run = template.runs[0] if template.runs else None
    run = new_paragraph.add_run(text)
    copy_run_properties(source_run, run)
    return new_paragraph


def normalise_blank_paragraph(paragraph, template_paragraph):
    """Clear an inherited list paragraph without leaving a visible list marker."""
    paragraph._p.clear_content()
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if template_paragraph._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template_paragraph._p.pPr))


def set_cell_text(cell, text: str, bold: bool = False):
    paragraph = cell.paragraphs[0]
    paragraph._p.clear_content()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(7 if not bold else 7.25)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        tag = f"w:{side}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), "45")
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches: float):
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def build_literature_table(document, old_table, caption_paragraph):
    headers = [
        "Reference and year",
        "Research objective",
        "Method, architecture, or setting",
        "Reported contribution",
        "Relevant boundary and relationship to this study",
    ]
    rows = [
        ["[1], 2024", "Residential energy-efficiency monitoring", "Microservices architecture", "Separates monitoring responsibilities into services.", "Motivates component separation; its deployment findings are not reused as benchmarks."],
        ["[2], 2025", "Home energy-management architecture", "Cloud HEMS conceptual model", "Organises ingestion, operational, and analytical environments.", "Supports the architectural framing; the present study does not claim its cloud-scale results."],
        ["[3], 2026", "Smart-grid energy management", "Hybrid communication with microservices", "Presents and validates a modular smart-grid platform in its own setting.", "Supports decoupled design; communication and deployment metrics are not transferred."],
        ["[4], 2024", "Anomaly detection from smart-meter data", "Daily-activity analysis with a real household meter", "Links anomaly interpretation to a documented measurement and evaluation context.", "Reinforces the need to keep data, labels, and protocol explicit; task differs from the legacy classifier."],
        ["[8], 2024", "Fault and misuse identification", "Smart-meter data analytics and machine learning", "Uses anomaly detection for faults, billing errors, and misuse context.", "Relevant anomaly application context; labels, features, and data are not those of the present study."],
        ["[10], 2024 (preprint)", "Real-time energy-anomaly detection", "Convolutional autoencoder with dynamic threshold", "Presents an alternative model using meter data and an adaptive threshold.", "A preprint and not a direct baseline; no performance comparison is claimed with the legacy Random Forest."],
        ["[11], 2025", "Smart energy monitoring", "IoT-oriented monitoring system", "Provides application-level monitoring context.", "Used only for monitoring context; source-level empirical detail is limited for a direct comparison."],
        ["[12], 2026", "Home energy management and anomaly detection", "Smart plugs, Zigbee/MQTT, and regression analysis", "Reports an IoT HEMS with predictive and anomaly-analysis functions.", "Shows a physical smart-plug approach; it is not a benchmark for simulator-compatible telemetry."],
        ["[13], 2025", "Household consumption monitoring", "ESP32 and PZEM-004T prototype", "Implements physical measurement of electrical parameters and IoT reporting.", "Clarifies the distinction between physical-meter validation and the present Wokwi-compatible interface test."],
        ["[14], 2024", "Cooling-energy conservation", "ANNs-PSO analytics microservice platform", "Combines prediction, optimisation, and smart-home control.", "Control optimisation is outside the present tracker scope; no control-performance comparison is made."],
        ["[15], 2024", "Real-time energy-meter monitoring", "Arduino/GSM smart-meter prototype", "Reports a physical meter with monitoring and alert functions.", "Provides physical-meter context; the present work does not claim meter calibration or SMS-alert validation."],
        ["[17], 2025", "Smart-city energy management", "Distributed microservices, Kafka, and InfluxDB", "Uses service and storage technologies relevant to energy data processing.", "Technology overlap supports design rationale; city-scale results do not establish local residential performance."],
        ["[19], 2026", "Appliance monitoring and anomaly detection", "Fast DDS and multi-task neural inference", "Reports a real-time sensing-to-inference pipeline with a different middleware and model.", "Illustrates a more extensively evaluated pipeline; the present study does not claim comparable appliance-level or field results."],
        ["[20], 2025", "Residential consumption and power-quality monitoring", "IoT smart meter with power-quality analysis", "Combines consumption monitoring with disturbance analysis in a physical prototype.", "Power-quality classification and physical validation are outside the demonstrated scope of the present study."],
    ]

    table = document.add_table(rows=1, cols=len(headers))
    table.style = old_table.style.name if old_table.style is not None else "Table Grid"
    table.autofit = False
    widths = [0.73, 1.02, 1.38, 1.28, 1.85]
    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        set_cell_text(cell, header, bold=True)
        set_cell_width(cell, width)
    set_repeat_table_header(table.rows[0])

    for row_values in rows:
        row = table.add_row()
        for cell, value, width in zip(row.cells, row_values, widths):
            set_cell_text(cell, value)
            set_cell_width(cell, width)
        set_cant_split(row)

    caption_paragraph._p.addnext(table._tbl)
    old_table._element.getparent().remove(old_table._element)
    return table


def main():
    document = Document(SOURCE)
    p = document.paragraphs

    replacements = {
        9: "Residential energy monitoring requires a traceable path from telemetry admission to storage, analysis, and controlled notification. This paper presents a local Home Energy Tracker implemented as event-driven microservices with gateway-mediated access control, device-bound telemetry tokens, Apache Kafka, InfluxDB, MySQL, deterministic insights, a duplicate-aware alert path, a legacy Random Forest inference service, and service observability. The repository classifier was previously evaluated on a 5,000-record labelled CSV and reproduced offline; its held-out F1 score of 0.902439 is therefore historical-dataset evidence, not REFIT, field, or household-level accuracy. A Wokwi-compatible request confirms that the inference interface can load the model and return a response, but it is only a smoke test. For data provenance, five official REFIT cleaned files were checksum-verified and the H1/H2/H4 primary cohort was transformed into 468,847 interval-energy records, totalling 386.398411389 kWh under the documented actual-delta sample-hold rule. H3/H11 remain a separate solar-affected sensitivity cohort. End-to-end REFIT ingestion, replay, alert-deduplication validation, and REFIT anomaly-accuracy evaluation are pending. The contribution is thus an implemented, reproducibility-oriented system with explicit evidence limits.",
        12: "Residential energy monitoring is more than a display of consumption: useful systems must admit telemetry, preserve an audit trail, derive operational insights, and issue bounded notifications without being mistaken for calibrated utility metering. Microservice architectures provide a practical way to separate these responsibilities into inspectable components [1], [2].",
        13: "This study concerns a local Home Energy Tracker, not a calibrated utility-metering product. Simulator-compatible telemetry is used to examine software-path execution, reproducibility, and data preparation; it does not provide physical sensor calibration or utility-billing accuracy. Physical-meter prototypes therefore represent a separate validation category [13], [15].",
        14: "An event-driven design decouples telemetry admission from downstream aggregation, storage, inference, and alert processing. Recent energy-management architectures use modular services and asynchronous communication for this separation [1], [3]. Their deployment conditions are not treated as numerical baselines for the present local prototype.",
        15: "In the present system, Keycloak supplies JWT-, role-, and ownership-based control for user-facing requests, while simulator-compatible telemetry is admitted through a device-bound token. These mechanisms define an admission boundary before event processing begins.",
        16: "Validated telemetry is published to Kafka. The Usage Service records time-stamped usage and windowed aggregates; InfluxDB stores time-series usage, while MySQL stores household, device, and alert state. This separation makes the processing path inspectable and keeps relational state distinct from time-series data.",
        17: "The implementation also exposes deterministic bill forecasting, appliance ranking, standby detection, saving tips, and related insights. Their usefulness depends on being traceable to admitted telemetry, so this study prioritises record handling and evidence boundaries over dashboard presentation alone.",
        18: "Smart-meter anomaly research links a model result to the measurement source, label definition, feature construction, and evaluation protocol. Studies of daily activities, faults or misuse, dynamic anomaly detection, and real-time AI pipelines therefore provide context for interpreting—not inflating—the evidence of the present classifier [4], [8], [10], [19].",
        19: "The installed Random Forest is a legacy repository artefact. Although its local reconstruction follows a documented protocol, the historical labelled CSV has incomplete source, licence, label-generation, and causal-feature provenance. Its metrics are consequently treated as reproducibility evidence rather than evidence of generalisation.",
        20: "REFIT provides a more thoroughly documented household-data source for future evaluation. The descriptor and cleaned-data records establish measurement provenance, but aggregate-power preparation and anomaly-label alignment still require verification before the data can support end-to-end anomaly claims [5]–[7].",
        21: "This paper documents the implemented architecture, a reproducible historical baseline, and checksum-verified REFIT preparation. It distinguishes those completed activities from isolated provisioning, replay, alert-deduplication validation, and REFIT anomaly-accuracy evaluation, which remain future validation work.",
        23: "Objective 1: Design a secure microservices data pipeline with telemetry validation, device-token checking, inspectable service boundaries, and invalid-message handling.",
        24: "Objective 2: Integrate Kafka with the microservices to support event-driven telemetry processing, baseline inference, alert handling, and observability for failures, replay, consumer lag, and tracing.",
        25: "Objective 3: Establish a reproducible Random Forest baseline by recording the legacy features, random seed, and chronological split, while clearly separating offline metrics from real-world validation.",
        26: "Objective 4: Prepare and integrity-check REFIT records, document data-quality outcomes, and define evaluation steps for future end-to-end anomaly-detection validation.",
        30: "Rovetto et al. [1] use a microservices architecture for residential energy-efficiency monitoring. Their work motivates the separation of monitoring responsibilities; the present study applies that principle to a local prototype and does not reuse the reported deployment outcomes as performance evidence.",
        31: "Motta et al. [2] organise home-energy functions around ingestion, operational, and analytical environments. This layered view helps frame the present gateway, event-processing, storage, inference, and alert components, while the present study makes no claim to replicate the referenced cloud-scale setting.",
        32: "Veloso et al. [3] combine hybrid communication with a modular microservice platform for smart-grid energy management. Their work demonstrates why communication, processing, and service responsibilities should be separated, whereas the present study confines its claim to local service implementation and leaves communication-scale and fault-tolerance evaluation outside the reported evidence.",
        33: "At the anomaly-detection level, Hernández et al. [4] connect smart-meter measurements with a defined evaluation context for activity anomalies. That connection is important here because the legacy CSV's labels and provenance are incomplete: its reported metrics can reproduce a historical experiment, but cannot establish household-level anomaly accuracy.",
        34: "Kulkarni et al. [8] frame anomaly detection in smart electric meters around faults, billing errors, and misuse. Their application context is relevant, but its data, labels, and method do not constitute a benchmark for the fixed legacy Random Forest used in this study.",
        35: "Maitra et al. [10] describe a convolutional-autoencoder approach with a dynamic threshold for real-time meter data. It provides an alternative modelling strategy, but it is an arXiv preprint and is not used as a direct performance comparator for the historical Random Forest.",
        36: "IoT-monitoring studies also focus on user-facing and physical-meter settings. Karthika et al. [11] provide smart-energy-monitoring context; Liu and Hsiao [12] report smart-plug-based home energy management with regression-based anomaly detection; and Hidayat and Al-Faris [13] use an ESP32/PZEM-004T monitoring prototype. These physical-monitoring contexts underscore why the Wokwi-compatible request in this study must not be presented as calibrated meter evidence.",
        37: "Sirisumrannukul et al. [14] integrate ANNs and particle-swarm optimisation in a cooling-management microservice platform, making optimisation and autonomous control central to their contribution. Hossain and Hassan [15] report a real-time smart-energy-meter system, and Zhu and Yue [17] describe smart-city energy management using technologies that include Docker, Kafka, and InfluxDB. These studies provide complementary application and architectural context, but their deployment settings and outcomes are not transferred to the present local tracker.",
        38: "Shuaibu et al. [19] present a smart-metering pipeline that uses DDS middleware and multi-task neural inference for appliance monitoring and anomaly detection. De Sousa and Rabelo [20] combine residential consumption monitoring with power-quality analysis. Both works illustrate broader operational goals than those demonstrated here: the present system does not claim appliance-level inference, power-quality classification, field calibration, or comparable real-time performance.",
        39: "REFIT is treated separately from these research studies because it is a dataset and data-provenance resource. Murray et al. [5] document the household measurements, while the University of Strathclyde and Zenodo records describe cleaned data access [6], [7]. The present work uses these sources for checksum verification, cohort selection, and repeatable interval-energy preparation rather than as a source of validated aggregate anomaly labels.",
        40: "Taken together, the reviewed work establishes the value of modular energy architectures, IoT monitoring, and anomaly-analysis methods. It also shows that evidence is tied to a particular data source, measurement system, and protocol. The present study therefore contributes an evidence-aware local architecture and a reproducible data-preparation path, while retaining clear boundaries around its untested runtime and external-validation stages.",
        41: "The comparison in Table I summarises representative studies by objective, approach, contribution, and the evaluation boundary relevant to the present study.",
        42: "Table I. Comparative summary of representative research on home energy monitoring and anomaly detection",
        45: "Research Gap 1: Traceability across the monitoring pipeline",
        46: "Prior work establishes modular cloud and microservice designs for energy data, including real-time monitoring deployments [1]–[3], [17], [19]. However, the reviewed literature does not always report a record-level, inspectable connection from telemetry admission through broker delivery, persistence, inference, and alert decision in the setting of a local residential prototype. This matters because a missing or duplicated alert cannot otherwise be localised to admission, delivery, storage, inference, or policy handling. The present architecture exposes these boundaries through the gateway, Kafka, services, InfluxDB, and MySQL; end-to-end reconciliation, duplicate suppression, delivery, and fault recovery under recorded-data replay remain unverified.",
        47: "Research Gap 2: Separation of reproducibility from model validity",
        48: "Smart-meter studies report outcomes under their own datasets and protocols [4], [8], [10], [12], [19]. What is not interchangeable is the meaning of a historical metric, an interface response, and a source-integrity check. This study separates the historical Random Forest evaluation, the Wokwi-compatible smoke test, and REFIT preprocessing so that a repeated training run, endpoint response, or checksum is not misrepresented as household-level anomaly accuracy. The remaining limitation is that this separation improves interpretation, not external model validity.",
        49: "Research Gap 3: Reproducible recorded-household preparation and later runtime evaluation",
        50: "REFIT establishes documented household measurements and cleaned-data records [5]–[7], but a reproducible recorded-household experiment also requires explicit cohort selection, timestamp handling, power-to-energy transformation, source integrity, device provisioning, and label alignment. These requirements matter because file integrity alone does not show that authenticated ingestion succeeds, that events reconcile through the runtime, or that an anomaly label belongs to the aggregate reading being evaluated. The present study contributes checksum-verified inputs, a defined H1/H2/H4 primary cohort, a separate solar-affected H3/H11 sensitivity cohort, and deterministic interval-energy artefacts. Device provisioning, live replay, aggregate-label validation, and REFIT accuracy remain outside the current evidence.",
        51: "Synthesis of Research Gaps",
        52: "The three gaps are connected: inspectable software boundaries support traceability; evidence separation prevents an architecture or smoke test from being mistaken for model validation; and reproducible REFIT preparation supplies a controlled basis for a later recorded-household experiment. Together they motivate the four objectives: implement the pipeline, provide observability, reproduce the legacy baseline without overstating it, and prepare REFIT data for future validation.",
        53: "Problem Statement",
        54: "The research problem is to document and evaluate a residential energy-monitoring system in a way that preserves the distinction between implemented software, reproduced historical results, data-integrity evidence, and unperformed runtime experiments. The system must support simulator-compatible telemetry, event processing, storage, analytics, inference, and alert handling through inspectable boundaries, while REFIT preparation must remain reproducible and ready for a later controlled end-to-end evaluation.",
        59: "The system accepts a Wokwi-compatible telemetry payload containing deviceId, energyConsumed, and timestamp. User-facing requests are governed by JWT, role, and ownership checks; device telemetry uses a device-bound token. The Ingestion Service validates the telemetry contract before publishing energy-usage events to Kafka.",
        60: "The Usage Service consumes energy-usage events and records time-stamped usage together with defined-period aggregates. InfluxDB stores time-series usage, whereas MySQL stores household, device, and alert state. The baseline inference service consumes the appropriate event type and returns anomaly results. The Alert Service implements threshold and anomaly pathways with a configured 15-minute deduplication window. Prometheus and Grafana provide service observation.",
        61: "Architecture implementation, historical-baseline reproduction, live interface testing, and REFIT preprocessing are treated as separate evidence streams. This prevents a successful service response or one inference result from being counted as an accuracy, reliability, or alert-correctness result.",
        64: "Figure 1 presents the implemented Home Energy Tracker architecture and the separately prepared, but not runtime-connected, REFIT branch.",
        65: "The solid path runs from simulator-compatible telemetry through gateway admission, ingestion, Kafka, usage aggregation, time-series storage, baseline inference, alert processing, relational state, and observability. The dashed REFIT path represents preparation for isolated provisioning, replay, reconciliation, and labelled evaluation; those stages have not yet been completed.",
        66: "Table II. System components and evidence boundaries",
        69: "The gateway defines the API boundary: user-facing routes use JWTs with role and ownership checks, whereas simulator-compatible device telemetry uses a device-bound token. A REFIT household identifier is not assumed to equal a backend device identifier. Any future replay must provision households, devices, and credentials separately.",
        70: "After contract validation, energy-usage events are published to Kafka for downstream aggregation and baseline inference. Decoupling enables component-level inspection, but component availability or Kafka connectivity does not by itself prove end-to-end event reconciliation.",
        72: "The retained Random Forest artefact, random_forest-csv-20261003T205417Z, uses Electricity_Consumed, Temperature, Humidity, Wind_Speed, Avg_Past_Consumption, Hour, DayOfWeek, and ConsumptionDeviation. The historical protocol uses a chronological split of 3,500 training, 750 validation, and 750 test records; the decision threshold is selected from validation F1.",
        73: "The model uses 300 trees, min_samples_leaf = 2, class_weight = balanced_subsample, random_state = 42, and a 0.31 threshold. For Wokwi-compatible payloads, four contextual inputs are median-imputed and temporal features are derived from the timestamp. The resulting response is an interface smoke test, not a statistical evaluation of the historical model.",
        74: "The Alert Service has a configured 15-minute deduplication window for threshold and anomaly pathways. No replay or fault experiment has yet established duplicate suppression, notification delivery, latency, recovery behaviour, or fault tolerance.",
        76: "The five local cleaned REFIT files (H1, H2, H3, H4, and H11) were checked against the official Zenodo MD5 values. Aggregate total-house active power is expressed in watts. H1, H2, and H4 form the primary cohort, while H3 and H11 remain a solar-affected sensitivity cohort because no solar or net-export correction was applied. Source-file MD5 verification is distinct from the SHA-256 hash recorded for the derived primary output.",
        77: "For each eligible pair of consecutive records, interval energy is estimated from the first aggregate-power sample and the actual elapsed time. The following formula preserves the documented method and units.",
        79: "Only intervals satisfying the condition below are retained.",
        81: "Cross-gap intervals are excluded rather than interpolated, filled, capped, or synthesised. Negative-power records are excluded from the combined interval-energy output but retained in the source audit. IQR flags are diagnostic only, and REFIT Issues values are treated as contextual information rather than anomaly labels.",
        84: "Input: a cleaned, verified REFIT household file for a selected 14-day period, checked against its MD5 checksum.",
        85: "Sort records by preserved Unix timestamp and retain the associated Time value.",
        86: "Compute the actual elapsed interval Δt for every adjacent pair.",
        87: "If 0 < Δt_s ≤ 16 s and the first aggregate-power sample is non-negative, apply the left-rectangle/sample-hold rule to estimate interval energy.",
        88: "Otherwise, retain the source record for audit and record the interval in the excluded-record log; do not interpolate, cap, or synthesise values.",
        89: "Produce the cleaned cohort, preprocessing summary, quality report, timestamp-interval report, and excluded-record log.",
        90: "Create the H1/H2/H4 primary cohort and retain H3/H11 as a separate solar-affected sensitivity cohort.",
        93: "Figure 2. Evidence and reproducibility workflow. The REFIT runtime phase remains pending until isolated device provisioning, credential creation, pilot reconciliation, and verified evaluation are completed.",
        94: "The upper workflow records the historical 5,000-record protocol, validation-threshold selection, offline reproduction, and endpoint smoke-test boundary. The lower workflow records official REFIT verification, cohort separation, actual-delta normalisation, and byte-identical repeatability. The final stage is deliberately dashed because ingestion and anomaly evaluation have not yet occurred.",
        99: "The reported observations come from distinct evidence classes. The architecture identifies implemented services and event paths; the historical analysis concerns the supplied labelled CSV; offline reproduction checks whether the historical experiment can be repeated; and REFIT preprocessing checks source integrity and deterministic preparation. None of these observations alone is a new anomaly-accuracy result. In particular, a service response or repeated training run does not establish accurate operation on household data.",
        100: "The Windows laptop used for local work has 8 GB RAM and a Docker memory limit of 3.696 GiB; the observed main-stack requirement is approximately 2.89 GiB. The isolated REFIT provisioning stack should therefore not run concurrently with the main stack on this machine. This resource constraint is a deployment limitation, not evidence of scientific validation.",
        102: "The local composition includes an API gateway, Keycloak, device-token admission, Kafka, InfluxDB, MySQL, a FastAPI Random Forest inference service, the Alert Service, Prometheus, and Grafana. Version numbers not established by the available evidence are not assumed. The legacy supervised CSV is repository material, whereas REFIT is an official recorded-household source with separate dataset records [5]–[7].",
        104: "The historical protocol evaluates the original Random Forest with a chronological 70/15/15 split: 3,500 training, 750 validation, and 750 test records. The threshold is selected using validation F1. The logistic-regression comparator uses the same held-out split with a threshold of 0.28.",
        106: "The deterministic trainer was run on the 5,000 copied historical records. All reported metric differences were zero, and the SHA-256 hash of the reproduced Random Forest artefact matched the legacy artefact, showing that the locally available historical experiment was reproduced. A separate Wokwi-compatible request used deviceId 1, energyConsumed = 0.55, and timestamp 2026-10-11T00:00:00Z. Its response tests model loading and payload handling only; it is not a latency, alert-delivery, deduplication, recovery, or accuracy test.",
        108: "The H1, H2, and H4 primary-cohort windows are reported in Table IV. The preprocessing pipeline was run twice with the same five validated source files, criteria, and configuration. Both artefact sets were byte-identical and included the cleaned primary and solar-sensitivity cohorts, manifest, excluded-record log, quality report, preprocessing summary, and timestamp-interval report.",
        112: "Table III. Historical baseline, offline reproduction, and live smoke-test evidence",
        114: "Table IV. REFIT primary-cohort normalisation results",
        116: "Here, aggregate denotes total household active power in watts. Interval energy is estimated by the actual-delta, left-rectangle, sample-hold method. H3 and H11 are excluded from the primary summary because of the documented solar-interference sensitivity.",
        119: "Table III shows that the Random Forest exceeds the logistic-regression comparator on historical held-out F1, ROC-AUC, and PR-AUC. Exact offline reproduction and a matching artefact hash establish that the legacy experiment can be regenerated from the locally available copied input and configuration. They do not resolve the historical CSV's missing provenance, licence, label-generation procedure, household identity, or the causal construction of Avg_Past_Consumption.",
        120: "The smoke test shows that the deployed interface can load the fixed model and handle a Wokwi-compatible payload with declared imputation. It provides no response-time distribution, throughput, alert-delivery, deduplication, recovery, or statistical-accuracy evidence. Five of six examined simulator boundaries also fall outside the historical 0–1 Electricity_Consumed range, motivating future input-range and out-of-distribution controls.",
        121: "REFIT preprocessing produced a traceable primary cohort of 468,847 retained intervals and an estimated 386.398411389 kWh across three selected 14-day periods. The selected intervals contained no duplicate or out-of-order Unix timestamps, and each household satisfied the reported cadence, coverage, missing-slot, and maximum-gap checks. The two preprocessing outputs were byte-identical; the primary-output SHA-256 is dfdba435f35a1548661309fd67b267548f419edd489cacf25c7566098f189163.",
        122: "This result is a preparation and reproducibility result, not a model-effectiveness claim. REFIT appliance-level annotations are not treated as aggregate anomaly ground truth: only H1 overlaps the selected primary cohort, and annotation semantics, counts, timestamps, and alignment with aggregate readings remain unverified. The derived interval-energy mapping has also not been accepted by the live ingestion API.",
        123: "Table V. Evidence boundaries and current validation status",
        127: "Historical CSV provenance and simulator boundary. The 5,000-record labelled CSV does not provide enough information to independently establish its complete provenance, licence, label-generation procedure, household identity, or the causal construction of Avg_Past_Consumption. Its historical metrics, including their exact offline reproduction, are reproducibility evidence for that legacy dataset rather than generalisation evidence. Likewise, Wokwi-compatible requests demonstrate interface execution only; they do not calibrate a physical sensor, verify utility-bill accuracy, or validate real-household telemetry.",
        128: "REFIT labels and runtime integration. REFIT provides documented household aggregate-power records, but file integrity, cohort preparation, and byte-identical transformation do not establish aggregate anomaly ground truth. Annotation alignment, timestamp semantics, and the relationship between available context records and selected aggregate intervals still require verification. Runtime integration remains incomplete: isolated H1/H2/H4 profiles, device credentials, data-contract mapping, pilot replay, event reconciliation, duplicate suppression, alert delivery, and fault/recovery behaviour have not been experimentally tested.",
        129: "Resource, performance, and security boundary. The 8 GB Windows laptop, 3.696 GiB Docker limit, and approximately 2.89 GiB observed main-stack requirement prevent safe concurrent operation of the isolated REFIT provisioning stack and the main stack. Comprehensive latency, throughput, scalability, fault-tolerance, and independent security tests are also absent. These constraints limit the scope of the present evidence and do not indicate a production deployment.",
        133: "This study documents a local, evidence-aware Home Energy Tracker with gateway-mediated admission, Kafka event processing, time-series and relational storage, deterministic insights, a legacy Random Forest inference service, a duplicate-aware alert path, and observability. The historical classifier was reproduced offline with its dataset-specific metrics, while a separate endpoint smoke test was correctly limited to interface execution. REFIT work established checksum-verified inputs, a defined H1/H2/H4 primary cohort, actual-delta interval-energy normalisation, quality controls, and deterministic reproduction of derived artefacts.",
        134: "The contribution is therefore architectural traceability, historical-experiment reproducibility, and recorded-household data preparation—not real-household anomaly accuracy or end-to-end runtime reliability. Future validation must proceed through isolated provisioning, data-contract verification, a small pilot replay with event reconciliation, label-alignment review, and only then a controlled H1/H2/H4 replay with deduplication, recovery, and performance tests.",
        136: "CRediT Authorship Contribution Statement",
        138: "Rohit Shrivastava: Conceptualization, methodology, software, REFIT data preprocessing, formal analysis, visualisation, and writing—original draft. Harsh Gupta: Microservice implementation, Docker integration, validation and testing, and writing—review and editing. Anshika Rai: Literature review, data curation, documentation, result-table and visualisation preparation, and writing—review and editing. Syamasudha Veeragandham: Methodology review, validation, writing—review and editing, project coordination, and correspondence.",
        142: "No external funding, grant, or other support was received for this work.",
        146: "The authors declare no competing financial interests or personal relationships that could have influenced the work reported in this paper.",
        150: "The cleaned REFIT data are identified through the dataset records in [6] and [7]. The historical labelled CSV and locally generated preprocessing artefacts are not declared as a public release in the available materials; their availability therefore cannot be guaranteed and remains subject to applicable source terms and author approval.",
        164: "[11] K. Karthika, T. C. G. Thrupthi Chandana, T. Sheikh, Sankeerthana, and Umme, “Smart Energy Monitoring System,” International Journal of Innovative Research in Information Security, vol. 11, no. 9, Oct. 2025, doi: 10.26562/ijiris.2025.v1109.15.",
    }

    for index, text in replacements.items():
        replace_paragraph(p[index], text)

    # Make the objective-to-method-to-evidence alignment visible immediately after Objective 4.
    alignment = (
        "Objective-to-evidence alignment: Objective 1 is addressed by the gateway and device-token admission boundary, with independent security testing pending. "
        "Objective 2 is addressed by the Kafka-integrated service design and monitoring components, with replay, lag, fault, and recovery tests pending. "
        "Objective 3 is supported by deterministic reproduction of the historical baseline, with external and household-level evaluation pending. "
        "Objective 4 is supported by REFIT integrity checking, cohort selection, and normalisation, with ingestion, label alignment, replay, and accuracy evaluation pending."
    )
    # The source includes a blank list-formatted paragraph after the objectives.
    # Normalize it before adding prose so Word does not render a stray marker.
    normalise_blank_paragraph(p[27], p[21])
    insert_paragraph_after(p[26], alignment, template_paragraph=p[21])

    staged_validation = (
        "Staged future validation. The next experiment should first verify the derived_interval_energy_kwh-to-energyConsumed data contract and provision isolated H1/H2/H4 profiles with device-bound credentials. It should then run a small pilot replay, reconcile every accepted and rejected event, validate label alignment, and only then consider a controlled full-household replay. Repeated latency and throughput measurements, input-range controls, deduplication checks, and fault/recovery tests should remain separate acceptance gates."
    )
    insert_paragraph_after(p[129], staged_validation)

    # Replace the broad thematic Table I with a concise study-level comparison.
    old_table = document.tables[0]
    build_literature_table(document, old_table, p[42])

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
