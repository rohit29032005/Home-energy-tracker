from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def xml_text(element) -> str:
    return "".join(element.itertext()).strip()


def table_rows(table):
    rows = []
    for row in table.rows:
        rows.append([" ".join(cell.text.split()) for cell in row.cells])
    return rows


def build_report(path: Path) -> dict:
    document = Document(path)
    section_data = []
    for index, section in enumerate(document.sections, start=1):
        section_data.append(
            {
                "index": index,
                "page_width_inches": round(section.page_width.inches, 3),
                "page_height_inches": round(section.page_height.inches, 3),
                "left_margin_inches": round(section.left_margin.inches, 3),
                "right_margin_inches": round(section.right_margin.inches, 3),
                "top_margin_inches": round(section.top_margin.inches, 3),
                "bottom_margin_inches": round(section.bottom_margin.inches, 3),
                "columns": len(section._sectPr.xpath('./w:cols')),
                "columns_xml": xml_text(section._sectPr),
            }
        )

    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs):
        text = " ".join(paragraph.text.split())
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else "",
                    "alignment": str(paragraph.alignment),
                    "text": text,
                }
            )

    tables = []
    for index, table in enumerate(document.tables, start=1):
        tables.append({"index": index, "rows": len(table.rows), "columns": len(table.columns), "data": table_rows(table)})

    images = []
    for rel_id, relationship in document.part.rels.items():
        if "image" in relationship.reltype:
            images.append({"rel_id": rel_id, "target": str(relationship.target_ref)})

    styles = sorted({p["style"] for p in paragraphs})
    references_start = next((p["index"] for p in paragraphs if p["text"].lower() in {"references", "reference"}), None)
    return {
        "file": str(path),
        "sha256": sha256(path),
        "core_properties": {
            "title": document.core_properties.title,
            "author": document.core_properties.author,
            "subject": document.core_properties.subject,
            "keywords": document.core_properties.keywords,
        },
        "paragraph_count": len(document.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "image_count": len(images),
        "styles": styles,
        "sections": section_data,
        "references_start_index": references_start,
        "paragraphs": paragraphs,
        "tables": tables,
        "images": images,
    }


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: inspect_docx.py INPUT.docx OUTPUT.json")
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    output.write_text(json.dumps(build_report(source), indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
