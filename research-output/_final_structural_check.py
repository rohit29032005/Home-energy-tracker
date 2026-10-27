from docx import Document
from pathlib import Path
src = Document(r"C:\Users\anju2\Desktop\Advance_java_research\version-1.1aHome_Energy_Tracker-publication_updated.docx")
out = Document(r"C:\Users\anju2\Desktop\Home-Energy-Tracker\research-output\Home_Energy_Tracker_Publication_Updated_Final.docx")
assert [p.text for p in src.paragraphs[:4]] == [p.text for p in out.paragraphs[:4]]
assert [p.text for p in src.paragraphs[73:]] == [p.text for p in out.paragraphs[73:]]
assert len(src.sections) == len(out.sections) == 12
for a,b in zip(src.sections,out.sections):
    assert (a.page_width,a.page_height,a.left_margin,a.right_margin,a.top_margin,a.bottom_margin,a.gutter,a._sectPr.xpath('./w:cols/@w:num'),a._sectPr.xpath('./w:cols/@w:space')) == (b.page_width,b.page_height,b.left_margin,b.right_margin,b.top_margin,b.bottom_margin,b.gutter,b._sectPr.xpath('./w:cols/@w:num'),b._sectPr.xpath('./w:cols/@w:space'))
text = '\n'.join(p.text for p in out.paragraphs) + '\n' + '\n'.join(c.text for t in out.tables for r in t.rows for c in r.cells)
needles = [
    'random_forest-csv-20261003T205417Z', 'anomalyProbability = 0.0',
    'Temperature, Humidity, Wind_Speed, and Avg_Past_Consumption',
    'REFIT controlled branch was disabled', 'Historical and Current Local Evaluation Evidence',
    '0.00 to 0.56 kWh', '46 passing tests', '18-test local Python ML suite'
]
for needle in needles:
    assert needle in text, needle
for forbidden in ['REFIT model deployed','real-world accuracy','production-ready','fully resilient']:
    assert forbidden.lower() not in text.lower(), forbidden
print('STRUCTURE_MATCH=TRUE')
print('AUTHOR_BLOCK_MATCH=TRUE')
print('REFERENCES_MATCH=TRUE')
print('SECTION_LAYOUT_MATCH=TRUE')
print('CLAIM_SCAN=PASS')
print('FINAL_TABLE_III_ROWS=', len(out.tables[2].rows))
