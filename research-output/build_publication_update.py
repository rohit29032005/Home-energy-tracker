from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker")
SOURCE = Path(r"C:\Users\anju2\Desktop\Advance_java_research\version-1.1aHome_Energy_Tracker-publication_updated.docx")
OUT_DIR = ROOT / "research-output"
ASSET_DIR = OUT_DIR / "_assets"
FINAL_MANUSCRIPT = OUT_DIR / "Home_Energy_Tracker_Publication_Updated_Final.docx"
AUDIT_DOC = OUT_DIR / "Home_Energy_Tracker_Publication_Update_Audit.docx"
AUDIT_TEXT = OUT_DIR / "Home_Energy_Tracker_Publication_Update_Audit.txt"

NAVY = (40, 82, 120)
PALE = (232, 240, 247)
ACCENT = (56, 121, 174)
DARK = (67, 78, 90)
GRAY = (102, 112, 121)
WHITE = (255, 255, 255)


def _font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf") if bold else Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf") if bold else Path(r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _wrapped(draw, text, font, max_width):
    words = text.split()
    lines, line = [], ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if draw.textlength(candidate, font=font) <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines or [""]


def _draw_centered_wrapped(draw, box, text, font, fill=DARK, spacing=5):
    x0, y0, x1, y1 = box
    lines = _wrapped(draw, text, font, x1 - x0 - 26)
    bboxes = [draw.textbbox((0, 0), line, font=font) for line in lines]
    line_heights = [bb[3] - bb[1] for bb in bboxes]
    total = sum(line_heights) + spacing * max(0, len(lines) - 1)
    y = y0 + max(0, (y1 - y0 - total) / 2)
    for line, bb, height in zip(lines, bboxes, line_heights):
        width = bb[2] - bb[0]
        draw.text((x0 + (x1 - x0 - width) / 2, y), line, font=font, fill=fill)
        y += height + spacing


def _block(draw, x, y, w, h, title, body, *, header_font, body_font, dashed=False, header_color=NAVY):
    if dashed:
        dash, gap = 14, 8
        for xx in range(x, x + w, dash + gap):
            draw.line((xx, y, min(xx + dash, x + w), y), fill=GRAY, width=3)
            draw.line((xx, y + h, min(xx + dash, x + w), y + h), fill=GRAY, width=3)
        for yy in range(y, y + h, dash + gap):
            draw.line((x, yy, x, min(yy + dash, y + h)), fill=GRAY, width=3)
            draw.line((x + w, yy, x + w, min(yy + dash, y + h)), fill=GRAY, width=3)
        draw.rectangle((x + 3, y + 3, x + w - 3, y + h - 3), fill=(248, 249, 250))
        header_fill = (215, 219, 223)
        border = GRAY
    else:
        draw.rectangle((x, y, x + w, y + h), fill=PALE, outline=header_color, width=4)
        header_fill = header_color
        border = header_color
    header_h = 58
    draw.rectangle((x, y, x + w, y + header_h), fill=header_fill, outline=border, width=3)
    title_fill = WHITE if not dashed else DARK
    _draw_centered_wrapped(draw, (x + 8, y + 7, x + w - 8, y + header_h - 4), title, header_font, title_fill, 1)
    _draw_centered_wrapped(draw, (x + 10, y + header_h + 6, x + w - 10, y + h - 8), body, body_font, DARK, 4)


def _arrow(draw, start, end, *, label=None, label_xy=None, dashed=False, width=5):
    color = GRAY if dashed else DARK
    if dashed:
        # Dashed straight segment; disabled/observation paths only.
        x0, y0 = start
        x1, y1 = end
        dx, dy = x1 - x0, y1 - y0
        length = max(1, (dx * dx + dy * dy) ** 0.5)
        unit_x, unit_y = dx / length, dy / length
        pos = 0
        while pos < length - 20:
            segment = min(16, length - pos - 20)
            draw.line((x0 + unit_x * pos, y0 + unit_y * pos, x0 + unit_x * (pos + segment), y0 + unit_y * (pos + segment)), fill=color, width=width)
            pos += 25
    else:
        draw.line((start, end), fill=color, width=width)
    x0, y0 = start
    x1, y1 = end
    dx, dy = x1 - x0, y1 - y0
    length = max(1, (dx * dx + dy * dy) ** 0.5)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (x1, y1)
    base = (x1 - ux * 22, y1 - uy * 22)
    points = [tip, (base[0] + px * 10, base[1] + py * 10), (base[0] - px * 10, base[1] - py * 10)]
    draw.polygon(points, fill=color)
    if label:
        f = _font(21)
        tx, ty = label_xy if label_xy else ((x0 + x1) / 2, (y0 + y1) / 2)
        bb = draw.textbbox((0, 0), label, font=f)
        draw.rectangle((tx - 5, ty - 4, tx + (bb[2] - bb[0]) + 5, ty + (bb[3] - bb[1]) + 4), fill=WHITE)
        draw.text((tx, ty), label, font=f, fill=ACCENT)


def create_architecture_figure(out_path: Path):
    image = Image.new("RGB", (2200, 1160), WHITE)
    draw = ImageDraw.Draw(image)
    heading = _font(31, True)
    title = _font(30, True)
    body = _font(22)
    compact = _font(20)
    small = _font(19)

    draw.text((74, 28), "Input and access layer", font=heading, fill=ACCENT)
    draw.text((620, 28), "Event-processing layer", font=heading, fill=ACCENT)
    draw.text((1120, 28), "Data and analytics layer", font=heading, fill=ACCENT)
    draw.text((1650, 28), "Persistence and observation layer", font=heading, fill=ACCENT)

    _block(draw, 45, 110, 420, 182, "Telemetry and frontend", "Wokwi-compatible ESP32 telemetry: deviceId, energyConsumed, timestamp; or authenticated web request", header_font=title, body_font=body)
    _block(draw, 45, 450, 420, 206, "API Gateway and Keycloak", "JWT role and ownership checks; device-bound X-Device-Token; trusted identity headers replaced", header_font=title, body_font=body)
    _block(draw, 570, 110, 420, 182, "Ingestion Service", "Validates telemetry and publishes the energy-usage event", header_font=title, body_font=body)
    _block(draw, 570, 450, 420, 206, "Kafka", "energy-usage consumed independently by Usage and ML; carries energy-alerts and anomaly-results", header_font=title, body_font=body)
    _block(draw, 1100, 110, 440, 182, "Usage Service and InfluxDB", "Persist time-series usage, compute selected-window aggregates, and publish threshold events", header_font=title, body_font=body)
    _block(draw, 1100, 450, 440, 214, "ML inference service (active baseline)", "KafkaInferenceWorker evaluates baseline RF random_forest-csv-20261003T205417Z; missing contextual fields are median-imputed; publishes anomaly-results", header_font=title, body_font=compact)
    _block(draw, 1100, 790, 440, 164, "Insight Service", "Deterministic bill, ranking, and standby outputs; Overview/Saving Tips use optional Ollama or a deterministic fallback", header_font=title, body_font=compact)
    _block(draw, 1660, 210, 480, 206, "Alert Service and 15-min deduplication", "Consumes threshold and anomaly paths; suppresses duplicate user/device/type keys", header_font=title, body_font=body)
    _block(draw, 1660, 560, 480, 166, "MySQL and Mailpit", "Application metadata and alert history; local SMTP capture for notification evidence", header_font=title, body_font=body)
    _block(draw, 1660, 835, 480, 116, "Prometheus and Grafana", "Service metrics and target-health observation", header_font=title, body_font=body)

    _arrow(draw, (255, 292), (255, 450), label="REST", label_xy=(183, 350))
    _arrow(draw, (465, 553), (570, 200))
    _arrow(draw, (780, 292), (780, 450), label="energy-usage", label_xy=(804, 352))
    _arrow(draw, (990, 543), (1100, 196), label="consume", label_xy=(1010, 348))
    _arrow(draw, (990, 553), (1100, 553), label="consume", label_xy=(1000, 505))
    _arrow(draw, (1540, 190), (1660, 310))
    _arrow(draw, (1540, 557), (1660, 332))
    _arrow(draw, (1900, 416), (1900, 560), label="persist / notify", label_xy=(1930, 480))
    # Prometheus/Grafana observation is described in its block to avoid suggesting that it scrapes MySQL/Mailpit.

    _block(draw, 435, 980, 1120, 130, "Experimental REFIT controlled branch — DISABLED", "Separate endpoint and topics only; it does not consume Wokwi energy-usage, publish legacy anomaly-results, or auto-create alerts", header_font=_font(25, True), body_font=_font(18), dashed=True)
    note = "Solid arrows: active verified service flow. Dashed arrows: observation or disabled experimental state. Wokwi values remain synthetic, not calibrated meter readings."
    draw.text((48, 1120), note, font=small, fill=GRAY)
    image.save(out_path, quality=95)


def create_workflow_figure(out_path: Path):
    image = Image.new("RGB", (2200, 860), WHITE)
    draw = ImageDraw.Draw(image)
    heading = _font(29, True)
    title = _font(24, True)
    body = _font(17)
    mini = _font(16)

    draw.text((45, 22), "Historical baseline artifact preparation", font=heading, fill=ACCENT)
    _block(draw, 45, 64, 330, 128, "Supplied CSV data", "5,000 labelled rows; provenance unavailable", header_font=title, body_font=body)
    _block(draw, 420, 64, 330, 128, "Preprocessing", "Order timestamps; derive Hour, DayOfWeek, and deviation", header_font=title, body_font=body)
    _block(draw, 795, 64, 360, 128, "Model selection", "Chronological 70/15/15 split; RF selected by validation F1", header_font=title, body_font=body)
    _block(draw, 1200, 64, 390, 128, "Saved baseline artifact", "RF v. random_forest-csv-20261003T205417Z; eight features; threshold 0.31", header_font=title, body_font=body)
    _arrow(draw, (375, 128), (420, 128))
    _arrow(draw, (750, 128), (795, 128))
    _arrow(draw, (1155, 128), (1200, 128))

    draw.text((45, 230), "Current runtime telemetry and inference path", font=heading, fill=ACCENT)
    _block(draw, 20, 280, 270, 130, "Input", "Wokwi: deviceId, energyConsumed, timestamp", header_font=title, body_font=body)
    _block(draw, 310, 280, 265, 130, "Access", "JWT ownership or device-bound token", header_font=title, body_font=body)
    _block(draw, 595, 280, 300, 130, "Ingest and Kafka", "Validate; publish energy-usage", header_font=title, body_font=body)
    _block(draw, 915, 280, 280, 130, "Usage and InfluxDB", "Persist/aggregate; publish energy-alerts", header_font=title, body_font=body)
    _block(draw, 1215, 280, 360, 130, "Active baseline ML", "Kafka worker; impute four context fields; RF; publish anomaly-results", header_font=title, body_font=body)
    _block(draw, 1595, 280, 585, 130, "Outputs and observation", "15-min deduplication; alerts/MySQL/Mailpit; deterministic insights; Prometheus/Grafana", header_font=title, body_font=body)
    _arrow(draw, (290, 345), (310, 345))
    _arrow(draw, (575, 345), (595, 345))
    _arrow(draw, (895, 345), (915, 345))
    # The two independent Kafka consumers are stated inside the Usage and ML blocks;
    # no Usage-to-ML arrow is drawn because that would misstate the event topology.
    _arrow(draw, (1575, 395), (1595, 395))

    _block(draw, 155, 490, 1890, 120, "Evidence boundary", "Historical: supplied-CSV held-out evaluation, synthetic E2E, deduplication, target snapshot, k6, and fault checks. Current (2026-10-11): public availability, authorization, 0.56-kWh event-path, baseline health, and one functional prediction. Not meter accuracy, latency, generalisation, or production readiness.", header_font=_font(23, True), body_font=_font(16), header_color=NAVY)
    _block(draw, 435, 655, 1340, 103, "Experimental REFIT controlled branch — DISABLED", "Separate offline/experimental input and output topics; no Wokwi energy-usage or legacy alert connection; no live prediction or empirical result reported", header_font=_font(22, True), body_font=_font(16), dashed=True)
    draw.text((45, 810), "Solid arrows denote active baseline workflow. The disabled REFIT box is intentionally not connected to the active runtime flow.", font=mini, fill=GRAY)
    image.save(out_path, quality=95)


def _first_run_properties(paragraph):
    if not paragraph.runs:
        return None
    rpr = paragraph.runs[0]._r.rPr
    return deepcopy(rpr) if rpr is not None else None


def _set_paragraph_text(paragraph, text, *, template_paragraph=None):
    template = template_paragraph or paragraph
    ppr = deepcopy(template._p.pPr) if template._p.pPr is not None else None
    rpr = _first_run_properties(template)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    if ppr is not None:
        existing = paragraph._p.pPr
        if existing is not None:
            paragraph._p.remove(existing)
        paragraph._p.insert(0, ppr)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return run


def _set_cell_text(cell, text, *, template_cell=None):
    target = cell.paragraphs[0]
    template = template_cell.paragraphs[0] if template_cell is not None else target
    _set_paragraph_text(target, text, template_paragraph=template)
    for p in cell.paragraphs[1:]:
        _set_paragraph_text(p, "", template_paragraph=template)


def _replace_figure(paragraph, image_path, width, height, alt_text, title_text):
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    for drawing in list(run._r.findall(qn("w:drawing"))):
        run._r.remove(drawing)
    shape = run.add_picture(str(image_path), width=width, height=height)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", title_text)


def _set_table_row(table, row_index, values):
    row = table.rows[row_index]
    for cell, value in zip(row.cells, values):
        _set_cell_text(cell, value)


def _add_table_row(table, values, *, template_row_index=-1):
    template = table.rows[template_row_index]
    row = table.add_row()
    for cell, source_cell, value in zip(row.cells, template.cells, values):
        _set_cell_text(cell, value, template_cell=source_cell)
    return row


def update_manuscript():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    fig1 = ASSET_DIR / "figure_1_updated.png"
    fig2 = ASSET_DIR / "figure_2_updated.png"
    create_architecture_figure(fig1)
    create_workflow_figure(fig2)

    doc = Document(str(SOURCE))
    assert len(doc.paragraphs) == 98 and len(doc.tables) == 3 and len(doc.inline_shapes) == 2 and len(doc.sections) == 12
    original_identity = [doc.paragraphs[i].text for i in range(4)]

    replacements = {
        6: "Abstract - Residential energy platforms require a traceable route from meter-like telemetry to controlled feedback. This study implements a local 17-service Docker composition combining Wokwi-compatible ESP32 telemetry, API-gateway admission control, Apache Kafka, InfluxDB, MySQL, deterministic insights, and duplicate-aware alert delivery. User-facing routes apply Keycloak JWT, role, and ownership checks, while simulated-device ingestion uses a device-bound token. A baseline Random Forest artifact, random_forest-csv-20261003T205417Z, is selected by validation F1 from the supplied labelled CSV; its historical held-out precision, recall, F1, ROC-AUC, and PR-AUC are 0.880952, 0.925000, 0.902439, 0.997430, and 0.948532, respectively. On 2026-10-11, the live ML container reported the baseline model UP and its legacy Kafka worker connected; a Wokwi-compatible three-field input with energyConsumed = 0.55 returned anomaly = false, anomalyProbability = 0.0, decisionThreshold = 0.31, and four explicit contextual-feature imputations. The isolated REFIT controlled branch was disabled and contributed no live result. Historical controlled local checks recorded duplicate suppression within 15 min, an eight-target Prometheus snapshot, and 450/450 accepted k6 ingestion requests. The results demonstrate implementation-level integration and bounded local evidence, not physical-meter accuracy, field-model generalisation, production security, or production-scale resilience.",
        10: "This work addresses the integration of a local, event-driven energy platform rather than the accuracy of a physical meter. The deployed composition separates user and device metadata, ingestion, usage, insight, alert, API gateway, and ML inference functions, and includes Kafka, InfluxDB, MySQL, Mailpit, Keycloak, Prometheus, Grafana, and a frontend. Wokwi supplies timestamped synthetic JSON payloads to the same HTTP ingestion contract intended for a future meter; Kafka decouples request acceptance from downstream persistence and analysis. The active ML component is a baseline FastAPI service and Kafka worker; the separately implemented REFIT controlled adapter is disabled in the current Compose configuration.",
        11: "The study makes four bounded contributions: it implements a 17-service local Docker path from simulated telemetry to time-series storage and access-controlled APIs; it trains and deploys a fixed baseline anomaly-inference artifact while retaining deterministic bill, ranking, standby, and threshold functions; it applies device-bound telemetry admission, downstream caller checks, and duplicate-aware notification; and it records controlled local evidence for end-to-end behavior, gateway security, alert deduplication, observability, ingress load, selected outage conditions, and current baseline runtime status. The scope excludes calibrated sensing, utility-bill validation, physical-field generalisation, production penetration testing, and production-scale resilience claims.",
        17: "The research gap is limited implementation-level linkage between residential telemetry, event-driven services, controlled access, historical storage, transparent analytics, anomaly inference, and alert delivery with explicit verification boundaries. The proposed system addresses that gap through Wokwi-compatible telemetry, device-bound gateway admission, Spring Boot microservices, Kafka, InfluxDB, MySQL, a deployed baseline Random Forest artifact, deterministic insight functions, and Mailpit-backed alerting with a configured deduplication window. Remaining gaps include calibrated sensing, documented dataset provenance, per-household contextual labels, field validation, production-grade security assessment, and distributed resilience evaluation.",
        20: "The runtime input contains deviceId, energyConsumed, and timestamp. Wokwi firmware emits this JSON payload at an approximately 8 s interval; Eco, Heavy, and Alert Test modes generate controlled values for integration demonstration and are not calibrated physical measurements. User-facing requests enter through the API Gateway with Keycloak JWT, realm-role, and ownership checks. The gateway replaces trusted identity headers, and Wokwi-compatible ingestion requires a device-bound X-Device-Token; downstream calls use separate static gateway or background-service credentials. The Ingestion Service validates accepted payloads and publishes energy-usage events to Kafka.",
        21: "The Usage Service and the FastAPI ML service independently consume energy-usage events. Usage writes time-stamped energy_usage points to InfluxDB and derives selected-window aggregates. The ML service hosts KafkaInferenceWorker, which invokes the fixed baseline Random Forest artifact random_forest-csv-20261003T205417Z and publishes anomaly-results; repeated processing failure is routed to anomaly-processing-failures before offset commitment. For the three-field Wokwi payload, Temperature, Humidity, Wind_Speed, and Avg_Past_Consumption are median-imputed and disclosed in the prediction response. User and Device Services maintain household, appliance, and alert-setting metadata in MySQL. Threshold and anomaly paths reach the Alert Service, which persists notifications and uses a configured 15-min deduplication window. Prometheus and Grafana provide service-level target and metric observation. The isolated REFIT controlled route is disabled and is not connected to Wokwi energy-usage or legacy alert topics.",
        24: "Fig. 1. Verified service architecture, active baseline inference path, and isolated disabled REFIT branch of the Home Energy Tracker.",
        27: "The Insight Service derives deterministic outputs from aggregated usage. It computes a daily rate and a 30-day energy projection from total energy E over D selected days. The configured local tariff uses a fixed 10.00 charge and slab rates of 0.12 per kWh for 0-100 kWh, 0.18 per kWh for 100-300 kWh, and 0.25 per kWh above 300 kWh. Device ranking uses each device total divided by the household total, and the standby estimate remains 20% of average load. Overview and Saving Tips may request optional Ollama text enrichment, but the verified 2026-10-11 response used deterministic fallback (isAiGenerated = false). These are demonstration rules, not a utility tariff, an appliance-level standby measurement, or a verified live generative-AI result.",
        29: "The anomaly branch uses the supplied smart_meter_data.csv with 5,000 records and seven columns, including Electricity_Consumed, weather/context fields, Avg_Past_Consumption, and Anomaly_Label (4,750 Normal and 250 Abnormal). Records are stably ordered by timestamp; Hour, DayOfWeek, and ConsumptionDeviation = Electricity_Consumed - Avg_Past_Consumption are derived. A chronological 3,500/750/750 train/validation/test split compares a balanced logistic-regression pipeline with a 300-tree balanced-subsample Random Forest. The saved baseline exposes eight features, uses threshold 0.31 selected by validation F1, and is loaded at FastAPI startup. For live Wokwi input, the four unavailable contextual fields are imputed with training-partition medians, while time-derived features are constructed from the timestamp. The 2026-10-11 direct prediction verifies loading, feature preparation, and response generation only; it is not a new accuracy, latency, generalisation, or complete telemetry-to-alert experiment. Dataset source, licence, and label-generation provenance are not documented.",
        31: "For threshold alerting, a scheduler evaluates the preceding one-hour household sum every 10 s. When alerting is enabled and the sum exceeds the user-defined threshold, an energy-alerts event is published. The Alert Service consumes threshold and ML anomaly paths, persists alerts in MySQL, and attempts SMTP delivery with Mailpit as the local demonstration inbox. Its deduplication key is userId:deviceId-or-HOUSEHOLD:alertType; matching keys are suppressed during the configured 15-min window. The active baseline ML worker includes a bounded default of three attempts and an anomaly-processing-failures topic as implemented handling, not as a separately evaluated recovery guarantee. The optional REFIT controlled endpoint and topic route remain isolated and disabled by default; they contributed no alert or empirical result in this paper.",
        38: "Fig. 2 separates historical baseline artifact preparation from the current runtime telemetry path. The upper lane prepares the supplied CSV, selects the baseline Random Forest and threshold on validation data, and produces the fixed inference artifact. The active runtime lane keeps authorised telemetry admission, ingestion, Kafka publication, usage persistence, baseline feature construction and imputation, anomaly-result publication, deterministic insight generation, alert deduplication, and output delivery distinct. The dashed REFIT box is a separate disabled experimental route, not a deployed Wokwi or legacy-alert path. Historical E2E, security, deduplication, target-health, ingress-load, and selected outage checks are distinguished from the current 2026-10-11 availability, authorisation, event-path, and functional baseline-inference checks.",
        41: "Fig. 2. Baseline artifact preparation, active telemetry workflow, and bounded historical/current verification evidence.",
        45: "The system was exercised as a local, single-host Docker deployment on Windows 11 Home Single Language (10.0.26300) with eight logical processors and 7.74 GiB RAM. The compose configuration defines 17 services covering relational, event, time-series, e-mail, identity, application, frontend, and observability functions. On 2026-10-11, Gateway health, frontend, Keycloak OpenID Connect discovery, Prometheus readiness, and Grafana login each returned HTTP 200; these checks establish host availability, not production-scale operation. The environment supports controlled implementation and verification; it is not a replicated cloud or production deployment.",
        52: "1. A prior controlled local E2E experiment submitted 0.25, 0.50, and 0.75 kWh analytics inputs for user 15/device 314 and a separate 1.00 kWh threshold input for user 16/device 315, whose configured threshold was 0.10 kWh. These historical synthetic REST observations were not silently re-presented as a post-recovery experiment.",
        53: "2. The historical offline procedure stably sorted the supplied CSV, produced chronological train/validation/test partitions of 3,500/750/750 records, selected the threshold by validation F1, and evaluated the chosen and logistic-regression baseline models once on the held-out partition. The current baseline artifact metadata retains the same recorded model identity and threshold; its metrics are reported as historical artifact evaluation, not as a newly repeated live study.",
        54: "3. On 2026-10-11, current runtime checks covered public availability, owned and cross-owner user access, device-bound telemetry admission, the Ingestion-to-Kafka-to-Usage-to-Influx path, protected insight endpoints, an administrator-only test alert, Mailpit capture, and temporary application-profile/device CRUD cleanup. A separate baseline ML health and direct prediction check established model loading, connected worker status, explicit imputation, and a response for a Wokwi-compatible payload; it did not rerun the complete ML-to-alert path after restart.",
        55: "4. Historical local ingress was measured with k6 at 2, 5, and 10 requests/s for 30 s stages plus ramp-down; separate Mailpit and ML-service outage checks recorded ingress availability and manual container recovery. Current source-level regression evidence consists of four separately executed Java commands with 46 passing tests, an 18-test local Python ML suite, and successful frontend lint and production build. These execution records are reported separately and are not merged into one production-readiness score.",
        58: "Table III separates historical offline and controlled results from current runtime verification because the measurements were collected by different procedures. Historical supplied-CSV assessment quantifies data-specific model performance, and historical E2E rows document synthetic integration conditions. Current 2026-10-11 evidence confirms availability, authorization/event-path behavior, and a live baseline endpoint/worker response, not a fresh statistical model evaluation or an ML-to-alert completion after restart. Source-level Java, Python, lint, and build records are likewise separate regression evidence rather than a unified production-readiness score.",
        65: "On the supplied held-out CSV, the Random Forest achieved substantially stronger precision, F1, ROC-AUC, and PR-AUC than the logistic-regression baseline. This supports selection of the saved Random Forest artifact for this dataset and split, but it does not validate field anomaly detection. The current direct Wokwi-compatible prediction confirms that the loaded baseline can construct features, explicitly impute missing context, and return a decision; it does not add an accuracy observation or prove the post-restart telemetry-to-ML-to-alert route. The historical E2E analytics and alert observations demonstrate a traceable local integration path from synthetic REST input to stored usage, deterministic outputs, and a captured notification; they do not establish physical-meter accuracy or utility-bill agreement.",
        66: "The historical gateway-status, deduplication, Prometheus target, k6, and selected outage records remain bounded to their stated dates and procedures. The current post-hardening run adds five user-route checks, five telemetry-admission checks, an observed 0.00-to-0.56 kWh Usage change after 0.37 and 0.19 kWh readings, protected insight responses, one administrator-only alert with Mailpit capture, and temporary CRUD cleanup. Static gateway and internal-service credentials provide defence in depth, but they are not mTLS, workload identity, secret rotation, replay protection, or penetration-test evidence. No fresh all-target Prometheus snapshot was captured after the ML service restart.",
        67: "The data and runtime boundaries remain material. The CSV has no documented source, licence, label provenance, household/device identifier, or provenance for Avg_Past_Consumption. Its Electricity_Consumed values are in an undocumented 0-1 range, whereas Wokwi modes send direct values in approximately the 1-40 range; optional Wokwi context is median-imputed and no range-validation or calibration experiment is reported. The REFIT controlled branch is disabled because its optional artifacts are unavailable to the current Docker build and contributes no result. No runtime inference-latency distribution was captured; the current single endpoint invocation is a functional check rather than a latency benchmark. Frontend rendering alone was not used as backend-data evidence because local fallback data may be displayed if an API request fails.",
        69: "The principal limitations are synthetic, non-calibrated Wokwi telemetry; the anonymous dataset and missing provenance; the training-to-Wokwi unit/range mismatch; and the absence of physical-meter, utility-bill, and external-field validation. The optional REFIT controlled branch is disabled and lacks a reported live or real-anomaly evaluation. The evaluation uses one local Docker host with a bounded ingress test, no distributed tracing or Kafka-lag study, and only selected manually recovered service-outage checks. Keycloak, gateway, device-bound token, and downstream-caller controls were functionally checked, but the local development deployment uses static shared credentials rather than TLS, mTLS, workload identity, credential rotation, replay protection, or penetration testing. No browser role/refresh acceptance test and no fresh post-restart all-target Prometheus snapshot were captured.",
        70: "Future work should collect consented, documented, calibrated interval-meter data with per-device and household context; establish label and feature provenance; and add range or out-of-distribution checks before model deployment. It should make the experimental REFIT artifacts accessible only after an audited dataset, label process, and isolated evaluation protocol are available, then compare it without transferring baseline metrics. Further work should repeat end-to-end and load measurements over longer workloads, add tracing and Kafka-lag instrumentation, and perform automated fault, retry, failure-topic, replay, and recovery experiments. Production hardening should address TLS, rotating secrets or workload identity, device provisioning, browser acceptance tests, and an independent security assessment before any real-world deployment claim.",
        72: "This paper presents a microservices-based Home Energy Tracker updated from an integration baseline to a locally evaluated event-driven system with gateway access controls, Kafka-mediated processing, InfluxDB and MySQL persistence, deterministic insights, a deployed baseline Random Forest inference artifact, alert deduplication, and Prometheus/Grafana observation. Historical evidence includes the supplied-CSV model assessment, synthetic E2E analytics and alert checks, gateway-security cases, duplicate suppression, an eight-target snapshot, local k6 ingress measurement, and two limited outage-availability observations. Current 2026-10-11 verification additionally shows public service availability, access-controlled event propagation, a live baseline ML container and connected worker, and one Wokwi-compatible prediction with explicit imputation; the REFIT controlled branch remained disabled. The work demonstrates implementation-level integration and bounded local verification; it does not establish meter accuracy, externally generalisable anomaly detection, production-scale performance, comprehensive resilience, or production security.",
    }
    for index, text in replacements.items():
        _set_paragraph_text(doc.paragraphs[index], text)

    _replace_figure(
        doc.paragraphs[23], fig1, doc.inline_shapes[0].width, doc.inline_shapes[0].height,
        "Verified service architecture. Wokwi-compatible telemetry and frontend requests enter an API gateway and Keycloak boundary, then pass through ingestion and Kafka to Usage/InfluxDB and active baseline ML inference. Alert Service applies 15-minute deduplication and persists/captures alerts. Prometheus/Grafana observes services. An isolated REFIT controlled experimental branch is visibly disabled and has no connection to Wokwi or legacy alerts.",
        "Verified service architecture with active baseline inference and disabled REFIT branch",
    )
    _replace_figure(
        doc.paragraphs[40], fig2, doc.inline_shapes[1].width, doc.inline_shapes[1].height,
        "Workflow figure separating historical baseline artifact preparation from the active runtime telemetry path. The active path processes a three-field Wokwi-compatible payload through controlled access, Kafka, usage persistence, baseline Random Forest inference with four contextual feature imputations, and alert/output services. Historical and current verification evidence are separated, and a disabled isolated REFIT branch is shown without a connection to Wokwi or legacy alert topics.",
        "Baseline workflow with current runtime evidence and disabled REFIT branch",
    )

    t0, t1, t2 = doc.tables
    _set_table_row(t0, 1, ["Access and telemetry", "Keycloak JWT, roles, and ownership protect user routes; Wokwi-compatible ingestion uses a device-bound X-Device-Token; gateway and internal callers use separate trusted credentials.", "Authorised user request or machine admission."])
    _set_table_row(t0, 2, ["Ingestion and event transfer", "Ingestion validates payloads and publishes Kafka energy-usage events. Usage and the ML worker consume the event independently.", "Asynchronous event for persistence, aggregation, and baseline inference."])
    _set_table_row(t0, 4, ["Baseline model and insight", "FastAPI hosts KafkaInferenceWorker and the fixed baseline Random Forest random_forest-csv-20261003T205417Z. Four absent Wokwi context fields are median-imputed; Insight functions are deterministic with optional Ollama fallback.", "Anomaly probability/decision plus deterministic bill, ranking, standby, overview, and saving outputs."])
    _set_table_row(t0, 5, ["Alerting, observation, and experimental boundary", "Threshold and baseline anomaly paths reach Alert Service, which applies 15-min deduplication and uses Mailpit; Prometheus/Grafana observe services. REFIT controlled is separate and disabled.", "Alert/email evidence, service metrics, and an explicitly inactive experimental branch."])

    _set_table_row(t1, 1, ["Runtime and deployment", "Windows 11 Home Single Language 10.0.26300; local single-host Docker with eight logical processors and 7.74 GiB RAM. The compose file defines 17 services. On 2026-10-11, Gateway, frontend, Keycloak OIDC, Prometheus readiness, and Grafana login each returned HTTP 200."])
    _set_table_row(t1, 2, ["Application and data stack", "Java 17 Spring Boot services and API Gateway; React/Vite frontend; FastAPI baseline ML service; Apache Kafka, MySQL, InfluxDB, Mailpit, Keycloak, Prometheus, and Grafana. The composition is locally deployed, not a cloud-scale result."])
    _set_table_row(t1, 4, ["CSV data and baseline model", "smart_meter_data.csv: 5,000 records, seven columns, and Normal/Abnormal labels of 4,750/250. A chronological 3,500/750/750 split selected the Random Forest. Artifact random_forest-csv-20261003T205417Z exposes eight features and threshold 0.31. Source, licence, and label provenance are undocumented."])
    _set_table_row(t1, 5, ["Historical controlled local E2E", "0.25, 0.50, and 0.75 kWh analytics inputs for user 15/device 314; a 1.00 kWh threshold input for user 16/device 315. These are prior synthetic REST observations, not a re-run after the current ML recovery."])
    _set_table_row(t1, 6, ["Evaluation protocols", "Historical held-out/model, E2E, deduplication, target-health, k6, and selected outage procedures; current 2026-10-11 availability, authorization/event-path, and direct baseline inference checks; separate 46-test Java, 18-test local Python, frontend lint, and build regressions."])

    _set_table_row(t2, 1, ["Historical Random Forest held-out test", "750 records; selected threshold = 0.31; artifact-evaluation condition.", "TN 705, FP 5, FN 3, TP 37; precision 0.880952; recall 0.925000; F1 0.902439; ROC-AUC 0.997430; PR-AUC 0.948532."])
    _set_table_row(t2, 2, ["Historical Logistic-Regression baseline", "Same held-out partition; threshold = 0.28.", "Precision 0.053333; recall 1.000000; F1 0.101266; ROC-AUC 0.422852; PR-AUC 0.059710."])
    _set_table_row(t2, 3, ["Historical controlled E2E analytics", "Three synthetic inputs total 1.50 kWh.", "3/3 HTTP 201; cost $0.18; daily 0.50 kWh; projected 15.00 kWh/$11.80; rank 100%; standby 4.17 W/LOW; observed within polling-bounded 3.79 s."])
    _set_table_row(t2, 4, ["Historical threshold alert and Mailpit", "1.00 kWh input exceeds 0.10 kWh threshold.", "HTTP 201; persisted sent THRESHOLD alert and matching Mailpit email; observed within polling-bounded 13.64 s."])
    _set_table_row(t2, 5, ["Historical gateway security", "Eight configured access-control cases.", "8/8 expected statuses: owned 200; missing/invalid/expired JWT 401; cross-user 403; administrator 200; device-token ingestion 201."])
    _set_table_row(t2, 6, ["Historical alert deduplication", "Two identical 1.00-kWh events at the same timestamp.", "One ML_ANOMALY alert was stored and a duplicate was explicitly suppressed within 15 min."])
    _set_table_row(t2, 7, ["Historical observability snapshot; current host availability", "Historical Prometheus target snapshot; current readiness/login routes.", "Historical 8/8 active targets up. Current Prometheus readiness and Grafana login returned HTTP 200; no new all-target snapshot is claimed after ML restart."])
    _set_table_row(t2, 8, ["Historical local ingress and fault checks", "k6 2/5/10 requests/s; Mailpit and ML separately stopped.", "450/450 accepted; mean 39.442 ms; p95 109.773 ms; 0% HTTP failure. During each outage, ingestion returned 201 and manual restart recovered the container."])
    _add_table_row(t2, ["Current active baseline ML runtime (2026-10-11)", "Live /health: baseline UP, legacy Kafka worker connected, REFIT controlled DISABLED. Direct Wokwi-compatible input: deviceId 1, energyConsumed 0.55, timestamp.", "anomaly = false; anomalyProbability = 0.0; decisionThreshold = 0.31; four imputed features: Temperature, Humidity, Wind_Speed, Avg_Past_Consumption. Functional check only."])
    _add_table_row(t2, ["Current authorization and event-path validation (2026-10-11)", "Five user-route checks; five telemetry-admission checks; accepted readings 0.37 and 0.19 kWh; protected insights and administrator alert checked.", "User checks: 401/200/403/403/200; telemetry: 401/401/201/201/403; Usage changed 0.00 to 0.56 kWh (>=0.559 lower bound). Alert ID 492 reported sent = true and Mailpit API returned 200."])

    _set_paragraph_text(doc.paragraphs[33], "TABLE I")
    _set_paragraph_text(doc.paragraphs[34], "Verified System Components, Current Baseline Flow, and Experimental Boundary")
    _set_paragraph_text(doc.paragraphs[47], "TABLE II")
    _set_paragraph_text(doc.paragraphs[48], "Local Experimental Environment and Evidence Conditions")
    _set_paragraph_text(doc.paragraphs[60], "TABLE III")
    _set_paragraph_text(doc.paragraphs[61], "Historical and Current Local Evaluation Evidence")

    # The master document's title, author information, structure, blank section-break paragraphs,
    # and references are intentionally unchanged.
    assert [doc.paragraphs[i].text for i in range(4)] == original_identity
    assert len(doc.paragraphs) == 98 and len(doc.tables) == 3 and len(doc.inline_shapes) == 2 and len(doc.sections) == 12
    doc.save(str(FINAL_MANUSCRIPT))


def _shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        from docx.oxml import OxmlElement
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _add_audit_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        _shade_cell(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()
    return table


def _audit_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    run.font.name = "Arial"
    return p


def _audit_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    return p


def create_audit_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.60)
    section.right_margin = Inches(0.60)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Home Energy Tracker Publication Update Audit")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(15)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Evidence-based revision package for the IEEE-style manuscript")
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(9)

    _audit_heading(doc, "1. Scope and source control")
    doc.add_paragraph(
        "The attached source manuscript version-1.1aHome_Energy_Tracker-publication_updated.docx was used because the filename with “(2)” named in the instruction was not present. The original was not overwritten. This audit accompanies the updated copy and documents the evidence status of substantive changes."
    )

    _audit_heading(doc, "2. Content-audit summary")
    _add_audit_table(doc, ["Area", "Status", "Action", "Reason/evidence"], [
        ("Title, author block, references", "Retained", "No content redesign", "Unchanged to preserve manuscript identity; 24 reference entries remain cited."),
        ("Abstract and contributions", "Revised", "Current baseline/REFIT distinction", "Live baseline health/prediction added; REFIT declared disabled; historical metrics labelled historical."),
        ("Architecture and methodology", "Revised", "Correct worker/data-flow wording", "FastAPI-hosted KafkaInferenceWorker, device-bound token, imputation, downstream credentials, and optional fallback clarified."),
        ("Figures 1–2", "Replaced in place", "Update active and disabled branches", "Figures retain numbering, placement, and navy/light-blue box language while adding baseline, imputation, Kafka topics, and disabled REFIT."),
        ("Tables I–III", "Revised", "Historical/current separation", "The old numeric results are retained under their conditions; current live runtime rows added without treating them as accuracy measurements."),
        ("Results, limitations, conclusion", "Revised", "Bound claims", "New inference is functional-only; security, observability, and REFIT limitations stated explicitly."),
    ])

    _audit_heading(doc, "3. Section-by-section change summary")
    _add_audit_table(doc, ["Section", "Substantive update"], [
        ("Abstract", "Adds the active artifact identity, current live functional prediction, four expected imputations, and disabled REFIT state; historical metrics retained as historical."),
        ("I–II", "Retains integration research gap; adds baseline/REFIT distinction and device-bound admission boundary without claiming a new ML algorithm."),
        ("III-A", "Corrects the ML description: KafkaInferenceWorker resides in the FastAPI ML service, consumes energy-usage, produces anomaly-results, and routes repeated failures to the failure topic."),
        ("III-B", "Documents baseline feature construction, training-median imputation, optional Ollama deterministic fallback, and REFIT isolation."),
        ("III-C", "Separates historical artifact/E2E evidence from current runtime verification and disabled experimental REFIT."),
        ("IV", "Labels older synthetic E2E/k6/outage records as historical and adds 2026-10-11 availability, authorization/event-path, ML functional, and separate regression evidence."),
        ("V", "Adds current runtime rows to Table III and makes the discussion distinguish functional inference from accuracy/latency/full post-restart event-to-alert proof."),
        ("VI–VII", "Adds disabled REFIT, static-shared-secret, no fresh all-target snapshot, and no browser acceptance-test limitations; preserves non-production boundary."),
    ])

    _audit_heading(doc, "4. Figure and table update summary")
    _add_audit_table(doc, ["Identifier", "Update"], [
        ("Fig. 1", "Retained in-place architecture figure; now shows the active baseline Random Forest and Kafka worker, device-bound token, anomaly-results/energy-alerts flow, optional Ollama fallback, Prometheus/Grafana, and a dashed disabled REFIT branch."),
        ("Fig. 2", "Retained in-place workflow figure; separates historical artifact preparation, current runtime feature imputation/inference, evidence boundary, and disabled REFIT route."),
        ("Table I", "Revised component/data-flow descriptions; identifies the baseline version and current REFIT-disabled boundary."),
        ("Table II", "Retains host/stack facts while adding the active artifact/threshold and historical versus current evidence conditions."),
        ("Table III", "Relabels old rows as historical and adds current live baseline plus authorization/event-path rows. It does not convert a single prediction into a performance metric."),
    ])

    _audit_heading(doc, "5. Claim, evidence, and status matrix")
    _add_audit_table(doc, ["Claim", "Evidence", "Status"], [
        ("17-service local Compose architecture", "docker-compose.yml and 2026-10-11 validation notes", "Verified configuration/local deployment"),
        ("Gateway/frontend/Keycloak/Prometheus/Grafana availability", "2026-10-11 HTTP 200 checks", "Current verified availability only"),
        ("Active baseline model", "model_metadata.json; live /health", "Current verified: random_forest-csv-20261003T205417Z"),
        ("Live baseline functional inference", "2026-10-11 /predict result", "Current verified functional check only"),
        ("Four missing runtime contextual features", "main.py feature preparation and response", "Current verified imputation behavior"),
        ("REFIT controlled branch", "Compose flags and live health", "Disabled; no live/empirical result used"),
        ("Historical RF metrics", "model_metadata.json and historical evaluation files", "Retained as historical supplied-CSV result"),
        ("Historical k6 and Prometheus values", "dated historical evidence", "Retained with dates/conditions; not repeated"),
        ("Authorization/event path", "2026-10-11 live matrix and 0.37+0.19 kWh run", "Current functional verification"),
        ("Source-level regressions", "46 Java tests, 18 local Python tests, frontend lint/build", "Separate test records; not a composite reliability score"),
    ])

    _audit_heading(doc, "6. Retained historical results")
    for text in [
        "Baseline Random Forest held-out CSV metrics and Logistic Regression comparison, explicitly identified as historical artifact evaluation on the repository-supplied labelled CSV.",
        "Synthetic controlled E2E analytics, threshold alert/Mailpit, gateway-status, deduplication, Prometheus target, k6, and selected outage observations, each retained with the original bounded conditions.",
        "No historical measurement has been described as a fresh 2026-10-11 rerun.",
    ]:
        _audit_bullet(doc, text)

    _audit_heading(doc, "7. Revised or removed obsolete/misleading claims")
    for text in [
        "Revised generic ML wording to the actual FastAPI-hosted KafkaInferenceWorker and baseline artifact flow.",
        "Removed any implication that REFIT is deployed, provides live predictions, consumes Wokwi traffic, or contributes final empirical metrics.",
        "Revised current-prediction wording so anomalyProbability = 0.0 is not presented as an accuracy, confidence, latency, or generalisation result.",
        "Revised observability wording: current Prometheus/Grafana HTTP availability is verified; the 8/8 target snapshot remains historical.",
        "Revised optional AI wording: current Overview/Saving Tips result was deterministic fallback, not verified live generative output.",
        "Revised security wording: current controls use local Keycloak/gateway/device-bound tokens/static internal secrets and are not production mTLS, workload identity, rotation, replay protection, or penetration-test evidence.",
    ]:
        _audit_bullet(doc, text)

    _audit_heading(doc, "8. Unresolved or manual-verification items")
    for text in [
        "The provided manuscript filename differs from the filename named in the instruction; this update used the attached accessible source document.",
        "No physical-meter calibration, utility-bill agreement, consented household field data, or external generalisation experiment is available.",
        "CSV provenance, licence, label-generation procedure, household/device identity, and unit semantics remain undocumented.",
        "The optional REFIT artifacts are unavailable to the current Docker build and its controlled benchmark must not be represented as deployed or real-anomaly performance.",
        "No fresh telemetry-to-Kafka-to-ML-to-alert completion, post-ML-restart all-target Prometheus sweep, browser role/refresh acceptance test, latency distribution, or production security assessment was captured.",
    ]:
        _audit_bullet(doc, text)

    _audit_heading(doc, "9. Final formatting and consistency audit")
    for text in [
        "Preserved: title, authors, affiliations, corresponding-author designation, two-column section structure, section-break paragraphs, figure/table numbering and captions, reference numbering, and 24-entry reference list.",
        "Verified in the edited DOCX structure: 98 body paragraphs, 3 tables, 2 inline figures, and 12 document sections, matching the source structural counts.",
        "Terminology scan completed for misleading REFIT deployment, real-world accuracy, production-ready, and fully resilient claims; replacement text uses bounded, evidence-specific language.",
        "Page-level rendering was completed with Microsoft Word PDF export and Poppler rasterization; all seven manuscript pages and all three audit pages were visually inspected before delivery.",
    ]:
        _audit_bullet(doc, text)

    doc.save(str(AUDIT_DOC))
    AUDIT_TEXT.write_text(
        "Companion audit document generated: " + str(AUDIT_DOC) + "\n"
        "Source master was preserved. The updated manuscript contains current verified baseline ML functionality, historical-result boundaries, and disabled REFIT status.\n",
        encoding="utf-8",
    )


def final_checks():
    doc = Document(str(FINAL_MANUSCRIPT))
    assert len(doc.paragraphs) == 98, len(doc.paragraphs)
    assert len(doc.tables) == 3, len(doc.tables)
    assert len(doc.inline_shapes) == 2, len(doc.inline_shapes)
    assert len(doc.sections) == 12, len(doc.sections)
    texts = "\n".join(p.text for p in doc.paragraphs)
    texts += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    for phrase in ["REFIT model deployed", "real-world accuracy", "production-ready", "fully resilient"]:
        assert phrase.lower() not in texts.lower(), phrase
    assert "random_forest-csv-20261003T205417Z" in texts
    assert "anomalyProbability = 0.0" in texts
    assert "REFIT controlled branch was disabled" in texts
    assert "8/8 active targets up" in texts
    # Confirm updated alternate text includes baseline and disabled REFIT state.
    descriptions = []
    for shape in doc.inline_shapes:
        descriptions.append(shape._inline.docPr.get("descr", ""))
    assert all("REFIT" in description and "baseline" in description.lower() for description in descriptions)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "sections": len(doc.sections),
        "table_iii_rows": len(doc.tables[2].rows),
        "figures_updated": len(descriptions),
    }


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    update_manuscript()
    create_audit_document()
    print(final_checks())
