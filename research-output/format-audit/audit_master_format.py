from __future__ import annotations

import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2])
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def qn(local: str) -> str:
    return f"{{{W}}}{local}"


def twips(value):
    return None if value is None else round(value / 20.0, 2)


def align_name(value):
    return None if value is None else str(value)


doc = Document(str(SRC))
paragraphs = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.replace("\n", " ")
    if not text.strip():
        continue
    pf = p.paragraph_format
    rinfo = []
    for run in p.runs:
        if not run.text:
            continue
        font = run.font
        rinfo.append({
            "text": run.text[:120],
            "bold": bool(font.bold),
            "italic": bool(font.italic),
            "font": font.name,
            "size_pt": None if font.size is None else round(font.size.pt, 2),
            "color": None if font.color.rgb is None else str(font.color.rgb),
            "caps": bool(font.all_caps),
        })
    paragraphs.append({
        "index": i,
        "text": text[:400],
        "style": p.style.name if p.style else None,
        "alignment": align_name(p.alignment),
        "left_indent_in": None if pf.left_indent is None else round(pf.left_indent.inches, 3),
        "right_indent_in": None if pf.right_indent is None else round(pf.right_indent.inches, 3),
        "first_line_in": None if pf.first_line_indent is None else round(pf.first_line_indent.inches, 3),
        "space_before_pt": None if pf.space_before is None else round(pf.space_before.pt, 2),
        "space_after_pt": None if pf.space_after is None else round(pf.space_after.pt, 2),
        "line_spacing": pf.line_spacing,
        "runs": rinfo,
    })

tables = []
for ti, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        rows.append([cell.text.replace("\n", " / ") for cell in row.cells])
    tbl_pr = table._tbl.tblPr
    grid = []
    grid_el = table._tbl.tblGrid
    if grid_el is not None:
        grid = [int(c.get(qn("w")) or 0) for c in grid_el.findall(qn("gridCol"))]
    tables.append({
        "index": ti,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "autofit": bool(table.autofit),
        "grid_twips": grid,
        "sample": rows,
        "style": table.style.name if table.style else None,
    })

with zipfile.ZipFile(SRC) as zf:
    doc_xml = etree.fromstring(zf.read("word/document.xml"))
    sects = []
    for i, sect in enumerate(doc_xml.findall(".//w:sectPr", NS), start=1):
        pg = sect.find("w:pgSz", NS)
        mar = sect.find("w:pgMar", NS)
        cols = sect.find("w:cols", NS)
        sects.append({
            "index": i,
            "page_w_twips": int(pg.get(qn("w"))) if pg is not None and pg.get(qn("w")) else None,
            "page_h_twips": int(pg.get(qn("h"))) if pg is not None and pg.get(qn("h")) else None,
            "page_orient": None if pg is None else pg.get(qn("orient")),
            "margins_twips": {} if mar is None else {k: int(mar.get(qn(k))) for k in ("top", "right", "bottom", "left", "header", "footer", "gutter") if mar.get(qn(k))},
            "columns": {} if cols is None else {etree.QName(k).localname: v for k,v in cols.attrib.items()},
            "type": None if sect.find("w:type", NS) is None else sect.find("w:type", NS).get(qn("val")),
        })
    section_boundaries = []
    paragraph_ordinal = -1
    for body_child_index, child in enumerate(doc_xml.find("w:body", NS)):
        if child.tag != qn("p"):
            continue
        paragraph_ordinal += 1
        sect = child.find("w:pPr/w:sectPr", NS)
        if sect is None:
            continue
        text = "".join(child.xpath(".//w:t/text()", namespaces=NS))
        cols = sect.find("w:cols", NS)
        section_boundaries.append({
            "body_child_index": body_child_index,
            "document_paragraph_index": paragraph_ordinal,
            "paragraph_text": text,
            "columns": {} if cols is None else {etree.QName(k).localname: v for k,v in cols.attrib.items()},
            "type": None if sect.find("w:type", NS) is None else sect.find("w:type", NS).get(qn("val")),
        })
    parts = []
    for name in zf.namelist():
        if name.startswith("word/") and not name.endswith("/"):
            parts.append({"part": name, "size": zf.getinfo(name).file_size})
    styles_xml = etree.fromstring(zf.read("word/styles.xml"))
    styles = []
    for s in styles_xml.findall("w:style", NS):
        styles.append({
            "styleId": s.get(qn("styleId")),
            "type": s.get(qn("type")),
            "name": (s.find("w:name", NS).get(qn("val")) if s.find("w:name", NS) is not None else None),
        })

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps({
    "source": str(SRC),
    "paragraph_count": len(doc.paragraphs),
    "nonempty_paragraph_count": len(paragraphs),
    "paragraphs": paragraphs,
    "tables": tables,
    "sections_xml": sects,
    "section_boundaries": section_boundaries,
    "parts": parts,
    "styles": styles,
}, indent=2), encoding="utf-8")
print(OUT)
