from docx import Document
from pathlib import Path
import zipfile
p = Path(r"C:\Users\anju2\Desktop\Advance_java_research\version-1.1aHome_Energy_Tracker-publication_updated.docx")
doc = Document(str(p))
out = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker\research-output\_source_analysis\document_inventory.txt")
lines=[]
lines.append(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)} sections={len(doc.sections)}")
for i, para in enumerate(doc.paragraphs):
    text=para.text.replace('\t',' <TAB> ')
    if text or para._p.xpath('.//w:drawing'):
        lines.append(f"P{i:03d} style={para.style.name!r} drawing={bool(para._p.xpath('.//w:drawing'))}: {text}")
for ti,t in enumerate(doc.tables):
    lines.append(f"TABLE {ti} rows={len(t.rows)} cols={len(t.columns)} style={t.style.name if t.style else None}")
    for ri,row in enumerate(t.rows):
        lines.append(' | '.join(f"R{ri}: {cell.text.replace(chr(10),' / ')}" for cell in row.cells))
with zipfile.ZipFile(p) as z:
    media=[n for n in z.namelist() if n.startswith('word/media/')]
    lines.append('MEDIA: ' + ', '.join(media))
    target=Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker\research-output\_source_analysis\media")
    target.mkdir(exist_ok=True)
    for n in media:
        (target/Path(n).name).write_bytes(z.read(n))
out.write_text('\n'.join(lines), encoding='utf-8')
print(out)
