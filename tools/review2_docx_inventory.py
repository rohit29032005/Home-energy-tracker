"""Extract a compact, ordered inventory from a DOCX for Review-2 planning."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def trunc(text: str, limit: int = 260) -> str:
    text = " ".join(text.split())
    return text if len(text) <= limit else text[: limit - 3] + "..."


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: review2_docx_inventory.py INPUT.docx")
    doc = Document(Path(sys.argv[1]))
    print(f"sections={len(doc.sections)} paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
    print("\nPARAGRAPHS:")
    for i, p in enumerate(doc.paragraphs):
        text = trunc(p.text)
        if text:
            print(f"P{i:03d} style={p.style.name!r}: {text}")
    print("\nTABLES:")
    for ti, table in enumerate(doc.tables):
        print(f"TABLE {ti}: rows={len(table.rows)} cols={len(table.columns)}")
        for ri, row in enumerate(table.rows):
            cells = [trunc(c.text, 180) for c in row.cells]
            print(f"  R{ri}: {cells}")
    print("\nSECTIONS:")
    for i, section in enumerate(doc.sections):
        print(
            f"S{i}: page=({section.page_width},{section.page_height}) "
            f"margins=({section.left_margin},{section.right_margin},{section.top_margin},{section.bottom_margin}) "
            f"columns={section._sectPr.xpath('./w:cols/@w:num') or ['1']}"
        )


if __name__ == "__main__":
    main()
