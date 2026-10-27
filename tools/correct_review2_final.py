"""Apply only evidence-backed citation and bibliography corrections to Review-2."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document


SOURCE = Path(
    r"C:\Users\anju2\Desktop\Home-Energy-Tracker\artifacts\review2_work\Home_Energy_Tracker_Review2_Final.docx"
)
OUTPUT = Path(
    r"C:\Users\anju2\Desktop\Home-Energy-Tracker\artifacts\review2_correction\Home_Energy_Tracker_Review2_Final_Corrected.docx"
)
CHANGE_LOG = OUTPUT.with_suffix(".changes.json")


def remove_trailing_citation(document: Document, index: int, expected: str) -> dict[str, object]:
    paragraph = document.paragraphs[index]
    if paragraph.text != expected:
        raise RuntimeError(
            f"Paragraph {index} does not match the audited source text: {paragraph.text!r}"
        )
    if len(paragraph.runs) != 2:
        raise RuntimeError(f"Paragraph {index} expected two runs, found {len(paragraph.runs)}")
    removed = paragraph.runs[1].text
    if not removed.startswith(" ["):
        raise RuntimeError(f"Paragraph {index} does not end in a standalone citation run: {removed!r}")
    paragraph.runs[1].text = ""
    return {"paragraph": index, "operation": "removed_misplaced_citation", "removed": removed}


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    changes: list[dict[str, object]] = []
    changes.append(
        remove_trailing_citation(
            document,
            11,
            "The scope of this study is a local Home Energy Tracker rather than a calibrated utility-metering product. "
            "Its telemetry is simulator-compatible, and the study therefore evaluates software-path evidence, reproducibility, "
            "and data preparation rather than physical sensor calibration or billing correctness. [13], [15]",
        )
    )
    changes.append(
        remove_trailing_citation(
            document,
            13,
            "In the implemented system, user-facing requests are mediated by JWT, role, and ownership controls backed by Keycloak. "
            "Simulator-compatible telemetry is admitted through a device-bound token. These mechanisms establish an explicit "
            "admission boundary before subsequent event processing. [1], [3]",
        )
    )
    changes.append(
        remove_trailing_citation(
            document,
            14,
            "Accepted telemetry is passed to Kafka, from which the Usage Service writes time-series usage and windowed aggregations. "
            "InfluxDB retains time-series usage, while MySQL retains household, device, and alert state. This division supports "
            "inspection of the processing path and separates relational state from time-series storage. [1], [2]",
        )
    )
    changes.append(
        remove_trailing_citation(
            document,
            15,
            "The system additionally exposes deterministic bill forecasting, appliance ranking, standby detection, saving tips, and related insights. "
            "These functions are useful only when their inputs can be traced to admitted telemetry; the project therefore emphasises "
            "record handling and evidence boundaries rather than dashboard presentation alone. [11], [12]",
        )
    )
    changes.append(
        remove_trailing_citation(
            document,
            17,
            "The installed Random Forest is a legacy repository artefact. It can be reproduced under its documented local protocol, "
            "but the source, licence, label-generation procedure, and causal provenance of the supplied historical labelled CSV remain incomplete. "
            "Consequently, its metrics are retained as reproducibility evidence rather than generalisation evidence. [4], [8]",
        )
    )
    changes.append(
        remove_trailing_citation(
            document,
            19,
            "This paper documents the implemented architecture, the reproducible historical baseline, and a checksum-verified REFIT preprocessing path. "
            "It explicitly separates completed work from pending isolated provisioning, replay, alert-deduplication validation, "
            "and REFIT anomaly-accuracy evaluation. [1]–[3], [5]–[7]",
        )
    )
    changes.append(
        remove_trailing_citation(
            document,
            44,
            "Historical baseline metrics, endpoint smoke tests, and data-preprocessing checks answer different questions. "
            "The historical Random Forest metrics describe a supplied labelled CSV; the smoke test shows interface execution; "
            "and REFIT preparation demonstrates source integrity and deterministic transformation. Treating any one of these "
            "as a replacement for the others would overstate the evidence. [4], [8]–[10]",
        )
    )

    reference = document.paragraphs[139]
    old_reference = reference.text
    expected_reference_prefix = "[11] Dr. Karthika K, Thrupthi Chandana G, Taahani Sheikh, Sankeerthana, and Umme, "
    if not old_reference.startswith(expected_reference_prefix) or len(reference.runs) != 1:
        raise RuntimeError(f"Reference [11] source text unexpectedly changed: {old_reference!r}")
    reference.runs[0].text = old_reference.replace("[11] Dr. Karthika K,", "[11] Karthika K,", 1)
    changes.append(
        {
            "paragraph": 139,
            "operation": "removed_honorific_from_verified_author_entry",
            "old": old_reference,
            "new": reference.text,
        }
    )

    document.save(OUTPUT)
    CHANGE_LOG.write_text(json.dumps(changes, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {OUTPUT}")
    print(f"Recorded {len(changes)} targeted changes in {CHANGE_LOG}")


if __name__ == "__main__":
    main()
