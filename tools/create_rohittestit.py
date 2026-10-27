from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUTPUT = r"C:\Users\anju2\Desktop\Home-Energy-Tracker\rohittestit.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F4F7FA"
LIGHT = "D9D9D9"
GREY = "666666"
CODE_BG = "F1F3F5"
WHITE = "FFFFFF"
BLACK = "000000"


def shade(element, fill):
    tc_pr = element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LIGHT, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_font(run, name="Aptos", size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_para_spacing(p, before=0, after=6, line=1.12):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def page_number_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=8.5, color=GREY)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
section.page_width = Inches(8.5)
section.page_height = Inches(11)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)

for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
    style = styles[name]
    style.font.name = "Aptos Display" if name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

# Word's built-in Title style can carry a theme-colored rule. The guide uses
# typography and whitespace only, so remove any inherited paragraph border.
title_ppr = styles["Title"]._element.find(qn("w:pPr"))
if title_ppr is not None:
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)

if "Code Block" not in styles:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
code_style.font.name = "Consolas"
code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
code_style.font.size = Pt(8.3)


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, name="Aptos", size={1: 15, 2: 12, 3: 10.5}[level], bold=True)
    p.paragraph_format.keep_with_next = True
    return p


def add_text(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r)
    set_para_spacing(p)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_font(r)
    set_para_spacing(p, after=3)
    return p


_numbered_step = 1


def reset_numbering():
    global _numbered_step
    _numbered_step = 1


def add_numbered(text):
    global _numbered_step
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.27)
    r = p.add_run(f"{_numbered_step}. ")
    set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r)
    set_para_spacing(p, after=3)
    _numbered_step += 1
    return p


def add_code(code, keep_with_next=False):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep_with_next
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    p_pr.append(shd)
    r = p.add_run(code)
    set_font(r, name="Consolas", size=8.3, color=BLACK)
    return p


def add_label_value(label, value):
    p = doc.add_paragraph()
    r = p.add_run(label)
    set_font(r, bold=True)
    r = p.add_run(value)
    set_font(r)
    set_para_spacing(p, after=3)
    return p


def add_table(headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_cells = table.rows[0].cells
    for i, heading in enumerate(headers):
        cell = header_cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell)
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_spacing(p, before=0, after=0, line=1.0)
        r = p.add_run(str(heading))
        set_font(r, size=font_size, bold=True, color=WHITE)
    set_repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                shade(cell, PALE)
            set_cell_margins(cell)
            set_cell_border(cell)
            if widths:
                cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 and len(str(value)) < 18) else WD_ALIGN_PARAGRAPH.LEFT
            set_para_spacing(p, before=0, after=0, line=1.0)
            r = p.add_run(str(value))
            set_font(r, size=font_size)
        prevent_row_split(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_expected(lines):
    add_text("", bold_lead="Expected result: ")
    # Replace the empty lead paragraph with a compact explanation under the test.
    p = doc.paragraphs[-1]
    p.clear()
    r = p.add_run("Expected result: ")
    set_font(r, bold=True)
    r = p.add_run(lines)
    set_font(r)
    set_para_spacing(p, after=4)


def add_proves(lines):
    p = doc.add_paragraph()
    r = p.add_run("What this proves: ")
    set_font(r, bold=True)
    r = p.add_run(lines)
    set_font(r)
    set_para_spacing(p, after=7)


_page_break_calls = 0


def add_page_break():
    """Keep the cover separate; let later sections flow naturally.

    Forced later breaks created sparse continuation pages when a preceding
    code block crossed a page boundary. Headings and code blocks now carry
    their own keep settings, so normal pagination is cleaner.
    """
    global _page_break_calls
    if _page_break_calls == 0:
        doc.add_page_break()
    _page_break_calls += 1


# Footer for all pages.
footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para_spacing(footer_p, before=0, after=0, line=1.0)
r = footer_p.add_run("Home Energy Tracker  |  Testing and Demo Guide  |  Page ")
set_font(r, size=8.5, color=GREY)
page_number_field(footer_p)

# Cover page
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Home Energy Tracker Testing and Demo Guide")
set_font(r, name="Aptos Display", size=24, bold=True, color=BLACK)
set_para_spacing(p, after=8)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Commands Expected Results and Full Pipeline Verification")
set_font(r, size=13, color=BLACK)
set_para_spacing(p, after=17)

add_label_value("Purpose: ", "A copy-paste Windows PowerShell guide for testing the implemented frontend, API gateway, security controls, IoT telemetry pipeline, analytics, alerting, Wokwi simulator, observability, performance, resilience, and full end-to-end path.")
add_label_value("How to use it: ", "Run the sections in order. Each test states the command, the expected output, and what the result proves. Use the short demo plan near the end when presenting to a professor.")
add_label_value("Project location: ", r"C:\Users\anju2\Desktop\Home-Energy-Tracker")
add_label_value("Execution environment: ", "Windows PowerShell, Docker Desktop, local Docker Compose deployment, and optionally VS Code with PlatformIO and Wokwi Simulator.")

add_heading("Important operating rules", 2)
add_bullet("Keep Docker Desktop running before starting any test. The commands are designed for the local project environment.")
add_bullet("Do not show .env contents, JWTs, or device tokens in a recording or presentation. The local demo token in this guide must be replaced before any real deployment.")
add_bullet("Some tests intentionally write local test data. The end-to-end, load, Wokwi, and alert tests send telemetry. Run those only when you want new local evidence.")
add_bullet("Do not run docker compose down -v during a demo. That deletes local volumes and data.")
add_bullet("Run one resilience test at a time. The script restores the selected container automatically, but wait until the recovery line appears before continuing.")

add_heading("What a successful full pipeline means", 2)
add_text("A valid energy reading is accepted by the gateway, authenticated using the machine token, validated by the Ingestion Service, published to Kafka, consumed by the Usage and ML services, persisted or aggregated for time-series analysis, and then exposed through usage, insight, alert, frontend, and monitoring views.")
add_code("Wokwi ESP32 or PowerShell\n  -> API Gateway :8080\n  -> Ingestion Service\n  -> Kafka energy-usage topic\n  -> Usage Service -> InfluxDB -> Usage and Insight APIs -> Frontend\n  -> ML Inference -> Kafka alert event -> Alert Service -> MySQL and Mailpit")

add_page_break()

add_heading("1  Quick reference", 1)
add_text("Use this table during the demo. These are local tools; opening a page proves that the corresponding local endpoint is reachable, but the API and event checks later provide stronger functional proof.")
add_table(
    ["Component", "Address", "Use during the demonstration"],
    [
        ["Dashboard", "http://localhost:3000", "Login and show household, device, insight, and alert panels."],
        ["API Gateway", "http://localhost:8080", "Single public entry point for API and health checks."],
        ["Keycloak", "http://localhost:8180", "Local identity provider and login page."],
        ["Kafka UI", "http://localhost:8070", "Show energy-usage and alert topics."],
        ["InfluxDB", "http://localhost:8072", "Inspect time-series bucket and data explorer."],
        ["Mailpit", "http://localhost:8025", "Show captured local alert email."],
        ["Prometheus", "http://localhost:9090", "Show scrape targets and metric queries."],
        ["Grafana", "http://localhost:3001", "Show configured observability dashboards if required."],
    ],
    widths=[1.1, 2.0, 3.85],
    font_size=8.5,
)

add_heading("Test order and effect", 2)
add_table(
    ["Step", "Test area", "Writes data", "Typical time"],
    [
        ["1", "Startup and health", "No", "1 to 3 min"],
        ["2", "Frontend and authenticated dashboard APIs", "No", "1 min"],
        ["3", "Profile and device CRUD", "Yes, temporary metadata", "1 min"],
        ["4", "Telemetry, Kafka, InfluxDB, insights", "Yes, telemetry", "20 to 40 sec"],
        ["5", "Alerts and Mailpit", "Reads existing evidence or creates controlled test data", "1 to 2 min"],
        ["6", "Security validation", "One test telemetry event", "about 95 sec"],
        ["7", "Observability", "Evidence files only", "under 1 min"],
        ["8", "Performance", "Yes, about 450 telemetry requests", "about 105 sec"],
        ["9", "Resilience", "One fault and recovery evidence file", "1 to 3 min"],
        ["10", "Controlled full E2E", "Yes, synthetic users/devices/readings", "1 to 3 min"],
        ["11", "Optional Wokwi hardware simulation", "Yes, telemetry", "1 to 2 min"],
    ],
    widths=[0.5, 2.1, 2.45, 1.9],
    font_size=8.4,
)

add_heading("2  Start the stack and verify readiness", 1)
add_text("Open PowerShell. The first command moves to the project root. The environment file is copied only if it does not already exist.")
add_code("Set-Location 'C:\\Users\\anju2\\Desktop\\Home-Energy-Tracker'\n\nif (-not (Test-Path .env)) {\n  Copy-Item .env.example .env\n}\n\ndocker compose up -d --build\ndocker compose ps")
add_expected("Docker prints each service as started or running. On a cold build, image download and Java startup can take several minutes. The final list should show the configured Compose services as Up or healthy.")
add_proves("The local microservice stack has been started. It is not yet proof that all HTTP APIs are ready, so run the health commands next.")

add_code("(Invoke-WebRequest -UseBasicParsing 'http://localhost:8080/actuator/health').StatusCode\n(Invoke-WebRequest -UseBasicParsing 'http://localhost:3000/').StatusCode\n(Invoke-WebRequest -UseBasicParsing 'http://localhost:9090/-/ready').StatusCode\n(Invoke-WebRequest -UseBasicParsing 'http://localhost:3001/api/health').StatusCode\n(Invoke-WebRequest -UseBasicParsing 'http://localhost:8025/api/v1/messages').StatusCode")
add_expected("Each line prints 200. If the Gateway is still starting, wait 10 seconds and run its health command again.")
add_proves("The dashboard, gateway, Prometheus, Grafana, and local email-capture API are reachable.")

add_heading("If the gateway is not ready", 3)
add_code("docker compose logs --tail=120 api-gateway\ndocker compose logs --tail=80 user-service device-service ingestion-service usage-service alert-service insight-service")
add_expected("Look for Spring Boot startup completion and no repeated connection errors. The gateway becomes usable only after its dependent services are available.")

add_page_break()

add_heading("3  Authenticate once for protected API tests", 1)
add_text("The API gateway protects user, device, usage, insight, and alert routes with Keycloak JWTs. Run this block once in the same PowerShell window. It obtains a new short-lived token for the seeded demo household and creates reusable variables.")
add_code("$tokenBody = @{\n  grant_type = 'password'\n  client_id = 'home-energy-cli'\n  username = 'demo-user'\n  password = 'DemoUser-ChangeMe-2026!'\n}\n\n$userToken = (Invoke-RestMethod -Method Post `\n  -Uri 'http://localhost:8180/realms/home-energy/protocol/openid-connect/token' `\n  -Body $tokenBody -ContentType 'application/x-www-form-urlencoded').access_token\n\n$userHeaders = @{ Authorization = \"Bearer $userToken\" }\n$api = 'http://localhost:8080/api/v1'\n$demoUserId = 1\n$demoDeviceId = 1\n$tokenLine = Get-Content .env | Where-Object { $_ -match '^\\s*WOKWI_INGESTION_TOKEN=' } | Select-Object -First 1\n$deviceToken = ($tokenLine -replace '^\\s*WOKWI_INGESTION_TOKEN=', '').Trim().Trim('\\\"')\nif ([string]::IsNullOrWhiteSpace($deviceToken)) {\n  throw 'Set WOKWI_INGESTION_TOKEN in .env before running machine-token tests.'\n}")
add_expected("No error is printed. The $userToken variable contains a JWT, but do not display or share it. A token normally expires quickly; rerun this block if a later protected call returns 401.")
add_proves("The Keycloak realm, the demo account, and gateway JWT validation are working together.")

add_heading("Demo credentials", 2)
add_bullet("demo-user: household dashboard and owned resource tests for user ID 1.")
add_bullet("demo-user-two: cross-user authorization test.")
add_bullet("demo-admin: administrator-only creation and listing tests.")
add_text("All passwords used below are local demo credentials. They are not suitable for a production deployment.")

add_heading("4  Frontend and live dashboard API checks", 1)
reset_numbering()
add_numbered("Open http://localhost:3000 in a browser.")
add_numbered("Choose Sign In. Keycloak opens; sign in as demo-user using the local demo password.")
add_numbered("After the dashboard loads, show the household information, registered devices, energy usage, forecast, ranking, standby detective, saving tips, and alerts panels.")
add_text("The UI is a useful presentation layer, but use the authenticated calls below to prove the displayed data is coming from the live backend rather than a browser-only fallback.")
add_code("$routes = @(\n  \"$api/user/$demoUserId\",\n  \"$api/device/user/$demoUserId\",\n  \"$api/usage/$demoUserId?days=1\",\n  \"$api/alert/user/$demoUserId\",\n  \"$api/insight/bill-forecast/$demoUserId?days=3\",\n  \"$api/insight/appliance-ranking/$demoUserId?days=3\",\n  \"$api/insight/standby-detective/$demoUserId?days=3\",\n  \"$api/insight/saving-tips/$demoUserId\"\n)\n\n$routes | ForEach-Object {\n  $reply = Invoke-WebRequest -UseBasicParsing -Uri $_ -Headers $userHeaders\n  [pscustomobject]@{ Route = $_; Status = $reply.StatusCode; Bytes = $reply.Content.Length }\n} | Format-Table -AutoSize")
add_expected("Eight rows report Status 200 and a positive Bytes value. Exact response size changes as local data changes.")
add_proves("The gateway forwards authenticated requests correctly and the dashboard's core data routes return non-empty live JSON.")

add_page_break()

add_heading("5  Create a new profile and device metadata", 1)
add_text("This is a controlled metadata CRUD check. It proves MySQL-backed profile and device APIs. It creates a temporary application profile, not a new Keycloak login account. Keep the values in the same terminal so the optional cleanup command can remove them.")
add_code("$adminBody = @{\n  grant_type = 'password'\n  client_id = 'home-energy-cli'\n  username = 'demo-admin'\n  password = 'DemoAdmin-ChangeMe-2026!'\n}\n$adminToken = (Invoke-RestMethod -Method Post `\n  -Uri 'http://localhost:8180/realms/home-energy/protocol/openid-connect/token' `\n  -Body $adminBody -ContentType 'application/x-www-form-urlencoded').access_token\n$adminHeaders = @{ Authorization = \"Bearer $adminToken\" }\n$stamp = Get-Date -Format 'yyyyMMddHHmmss'")

add_code("$newProfile = @{\n  name = \"Demo Profile $stamp\"\n  surname = 'Local'\n  email = \"demo.profile.$stamp@example.test\"\n  address = 'Local test only'\n  alerting = $false\n  energyAlertingThreshold = 999999\n} | ConvertTo-Json\n\n$createdProfile = Invoke-RestMethod -Method Post -Uri \"$api/user\" `\n  -Headers $adminHeaders -ContentType 'application/json' -Body $newProfile\n$createdProfile | ConvertTo-Json -Depth 5\n$temporaryUserId = $createdProfile.id")
add_expected("A JSON object is returned with a new numeric id, name, email, alerting, and energyAlertingThreshold. The ID is different on every run.")
add_proves("The admin gateway route, User Service, and MySQL persistence path are working.")

add_code("$newDevice = @{\n  name = \"Demo Plug $stamp\"\n  type = 'GENERAL'\n  location = 'Local Test Desk'\n  userId = $temporaryUserId\n} | ConvertTo-Json\n\n$createdDevice = Invoke-RestMethod -Method Post -Uri \"$api/device\" `\n  -Headers $adminHeaders -ContentType 'application/json' -Body $newDevice\n$createdDevice | ConvertTo-Json -Depth 5\n$temporaryDeviceId = $createdDevice.id\n\nInvoke-RestMethod -Uri \"$api/device/$temporaryDeviceId\" -Headers $adminHeaders |\n  ConvertTo-Json -Depth 5")
add_expected("The created device has a new numeric id and the same name, type, location, and userId when read back.")
add_proves("Device metadata is persisted and retrievable through the protected gateway route.")

add_heading("Optional cleanup", 3)
add_code("(Invoke-WebRequest -UseBasicParsing -Method Delete `\n  -Uri \"$api/device/$temporaryDeviceId\" -Headers $adminHeaders).StatusCode\n\n(Invoke-WebRequest -UseBasicParsing -Method Delete `\n  -Uri \"$api/user/$temporaryUserId\" -Headers $adminHeaders).StatusCode")
add_expected("Both lines normally return 204. Run cleanup only after you have finished showing the created objects.")

add_heading("6  Telemetry ingestion and event pipeline", 1)
add_text("The ingestion endpoint is intentionally different from a user API: it requires the local machine device token instead of a user JWT. This lets an ESP32 or Wokwi simulator submit a telemetry reading without having a human login. The energyConsumed field is kWh, not watts.")
add_code("$machineHeaders = @{\n  'X-Device-Token' = $deviceToken\n}\n\n$usageBefore = Invoke-RestMethod -Uri \"$api/usage/$demoUserId?days=3\" -Headers $userHeaders\n$totalBefore = ($usageBefore.devices | Measure-Object -Property energyConsumed -Sum).Sum\nif ($null -eq $totalBefore) { $totalBefore = 0 }\n\n$reading = @{\n  deviceId = $demoDeviceId\n  energyConsumed = 0.25\n  timestamp = (Get-Date).ToUniversalTime().ToString('o')\n} | ConvertTo-Json -Compress\n\n$response = Invoke-WebRequest -UseBasicParsing -Method Post `\n  -Uri \"$api/ingestion\" -Headers $machineHeaders `\n  -ContentType 'application/json' -Body $reading\n$response.StatusCode")
add_expected("The final line is 201. It means the ingestion API accepted the telemetry message.")
add_proves("Gateway machine-token authentication and Ingestion Service validation are working. The 201 acknowledgement is intentionally fast; later processing is asynchronous.")

add_code("Start-Sleep -Seconds 15\n\n$usageAfter = Invoke-RestMethod -Uri \"$api/usage/$demoUserId?days=3\" -Headers $userHeaders\n$totalAfter = ($usageAfter.devices | Measure-Object -Property energyConsumed -Sum).Sum\nif ($null -eq $totalAfter) { $totalAfter = 0 }\n\n[pscustomobject]@{\n  TotalBeforeKwh = [math]::Round([double]$totalBefore, 3)\n  TotalAfterKwh = [math]::Round([double]$totalAfter, 3)\n  ExpectedMinimumIncreaseKwh = 0.25\n} | Format-List")
add_expected("TotalAfterKwh is at least approximately 0.25 kWh higher in the same time window. Small variation is possible if other tests are also sending readings.")
add_proves("Kafka consumers and time-series aggregation processed the accepted telemetry event.")

add_heading("Show the asynchronous pipeline", 2)
add_code("docker exec kafka /opt/kafka/bin/kafka-topics.sh `\n  --bootstrap-server kafka:9092 --list\n\ndocker exec kafka /opt/kafka/bin/kafka-console-consumer.sh `\n  --bootstrap-server kafka:9092 --topic energy-usage `\n  --from-beginning --max-messages 5 --timeout-ms 5000\n\nInvoke-RestMethod 'http://localhost:8072/health'")
add_expected("Kafka topic list includes energy-usage. The consumer prints event JSON containing a deviceId, energyConsumed value, and timestamp. The InfluxDB health response indicates healthy/pass status.")
add_proves("The event broker and time-series store are part of the real pipeline, not only a direct REST dashboard.")

add_page_break()

add_heading("7  Analytics and explainable energy insights", 1)
add_text("Run the following after telemetry has been accepted and the short consumer wait has completed. These endpoints are deterministic analysis outputs over the time window selected by days.")
add_heading("Bill shock forecast", 2)
add_code("Invoke-RestMethod -Uri \"$api/insight/bill-forecast/$demoUserId?days=3\" `\n  -Headers $userHeaders | ConvertTo-Json -Depth 6")
add_expected("JSON includes currentKwh, currentEstimatedCost, dailyRateKwh, projectedMonthlyKwh, projectedMonthlyCost, and tariffSummary.")
add_proves("Observed usage is converted into a transparent monthly forecast using the configured tariff logic.")

add_heading("Appliance ranking", 2)
add_code("Invoke-RestMethod -Uri \"$api/insight/appliance-ranking/$demoUserId?days=3\" `\n  -Headers $userHeaders | ConvertTo-Json -Depth 8")
add_expected("A rankings collection lists devices in descending energy usage with consumption percentage and estimated cost fields.")
add_proves("The system identifies the highest contributing appliance over the queried period.")

add_heading("Standby Energy Detective", 2)
add_code("Invoke-RestMethod -Uri \"$api/insight/standby-detective/$demoUserId?days=3\" `\n  -Headers $userHeaders | ConvertTo-Json -Depth 8")
add_expected("The response contains baselineWatts, standbyDailyKwh, standbyMonthlyKwh, standbyMonthlyCost, riskLevel, recommendations, and probableCategories when sufficient usage data is available.")
add_proves("The project's novel explainable heuristic converts a low-activity baseline into estimated monthly energy cost and action-oriented recommendations. It is not a claim of exact appliance-level disaggregation.")

add_heading("Savings tips", 2)
add_code("Invoke-RestMethod -Uri \"$api/insight/saving-tips/$demoUserId\" `\n  -Headers $userHeaders | ConvertTo-Json -Depth 8")
add_expected("The response contains recommendations and energy/cost context. If a local AI model is unavailable, a safe deterministic fallback can return isAiGenerated as false; that fallback is correct operation, not a failure.")
add_proves("The recommendation endpoint remains usable even when its optional natural-language AI dependency is unavailable.")

add_heading("8  Alerts email delivery and deduplication", 1)
add_text("The safest way to show alert creation and email delivery is the controlled end-to-end suite in Section 13. For a quick read-only demonstration, show existing alert history and the local Mailpit inbox first.")
add_code("Invoke-RestMethod -Uri \"$api/alert/user/$demoUserId\" -Headers $userHeaders |\n  Select-Object id, alertType, message, threshold, energyConsumed, email, createdAt, sent |\n  Format-Table -AutoSize\n\nInvoke-RestMethod 'http://localhost:8025/api/v1/messages' |\n  ConvertTo-Json -Depth 6")
add_expected("Alert history contains persisted records with a sent flag. Mailpit returns a messages collection when alert email delivery has occurred.")
add_proves("Alert state is persisted separately from telemetry and local email delivery is observable without sending mail to an external address.")

add_code("docker compose logs --tail=200 alert-service |\n  Select-String 'Suppressed duplicate|Email sent|ML result'", keep_with_next=True)
add_expected("After an alert scenario, logs may show a sent email, an ML result, or Suppressed duplicate. Duplicate suppression is expected when an equivalent alert repeats within the configured 15-minute window.")
add_proves("Alert deduplication prevents repeated notifications for the same alert key within the configured window.")

add_page_break()

add_heading("9  Security validation", 1)
add_text("This script runs eight focused authorization and authentication checks through the gateway. It deliberately waits long enough for a real Keycloak token to expire. The wait is expected; do not interrupt it.")
add_code("Set-Location 'C:\\Users\\anju2\\Desktop\\Home-Energy-Tracker'\n\npowershell -NoProfile -ExecutionPolicy Bypass `\n  -File .\\tools\\security\\run-security-validation.ps1 `\n  -DeviceToken $deviceToken")
add_expected("After about 95 seconds, PowerShell prints Security validation passed. Evidence is written to evidence\\06_security. The CSV shows eight passed cases: owned profile 200, missing/malformed token 401, cross-user access 403, admin list 200, machine ingestion 201, and expired JWT 401.")
add_proves("Authentication, ownership enforcement, administrative access, IoT machine-token access, and actual token-expiry handling all work at the gateway boundary.")

add_heading("10  Observability and monitoring", 1)
add_code("powershell -NoProfile -ExecutionPolicy Bypass `\n  -File .\\tools\\observability\\collect-runtime-observability.ps1\n\nGet-ChildItem .\\evidence\\08_observability |\n  Sort-Object LastWriteTime -Descending | Select-Object -First 8 Name, LastWriteTime")
add_expected("The script prints Runtime observability evidence captured and creates timestamped gateway metrics, Prometheus targets, alert deduplication metric, ML metric, and Docker resource files.")
add_proves("Metrics and health evidence can be collected from the live running stack rather than inferred from UI appearance.")

add_text("Open http://localhost:9090/targets. A healthy demonstration shows active targets with health up. The final recorded validation captured 8 active targets and 0 down targets; this number should be confirmed again from the current page during the presentation.")

add_heading("11  Performance test using the official k6 Docker image", 1)
add_text("This test sends real telemetry messages to device 1. It uses Docker because a native k6 installation is not required. Do not run it immediately before a clean-data demo, because it produces approximately 450 local events.")
add_code("powershell -NoProfile -ExecutionPolicy Bypass `\n  -File .\\tools\\performance\\run-load-test.ps1 `\n  -DeviceToken $deviceToken `\n  -UseDockerK6")
add_expected("The run takes about 105 seconds. The k6 summary displays http_req_duration and http_req_failed threshold checks. Success includes p(95) < 2000 ms and error rate < 0.05, then Load test completed with evidence paths.")
add_proves("The same authenticated ingestion contract used by Wokwi was exercised under a short local ramped arrival rate.")

add_table(
    ["Reference measurement", "Observed local validation on 05 Oct 2026"],
    [
        ["Accepted requests", "450 of 450"],
        ["HTTP failure rate", "0.00 percent"],
        ["Mean duration", "39.44 ms"],
        ["p95 duration", "109.77 ms"],
        ["Scenario scope", "Single-host local test, not a distributed production benchmark"],
    ],
    widths=[2.35, 4.55],
    font_size=8.7,
)

add_page_break()

add_heading("12  Resilience and recovery test", 1)
add_text("The fault-injection script stops one local Compose service, captures behavior while it is stopped, then starts it again and waits for container-state recovery. Start with Mailpit, which has the smallest user impact. Do not run multiple scenarios in parallel.")
add_code("powershell -NoProfile -ExecutionPolicy Bypass `\n  -File .\\tools\\resilience\\run-fault-injection.ps1 `\n  -Target mailpit `\n  -DeviceToken $deviceToken")
add_expected("Docker reports the selected container stopping and starting. The final output says Fault scenario captured in evidence\\10_resilience and recovered=True appears in the created report.")
add_proves("The system records one controlled dependency outage and container recovery. It does not by itself prove every alert was delivered during the outage.")

add_text("A second useful scenario is ML inference recovery. Run it only after Mailpit has returned to running state.")
add_code("powershell -NoProfile -ExecutionPolicy Bypass `\n  -File .\\tools\\resilience\\run-fault-injection.ps1 `\n  -Target ml-inference-service `\n  -DeviceToken $deviceToken")
add_expected("The log records an ingestion HTTP status while the ML service is stopped, then the worker starts again. The recorded local run showed gateway ingestion 201 during the fault and post-restart ML/Kafka processing.")
add_proves("Telemetry acceptance is decoupled from the optional ML worker, and the worker can recover after a controlled restart.")

add_heading("13  Controlled full end to end verification", 1)
add_text("This is the strongest automated proof. It creates isolated synthetic users and devices, sends telemetry through the simulator-compatible ingestion endpoint, checks analytics assertions, triggers one threshold alert, verifies its persisted alert record and a new matching Mailpit email, then disables alerting for the synthetic alert user.")
add_code("powershell -NoProfile -ExecutionPolicy Bypass `\n  -File .\\tools\\run_review2_e2e_verification.ps1 `\n  -DeviceToken $deviceToken `\n  -EvidenceDirectory .\\evidence\\13_end_to_end")
add_expected("The JSON result ends with overallPassed: true. Look for accepted telemetry statusCode 201, analytics assertions marked passed, deliveryObserved: true, and a timestamped review2_e2e_results JSON file.")
add_proves("One automated run checks the actual end-to-end pipeline from telemetry API through event processing, usage and insight responses, alert persistence, and email capture.")

add_heading("Reference values from the recorded successful local E2E run", 2)
add_table(
    ["Assertion", "Recorded value", "Interpretation"],
    [
        ["Analytics availability", "3.79 seconds", "Controlled telemetry became available to the analytics path."],
        ["Controlled usage", "1.50 kWh", "Sum of the three isolated telemetry readings."],
        ["Current estimated cost", "$0.18", "Tariff calculation for the controlled dataset."],
        ["Monthly forecast", "15.00 kWh and $11.80", "Projection for the fixed three-day controlled case."],
        ["Alert delivery observation", "13.64 seconds", "Threshold alert persisted and a new Mailpit email was observed."],
    ],
    widths=[1.95, 1.65, 3.3],
    font_size=8.5,
)
add_text("These are reference values from the saved evidence run. New runs create unique synthetic IDs and may have different timing; use overallPassed and the current evidence file as the primary pass/fail proof.")

add_page_break()

add_heading("14  Optional Wokwi ESP32 simulation", 1)
add_text("Use this section to demonstrate an IoT-style input rather than a PowerShell-injected reading. Open only the iot-simulator folder in VS Code. The current local simulator connects directly to the host API Gateway; no Cloudflare tunnel and no HTTPS configuration are required for this setup.")
reset_numbering()
add_numbered("Ensure the Docker Gateway health check in Section 2 returns 200.")
add_numbered("In VS Code, open C:\\Users\\anju2\\Desktop\\Home-Energy-Tracker\\iot-simulator as the workspace folder.")
add_numbered("Confirm that PlatformIO IDE and Wokwi Simulator extensions are installed. If required, use F1, then Wokwi: Request a new License.")
add_numbered("Review local include\\config.h. Keep the URL as host.wokwi.internal rather than localhost because the ESP32 runs in a virtual environment.")
add_code("constexpr char INGESTION_URL[] =\n  \"http://host.wokwi.internal:8080/api/v1/ingestion\";\n#define HOME_ENERGY_INGESTION_DEVICE_TOKEN \"<paste WOKWI_INGESTION_TOKEN from the root .env file>\"\nconstexpr long DEVICE_ID = 1;")
add_text("The config file is local and ignored by Git. Do not commit a real production token. DEVICE_ID must match a registered device; the seeded demo uses device 1.")
add_numbered("Press F1, run PlatformIO: Build, and wait for a successful ESP32 build. The first build downloads the framework and Adafruit libraries, so it can take several minutes.")
add_numbered("Press F1, run Wokwi: Start Simulator, and open the Serial Monitor at 115200 if it does not open automatically.")
add_numbered("Wait for WiFi connected and API: 201. Rotate the potentiometer or press a button to produce a new reading.")

add_table(
    ["Control", "Firmware mode", "Expected simulated value", "Presentation use"],
    [
        ["Top button", "ECO", "approximately 1 to 6 kWh", "Show a normal low-use household reading."],
        ["Middle button", "HEAVY", "approximately 8 to 20 kWh", "Show a higher but ordinary energy reading."],
        ["Bottom button", "ALERT TEST", "approximately 26 to 40 kWh", "Crosses the seeded user 1 threshold of 25 kWh."],
        ["Potentiometer", "All modes", "Adjusts the mode range", "Show an input value changing before it is sent."],
    ],
    widths=[1.15, 1.25, 1.9, 2.6],
    font_size=8.4,
)

add_expected("The OLED displays HOME ENERGY TRACKER, a mode, a kWh reading, and API: 201. The serial output shows Connecting to Wokwi WiFi, WiFi connected: <10.x.x.x>, JSON payload, and an accepted server response.")
add_proves("ESP32 firmware, WiFi, HTTP payload structure, machine-token authentication, and the same ingestion endpoint used by the web API are all connected.")
add_text("For an alert demonstration, press the bottom button once, wait for API: 201, stop the simulator to avoid repeated readings, then wait about 10 to 15 seconds and show alert history or Mailpit. Treat Wokwi as an API and firmware simulation, not as a calibrated physical mains meter.")

add_heading("15  Build and regression check", 1)
add_code("Set-Location 'C:\\Users\\anju2\\Desktop\\Home-Energy-Tracker\\api-gateway'\nmvn -B test")
add_expected("Maven reports BUILD SUCCESS. The focused gateway filter suite reports Tests run: 4, Failures: 0, Errors: 0, Skipped: 0.")
add_proves("The regression tests cover the gateway's JWT ownership filtering and IoT token forwarding behavior after the response-body fix.")

add_page_break()

add_heading("16  Professor demo sequence", 1)
add_text("Use this shorter order when time is limited. It shows the most meaningful proof without running every long test live. Keep a browser, one PowerShell terminal, and the Wokwi VS Code workspace ready before the presentation.")
add_table(
    ["Time", "Show", "Say or point out", "Evidence"],
    [
        ["0:00 to 0:45", "Dashboard login", "This is an event-driven household energy platform, not only a static dashboard.", "Live Keycloak login and dashboard."],
        ["0:45 to 1:30", "Gateway health and dashboard APIs", "The frontend reads authenticated data through one public gateway.", "HTTP 200 and non-empty JSON."],
        ["1:30 to 2:30", "One telemetry reading", "A reading is accepted quickly and then processed asynchronously.", "201 from ingestion."],
        ["2:30 to 3:15", "Kafka and usage", "Kafka decouples ingestion from analytics; time-series data is queried later.", "Kafka event plus updated usage."],
        ["3:15 to 4:15", "Forecast, ranking, standby", "The system estimates bill impact, identifies top devices, and explains standby waste.", "Three insight JSON results."],
        ["4:15 to 5:00", "Alert and Mailpit", "Alerts are persisted and email is captured locally for safe demonstration.", "Alert history and Mailpit message."],
        ["5:00 to 6:00", "Wokwi or E2E evidence", "The ESP32 simulator uses the same telemetry contract; the E2E suite verifies the full path.", "OLED API: 201 or overallPassed: true."],
        ["6:00 to 6:45", "Security and quality", "Users cannot cross access boundaries; the IoT machine token is separate.", "8 security cases passed and selected evidence files."],
    ],
    widths=[0.85, 1.55, 2.65, 1.75],
    font_size=8.1,
)

add_heading("Short presentation statements", 2)
add_bullet("The reading path is asynchronous: the client receives 201 after accepted ingestion, while Kafka consumers handle usage aggregation, ML inference, and alert events independently.")
add_bullet("MySQL stores users, devices, thresholds, and alert history. InfluxDB stores time-series energy measurements. Kafka carries events between services.")
add_bullet("The Standby Energy Detective is an explainable heuristic that estimates persistent low-activity load and converts it into monthly cost. It does not claim exact electrical disaggregation.")
add_bullet("The current proof is local Docker Compose. Do not claim cloud deployment or calibrated physical smart-meter accuracy unless those are separately deployed and measured.")

add_heading("17  Troubleshooting guide", 1)
add_table(
    ["Symptom", "Likely reason", "What to do"],
    [
        ["Vite proxy ECONNREFUSED", "Frontend is up but Gateway is not reachable.", "Check http://localhost:8080/actuator/health, then docker compose logs --tail=120 api-gateway."],
        ["401 on a protected endpoint", "JWT expired or missing.", "Rerun Section 3 to obtain a fresh demo-user token."],
        ["403 on user or device data", "Token belongs to a different household.", "Use demo-user for user ID 1 or demo-admin for administrative operations."],
        ["401 on ingestion", "Missing or incorrect X-Device-Token.", "Use the local machine token and do not replace it with a user JWT."],
        ["201 but usage has not changed", "Kafka consumer is still processing.", "Wait 15 seconds, then query usage again and inspect usage-service logs."],
        ["Wokwi API error negative code", "ESP32 cannot reach the host URL or Gateway is not ready.", "Keep host.wokwi.internal in config.h and verify Gateway health on port 8080."],
        ["No Wokwi display build", "First PlatformIO build is downloading packages.", "Wait for PlatformIO: Build to finish; confirm Adafruit GFX and SSD1306 dependencies."],
        ["Mailpit empty", "No new alert has completed or message was deduplicated.", "Use the E2E suite for controlled proof; then refresh http://localhost:8025."],
        ["Performance test cannot start", "Docker or k6 image unavailable.", "Confirm docker version, then rerun with -UseDockerK6."],
    ],
    widths=[1.45, 2.05, 3.3],
    font_size=8.1,
)

add_heading("18  Final readiness checklist", 1)
add_table(
    ["Check", "Pass criterion", "Done"],
    [
        ["Platform", "Gateway health is 200; dashboard and Keycloak login work.", "[ ]"],
        ["Dashboard APIs", "Eight protected routes show 200 and non-empty JSON.", "[ ]"],
        ["Telemetry path", "Ingestion returns 201; usage or Kafka event can be shown.", "[ ]"],
        ["Insights and alert", "Insight JSON plus alert history or Mailpit is available.", "[ ]"],
        ["Security and monitoring", "Security script passes 8 cases; Prometheus evidence is available.", "[ ]"],
        ["Wokwi optional", "OLED or serial output shows API: 201.", "[ ]"],
        ["Strongest proof", "E2E result reports overallPassed: true.", "[ ]"],
    ],
    widths=[1.8, 4.55, 0.45],
    font_size=8.2,
)

doc.core_properties.title = "Home Energy Tracker Testing and Demo Guide"
doc.core_properties.subject = "Commands Expected Results and Full Pipeline Verification"
doc.core_properties.author = "Home Energy Tracker Project"
doc.core_properties.comments = "Local project verification guide"

doc.save(OUTPUT)
print(OUTPUT)
