from docx import Document

SOURCE = r"C:\Users\anju2\Desktop\Home-Energy-Tracker\artifacts\review2_work\Home_Energy_Tracker_Review2_Final.docx"

document = Document(SOURCE)
print(f"PARAGRAPHS {len(document.paragraphs)} TABLES {len(document.tables)}")
needles = (
    "legacy repository artefact",
    "deterministic bill forecasting",
    "Historical baseline metrics",
)
for index, paragraph in enumerate(document.paragraphs):
    value = paragraph.text.strip()
    if value.startswith("[") or value == "References" or "[" in value or any(needle in value for needle in needles):
        print(f"IDX {index} RUNS {len(paragraph.runs)} STYLE {paragraph.style.name} {value!r}")
        if index in {11, 13, 14, 15, 16, 17, 19, 33, 37, 44, 139}:
            for run_index, run in enumerate(paragraph.runs):
                print(
                    f"  RUN {run_index} text={run.text!r} "
                    f"bold={run.bold} italic={run.italic} "
                    f"font={run.font.name} size={run.font.size}"
                )
