"""Static and PDF validation for the Review-2 manuscript deliverables.

This script is intentionally read-only.  It checks the supplied final DOCX and
the corresponding PDF for required manuscript structure, evidence-boundary
text, OOXML integrity, and basic render completeness.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import pdfplumber
from docx import Document
from lxml import etree
from pypdf import PdfReader


TITLE = (
    "A Microservices-Based Home Energy Tracker for Event-Driven Monitoring, "
    "Deduplicated Alerting, and ML-Based Anomaly Detection"
)
AUTHOR_TOKENS = (
    "Rohit Shrivastava",
    "Harsh Gupta",
    "Anshika Rai",
    "Syamasudha Veeragandham",
    "Vellore Institute of Technology",
    "rohit.lalbahadur2023@vitstudent.ac.in",
    "harsh.gupta2023@vitstudent.ac.in",
    "anshika.rai2023@vitstudent.ac.in",
    "syamasudha.v@vit.ac.in",
    "Corresponding author",
)
REQUIRED_HEADINGS = (
    "1. Introduction",
    "1.1 Research Objectives",
    "2. Related Work",
    "2.1 Research Gaps Identified",
    "Problem Statement",
    "3. Materials and Methods",
    "4. Experimental Results and Discussion",
    "5. Limitations and Future Work",
    "6. Conclusion",
    "CRediT Authorship Contribution Statement",
    "Funding Statement",
    "Declaration of Competing Interests",
    "Data Availability",
    "References",
)
PLACEHOLDERS = ("[TO BE PROVIDED]", "<TO BE PROVIDED>", "Lorem ipsum")
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text))


def abstract_body(paragraphs: list[str]) -> str:
    for paragraph in paragraphs:
        if paragraph.startswith("Abstract"):
            if "—" in paragraph:
                return paragraph.split("—", 1)[1].strip()
            if "-" in paragraph:
                return paragraph.split("-", 1)[1].strip()
    fail("abstract paragraph not found")


def keyword_items(paragraphs: list[str]) -> list[str]:
    for paragraph in paragraphs:
        if paragraph.startswith("Keywords"):
            segment = paragraph.split("—", 1)[-1] if "—" in paragraph else paragraph.split("-", 1)[-1]
            return [item.strip() for item in segment.split(";") if item.strip()]
    fail("keywords paragraph not found")


def validate_docx(docx_path: Path) -> dict[str, object]:
    try:
        with ZipFile(docx_path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                fail(f"DOCX ZIP corrupt at {bad_member}")
            document_xml = archive.read("word/document.xml")
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
    except BadZipFile as exc:
        fail(f"DOCX is not a valid package: {exc}")

    document = Document(docx_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    all_text = "\n".join(paragraphs + [cell.text for table in document.tables for row in table.rows for cell in row.cells])

    checks: list[tuple[str, bool]] = [("title", TITLE in all_text)]
    checks.extend((f"author detail: {token}", token in all_text) for token in AUTHOR_TOKENS)
    checks.extend((f"heading: {heading}", heading in all_text) for heading in REQUIRED_HEADINGS)
    checks.extend((f"citation [{index}]", f"[{index}]" in all_text) for index in range(1, 19))
    checks.extend((f"no placeholder: {placeholder}", placeholder not in all_text) for placeholder in PLACEHOLDERS)
    failures = [label for label, passed in checks if not passed]
    if failures:
        fail("DOCX check(s) failed: " + "; ".join(failures))

    abstract_words = word_count(abstract_body(paragraphs))
    if not 150 <= abstract_words <= 250:
        fail(f"abstract word count {abstract_words} is outside 150-250")
    keywords = keyword_items(paragraphs)
    if len(keywords) != 8:
        fail(f"keyword count {len(keywords)} is not eight")

    intro_start = paragraphs.index("1. Introduction") + 1
    objective_start = paragraphs.index("1.1 Research Objectives")
    introduction = paragraphs[intro_start:objective_start]
    uncited_intro = [text[:72] for text in introduction if not re.search(r"\[\d+", text)]
    if len(introduction) != 10:
        fail(f"introduction paragraph count {len(introduction)} is not ten")
    if uncited_intro:
        fail("introduction paragraph(s) without citations: " + " | ".join(uncited_intro))

    references = [text for text in paragraphs if re.match(r"^\[\d+\]", text)]
    if len(references) != 18:
        fail(f"reference count {len(references)} is not 18")
    if len(document.tables) != 5:
        fail(f"table count {len(document.tables)} is not 5")
    if len(media) != 2:
        fail(f"embedded media count {len(media)} is not 2")

    xml = etree.fromstring(document_xml)
    math_count = len(xml.xpath(".//m:oMath", namespaces={"m": M_NS}))
    if math_count < 2:
        fail(f"native equation count {math_count} is below two")
    columns = xml.xpath(".//w:sectPr/w:cols/@w:num", namespaces={"w": W_NS})
    if any(value != "1" for value in columns):
        fail(f"found non-single-column section setting(s): {columns}")

    # Evidence-boundary language must remain explicit rather than being converted into completion claims.
    boundary_terms = (
        "not been completed",
        "not establish real-world anomaly accuracy",
        "not run concurrently",
        "pending",
    )
    missing_boundary = [term for term in boundary_terms if term not in all_text]
    if missing_boundary:
        fail("evidence-boundary wording missing: " + ", ".join(missing_boundary))

    return {
        "paragraphs": len(paragraphs),
        "tables": len(document.tables),
        "media": len(media),
        "abstract_words": abstract_words,
        "keywords": len(keywords),
        "cited_introduction_paragraphs": len(introduction),
        "references": len(references),
        "native_equations": math_count,
    }


def validate_pdf(pdf_path: Path) -> dict[str, object]:
    reader = PdfReader(pdf_path)
    if reader.is_encrypted:
        fail("PDF is encrypted")
    page_count = len(reader.pages)
    if page_count != 12:
        fail(f"PDF page count {page_count} is not 12")
    sizes = []
    for page in reader.pages:
        box = page.mediabox
        sizes.append((round(float(box.width), 1), round(float(box.height), 1)))
    # ISO A4: 595.3 x 841.9 PostScript points.
    if any(abs(width - 595.3) > 2 or abs(height - 841.9) > 2 for width, height in sizes):
        fail(f"PDF page size is not A4: {sizes}")
    with pdfplumber.open(pdf_path) as pdf:
        text_lengths = [len((page.extract_text() or "").strip()) for page in pdf.pages]
    empty_pages = [str(index + 1) for index, length in enumerate(text_lengths) if length < 100]
    if empty_pages:
        fail("PDF has an empty or near-empty page: " + ", ".join(empty_pages))
    return {"pages": page_count, "page_size": sizes[0], "minimum_text_characters": min(text_lengths)}


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: validate_review2_final.py FINAL.docx FINAL.pdf")
    docx_path, pdf_path = map(Path, sys.argv[1:])
    if not docx_path.is_file():
        fail(f"DOCX missing: {docx_path}")
    if not pdf_path.is_file():
        fail(f"PDF missing: {pdf_path}")
    docx_results = validate_docx(docx_path)
    pdf_results = validate_pdf(pdf_path)
    print("PASS: Review-2 final deliverables validated")
    for key, value in docx_results.items():
        print(f"DOCX {key}: {value}")
    for key, value in pdf_results.items():
        print(f"PDF {key}: {value}")


if __name__ == "__main__":
    main()
