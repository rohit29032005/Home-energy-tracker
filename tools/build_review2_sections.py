from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker")
OUTPUT_DIR = ROOT / "research-output"
FIGURE_DIR = OUTPUT_DIR / "review2-figures"
OUTPUT_DOCX = OUTPUT_DIR / "research_review2_sections_3_to_6.docx"

NAVY = "1F4E79"
BLUE_FILL = "EAF2F8"
PALE_BLUE = "F4F8FB"
GRAY_FILL = "F2F2F2"
GRID = "D9D9D9"
TEXT = "000000"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=11, bold=False, italic=False, name="Times New Roman", color=TEXT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc, text, after=5, before=0, indent=0, justify=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Cm(indent)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text)
    set_run_font(r, 10, italic=True)
    return p


def add_figure(paragraph, image_path, width, alt_text):
    """Insert an accessible figure with Word title and description metadata."""
    shape = paragraph.add_run().add_picture(str(image_path), width=width)
    shape._inline.docPr.set("title", alt_text)
    shape._inline.docPr.set("descr", alt_text)
    return shape


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(11 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    set_keep_with_next(p)
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, 12, bold=True)
    elif level == 2:
        set_run_font(r, 11, bold=True)
    else:
        set_run_font(r, 10.5, bold=True)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(text)
    set_run_font(r, 9, italic=True)
    return p


def add_table(doc, caption, headers, rows, widths=None):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = header_cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, 8.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, 8.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def text_lines(text, font, max_width, draw):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else current + " " + word
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw, box, text, font, fill="#111111", line_gap=7):
    x1, y1, x2, y2 = box
    lines = text_lines(text, font, x2 - x1 - 26, draw)
    bboxes = [draw.textbbox((0, 0), line, font=font) for line in lines]
    heights = [b[3] - b[1] for b in bboxes]
    total_h = sum(heights) + line_gap * max(0, len(lines) - 1)
    y = y1 + ((y2 - y1 - total_h) / 2)
    for line, bbox, height in zip(lines, bboxes, heights):
        line_w = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - line_w) / 2, y), line, font=font, fill=fill)
        y += height + line_gap


def draw_box(draw, box, title, detail, title_font, body_font, fill=BLUE_FILL):
    if not fill.startswith("#"):
        fill = "#" + fill
    draw.rectangle(box, fill=fill, outline="#1F4E79", width=3)
    x1, y1, x2, y2 = box
    title_h = 32
    draw.rectangle((x1, y1, x2, y1 + title_h), fill="#1F4E79", outline="#1F4E79")
    draw_centered_text(draw, (x1 + 5, y1 + 2, x2 - 5, y1 + title_h - 2), title, title_font, fill="white", line_gap=2)
    draw_centered_text(draw, (x1 + 8, y1 + title_h + 3, x2 - 8, y2 - 5), detail, body_font)


def draw_arrow(draw, start, end, label=None, font=None, color="#3F4E5A"):
    draw.line((start[0], start[1], end[0], end[1]), fill=color, width=4)
    dx, dy = end[0] - start[0], end[1] - start[1]
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    head = 14
    p1 = (end[0], end[1])
    p2 = (end[0] - head * ux + 6 * px, end[1] - head * uy + 6 * py)
    p3 = (end[0] - head * ux - 6 * px, end[1] - head * uy - 6 * py)
    draw.polygon([p1, p2, p3], fill=color)
    if label and font:
        mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
        bbox = draw.textbbox((0, 0), label, font=font)
        draw.rectangle((mx - (bbox[2] - bbox[0]) / 2 - 4, my - 12, mx + (bbox[2] - bbox[0]) / 2 + 4, my + 10), fill="white")
        draw.text((mx - (bbox[2] - bbox[0]) / 2, my - 10), label, font=font, fill="#1F4E79")


def get_font(size, bold=False):
    base = Path(r"C:\Windows\Fonts")
    candidate = base / ("arialbd.ttf" if bold else "arial.ttf")
    if candidate.exists():
        return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def make_architecture_figure(path):
    image = Image.new("RGB", (2200, 1160), "white")
    draw = ImageDraw.Draw(image)
    title_font = get_font(19, True)
    body_font = get_font(18, False)
    stage_font = get_font(21, True)
    label_font = get_font(16, False)

    for x, label in [(80, "Input and Interface Layer"), (500, "Event Processing Layer"), (1220, "Data and Analytics Layer"), (1710, "Persistence and Delivery Layer")]:
        draw.text((x, 48), label, font=stage_font, fill="#1F4E79")

    draw_box(draw, (80, 145, 390, 345), "Virtual Smart Meter", "Wokwi ESP32 sends deviceId, energyConsumed, and timestamp", title_font, body_font)
    draw_box(draw, (80, 515, 390, 715), "Web Dashboard", "Household profiles, device registration, usage views, insights, and alert history", title_font, body_font, fill="F7FBFE")
    draw_box(draw, (490, 300, 780, 510), "API Gateway", "Single REST entry point and routing boundary", title_font, body_font)
    draw_box(draw, (890, 145, 1190, 345), "Ingestion Service", "Validates JSON request and publishes energy-usage", title_font, body_font)
    draw_box(draw, (890, 515, 1190, 715), "User / Device Services", "Manage household and appliance metadata", title_font, body_font, fill="F7FBFE")
    draw_box(draw, (1300, 145, 1600, 345), "Kafka", "Topics: energy-usage and energy-alerts", title_font, body_font, fill="F5FAFD")
    draw_box(draw, (1300, 515, 1600, 735), "Usage Service", "Stores time-series points and aggregates device and household usage", title_font, body_font)
    draw_box(draw, (1710, 145, 2010, 345), "InfluxDB", "energy_usage measurement with deviceId tag", title_font, body_font, fill="F7FBFE")
    draw_box(draw, (1710, 515, 2010, 735), "Insight Service", "Forecast, device ranking, standby heuristic, optional Ollama wording", title_font, body_font, fill="F7FBFE")
    draw_box(draw, (1300, 875, 1600, 1075), "Alert Service", "Persists breach alert and dispatches notification email", title_font, body_font, fill="F7FBFE")
    draw_box(draw, (1710, 875, 2010, 1075), "MySQL and Mailpit", "Relational metadata, alert history, and local email capture", title_font, body_font, fill="F7FBFE")

    draw_arrow(draw, (390, 245), (490, 355), "JSON REST", label_font)
    draw_arrow(draw, (390, 615), (490, 455), "API calls", label_font)
    draw_arrow(draw, (780, 355), (890, 245), "route", label_font)
    draw_arrow(draw, (780, 455), (890, 615), "route", label_font)
    draw_arrow(draw, (1190, 245), (1300, 245), "energy-usage", label_font)
    draw_arrow(draw, (1450, 345), (1450, 515), "consume", label_font)
    draw_arrow(draw, (1600, 575), (1710, 270), "persist", label_font)
    draw_arrow(draw, (1190, 615), (1300, 615), "ownership lookup", label_font)
    draw_arrow(draw, (1600, 625), (1710, 625), "usage query", label_font)
    draw_arrow(draw, (1450, 735), (1450, 875), "energy-alerts", label_font)
    draw_arrow(draw, (1600, 975), (1710, 975), "store / email", label_font)

    note = "Solid arrows denote implemented data or control flow. Optional Ollama output supplements textual recommendations only; numerical analytics remain deterministic."
    draw.text((80, 1110), note, font=label_font, fill="#404040")
    image.save(path, dpi=(300, 300))


def make_workflow_figure(path):
    image = Image.new("RGB", (2200, 860), "white")
    draw = ImageDraw.Draw(image)
    title_font = get_font(17, True)
    body_font = get_font(16, False)
    label_font = get_font(16, False)
    stage_font = get_font(18, True)

    labels = [
        (60, "Input"), (315, "Preprocessing"), (570, "Event Processing"), (825, "Data Storage"),
        (1080, "Aggregation"), (1335, "Analysis and Alert Rule"), (1590, "Evaluation"), (1845, "Output"),
    ]
    for x, label in labels:
        draw.text((x, 58), label, font=stage_font, fill="#1F4E79")

    boxes = [
        (45, 160, 285, 405, "Telemetry Input", "Wokwi ESP32 generates a JSON reading every 8 s"),
        (310, 160, 550, 405, "Input Validation", "Require deviceId and timestamp; energyConsumed is non-negative"),
        (575, 160, 815, 405, "Event Publication", "Map reading to energy-usage event and publish to Kafka"),
        (840, 160, 1080, 405, "Time-Series Storage", "Usage Service writes energy_usage point to InfluxDB"),
        (1105, 160, 1345, 405, "Usage Aggregation", "Sum device and household usage within selected windows"),
        (1370, 160, 1610, 405, "Analytics and Threshold Rule", "Forecast, ranking, standby heuristic, and alert-breach evaluation"),
        (1635, 160, 1875, 405, "Functional Evaluation", "Record API, Kafka, storage, insight, and alert evidence"),
        (1900, 160, 2140, 405, "User Output", "Dashboard insight, alert history, and Mailpit email on breach"),
    ]
    for x1, y1, x2, y2, title, detail in boxes:
        draw_box(draw, (x1, y1, x2, y2), title, detail, title_font, body_font, fill="F7FBFE" if x1 >= 1370 else BLUE_FILL)
    for i in range(len(boxes) - 1):
        x1 = boxes[i][2]
        x2 = boxes[i + 1][0]
        draw_arrow(draw, (x1, 282), (x2, 282), None, label_font)
    note = "No supervised learning, train-test split, or empirical anomaly classifier is implemented in the current prototype."
    draw.text((45, 750), note, font=label_font, fill="#404040")
    image.save(path, dpi=(300, 300))


def add_algorithm(doc, title, input_text, output_text, steps):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 10.5, bold=True)
    for label, value in (("Input", input_text), ("Output", output_text)):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(1)
        r1 = p.add_run(label + ": ")
        set_run_font(r1, 9.5, bold=True, name="Courier New")
        r2 = p.add_run(value)
        set_run_font(r2, 9.5, name="Courier New")
    for idx, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{idx}. {step}")
        set_run_font(r, 9.5, name="Courier New")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    for name, size in (("Heading 1", 12), ("Heading 2", 11), ("Heading 3", 10.5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(TEXT)
    return doc


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    architecture = FIGURE_DIR / "fig_1_proposed_system_architecture.png"
    workflow = FIGURE_DIR / "fig_2_proposed_workflow.png"
    make_architecture_figure(architecture)
    make_workflow_figure(workflow)

    doc = setup_document()

    # Section 3
    add_heading(doc, "3. PROPOSED METHODOLOGY AND PROPOSED SYSTEM", 1)
    add_heading(doc, "3.1 Overall Methodology", 2)
    add_body(doc, "The proposed work is an event-driven software-system prototype for household energy monitoring. It is designed primarily to address the integration gap identified in Section 2: the need to combine device telemetry, scalable service separation, time-series storage, deterministic energy analytics, and threshold-based notification within a single residential energy-management platform. The present implementation is not a supervised machine-learning experiment. Its current energy input is a controlled Wokwi ESP32 simulation, and its numerical outputs are generated by deterministic calculations.")
    add_body(doc, "The input to the system is an energy-telemetry record containing deviceId, energyConsumed, and timestamp. In the current prototype, the Wokwi firmware produces this JSON record every 8 s. The virtual potentiometer and three modes generate controlled scenarios: Eco (approximately 1-6 kWh), Heavy (approximately 8-20 kWh), and Alert Test (approximately 26-40 kWh). These values are scenario inputs for integration testing and must not be interpreted as calibrated physical meter measurements.")
    add_body(doc, "Data acquisition begins when the virtual ESP32 sends a REST request through the API Gateway to the Ingestion Service. The service validates that deviceId and timestamp are present and that energyConsumed is non-negative. It then transforms the request into an energy-usage event and publishes that event to Apache Kafka. The Usage Service consumes the event, stores deviceId as a tag and energyConsumed as a field in the InfluxDB energy_usage measurement, and preserves the event timestamp for time-window queries.")
    add_body(doc, "Feature transformation is performed through aggregation rather than feature learning. For a selected user and observation window of D days, the Usage Service sums energyConsumed for each registered device. The Insight Service derives household total energy, daily consumption rate, projected 30-day energy, slab-based cost, device contribution percentage, estimated device cost, and a standby-load estimate. The Alert path separately aggregates the preceding one-hour usage window every 10 s, maps device readings to their users, and compares each household total with its configured energy threshold.")
    add_body(doc, "The outputs are (i) device-level and household-level usage summaries, (ii) a bill forecast, (iii) ranked appliances, (iv) a standby-risk estimate with recommendations, and (v) persisted alert records with local email capture through Mailpit. Evaluation at the current review stage focuses on functional correctness and end-to-end traceability: accepted ingestion request, Kafka event propagation, InfluxDB persistence, insight retrieval, and alert delivery. One controlled functional run is reported in Section 5.1. Quantitative performance, forecast-error, and classification metrics still require larger experiments and ground-truth data.")

    add_heading(doc, "3.2 Proposed System Architecture", 2)
    add_body(doc, "Fig. 1 presents the proposed architecture. The architecture separates input, event processing, storage, analytics, and presentation responsibilities. The separation allows the telemetry-facing ingestion path to acknowledge valid readings without waiting for aggregation, analytics, or notification delivery.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_figure(
        p,
        architecture,
        Inches(6.55),
        "Block diagram of the proposed event-driven Home Energy Tracker architecture.",
    )
    add_caption(doc, "Fig. 1. Proposed system architecture of the Home Energy Tracker.")
    add_body(doc, "The virtual smart-meter block represents the current Wokwi ESP32 implementation. It supplies the same REST payload shape expected from a future physical meter but does not replace calibrated sensing hardware. The API Gateway is the public routing boundary for the browser and virtual meter. The Ingestion Service validates the payload and emits the energy-usage event. Kafka decouples this producer from the Usage Service, which persists time-series points to InfluxDB.")
    add_body(doc, "User and Device Services maintain household and appliance metadata in MySQL. During aggregation, the Usage Service retrieves device ownership and user alert settings. The Insight Service queries aggregated usage data and produces deterministic forecasts, rankings, and standby estimates. The Alert Service consumes energy-alerts events, persists alert history in MySQL, and sends a notification through the configured SMTP endpoint; Mailpit is used as a local demonstration inbox. Ollama, when reachable, is used only to express recommendations in natural language. A deterministic fallback is returned if the optional AI service is unavailable.")

    add_heading(doc, "3.3 Detailed Proposed Approach", 2)
    add_body(doc, "The proposed approach is executed as the following connected stages. Each stage produces an output consumed by the next stage, thereby making the source of dashboard values and alerts traceable.")
    add_table(doc, "Table 1. Detailed proposed approach and component interfaces.",
              ["Component", "Purpose", "Input", "Processing technique", "Output and next stage"],
              [
                  ["Profile and device registration", "Associate a household with its appliances and alert setting.", "User details, device name, type, location, and threshold.", "REST CRUD operations stored in MySQL.", "User and device identifiers used by usage aggregation."],
                  ["Wokwi telemetry acquisition", "Generate controlled end-to-end test telemetry.", "Potentiometer and selected Eco, Heavy, or Alert Test mode.", "ESP32 firmware maps analog mode state to a kWh value and creates a timestamped JSON request every 8 s.", "{deviceId, energyConsumed, timestamp} sent to API Gateway."],
                  ["Ingestion", "Validate and decouple device-facing requests.", "Telemetry JSON payload.", "Bean validation: non-null deviceId/timestamp and non-negative energyConsumed; event mapping and Kafka publication.", "energy-usage event for Usage Service."],
                  ["Usage storage and aggregation", "Persist readings and derive user/device window sums.", "Kafka event plus user/device metadata.", "Write InfluxDB point; Flux time-window sum grouped by deviceId; ownership resolution through Device Service.", "Per-device energy totals and household totals."],
                  ["Insight computation", "Convert usage totals into interpretable decision support.", "Aggregated usage over selected days and configured tariff values.", "Daily-rate forecast, slab-cost calculation, descending device sort, 20% average-load standby heuristic.", "Forecast, appliance ranking, standby risk, recommendations."],
                  ["Threshold alerting", "Notify a household when preceding one-hour summed usage is over its setting.", "One-hour household aggregate, threshold, and alerting flag.", "Scheduled rule evaluation every 10 s and energy-alerts publication on breach.", "Persisted alert and local email notification."],
                  ["Evaluation", "Demonstrate traceability and quantify performance when experimental data are collected.", "HTTP logs, Kafka records, InfluxDB query results, alert history, timing data.", "Functional checks and [TO BE PROVIDED] metric calculation.", "Evidence table and results/discussion inputs."],
              ], widths=[1.15, 1.37, 1.1, 1.55, 1.38])

    add_heading(doc, "3.4 Algorithms and Techniques", 2)
    add_body(doc, "The implemented system uses deterministic rules and mathematical transformations rather than a trained prediction or anomaly-detection model. Therefore, accuracy, F1-score, and train-test split claims are not appropriate for the current prototype. The following techniques are implemented or directly represented in the source code.")

    add_heading(doc, "3.4.1 Event-driven ingestion and time-window aggregation", 3)
    add_body(doc, "Purpose: accept a validated telemetry request and make it asynchronously available to downstream processing. Input: deviceId, energyConsumed, and timestamp. Working principle: the Ingestion Service maps the valid request to an energy-usage Kafka event. The Usage Service consumes the event and writes an InfluxDB point with deviceId as a tag, energyConsumed as a field, and the supplied timestamp. For retrieval, a Flux query restricts the time range, filters the energy_usage measurement and energyConsumed field, groups by deviceId, and sums the values. Output: per-device energy totals for a selected user and time range. This approach is appropriate because telemetry storage and aggregation can be performed independently of client request acknowledgement.")

    add_heading(doc, "3.4.2 Slab-based bill forecast", 3)
    add_body(doc, "Purpose: estimate a 30-day energy quantity and cost from observed usage. Let E_tot be the summed energy over D selected days. The system computes the average daily rate and 30-day projection as follows:")
    add_equation(doc, "E_day = E_tot / D,     E_30 = 30 x E_day")
    add_body(doc, "The projected cost is calculated using configured tariff slabs. For slab limits L1 and L2, slab rates r1, r2, r3, and fixed charge F, the implementation allocates consumption sequentially to the slab ranges and adds F only to the projected monthly cost. The local prototype defaults are F = 10.00, L1 = 100 kWh, r1 = 0.12, L2 = 300 kWh, r2 = 0.18, and r3 = 0.25. These are configuration defaults rather than a validated tariff for a particular utility and should be replaced by [TO BE PROVIDED: applicable tariff source] before a real-billing claim is made.")
    add_body(doc, "Output: current-period estimated cost, daily rate, projected 30-day kWh, and projected 30-day cost. The method is appropriate for a transparent prototype because every estimate can be reproduced from the aggregated input and tariff configuration. It is not a substitute for a utility-bill forecast validated against real billing data.")

    add_heading(doc, "3.4.3 Appliance ranking", 3)
    add_body(doc, "Purpose: identify devices that contribute the largest proportion of observed consumption. Input: device-level sums for a selected time window. The devices are sorted in descending order of energyConsumed. For device i, contribution is calculated as p_i = (E_i / E_tot) x 100 when E_tot is positive. Estimated device cost is allocated proportionally as c_i = p_i x C_tot / 100, where C_tot is the current-period slab-based cost. Output: a ranked device list with kWh, contribution percentage, and estimated cost. This method is appropriate because the application already records device-associated telemetry and does not claim appliance disaggregation from a single aggregate meter.")

    add_heading(doc, "3.4.4 Standby Energy Detective heuristic", 3)
    add_body(doc, "Purpose: translate continuous low-level demand into an explainable estimate of potential standby cost. The current implementation does not detect night-time or low-activity periods directly. Instead, it estimates baseline power as 20% of the average load across the selected observation period. If the selected window has D days and total energy E_tot, then:")
    add_equation(doc, "P_avg = (E_tot / (24D)) x 1000 W,     P_base = 0.20 x P_avg")
    add_equation(doc, "E_standby_day = 24 x P_base / 1000,     E_standby_month = 30 x E_standby_day")
    add_body(doc, "The system labels standby risk as LOW when P_base is at most 50 W, MEDIUM when it is greater than 50 W and at most 150 W, and HIGH when it is greater than 150 W. It calculates a monthly standby cost through the same tariff function and returns generic probable device categories and practical recommendations. Output: baseline watts, daily/monthly standby kWh, monthly estimated cost, risk level, and recommendations. This heuristic is appropriate for an explainable prototype, but it is not a ground-truth appliance-level standby detector.")

    add_heading(doc, "3.4.5 Threshold rule for alerts", 3)
    add_body(doc, "Purpose: generate a notification when aggregated household usage exceeds the user-defined limit. At each scheduled evaluation, the system sums the preceding one-hour device readings associated with a user. If alerting is enabled and the sum is greater than threshold theta_u, an energy-alerts event is published:")
    add_equation(doc, "Alert(u) = 1 if alerting_u = true and sum(E_i over preceding hour for user u) > theta_u; otherwise Alert(u) = 0")
    add_body(doc, "The scheduler runs every 10 s. Output: an alert event consumed by the Alert Service, which persists an alert record and attempts notification delivery. The technique is appropriate for transparent threshold monitoring; however, the present implementation does not deduplicate repeated breaches, so this limitation must be retained in the experimental discussion.")

    add_algorithm(doc, "Algorithm 1: Event ingestion, aggregation, and threshold evaluation",
                  "Telemetry record r = {deviceId, energyConsumed, timestamp}; registered devices; user thresholds.",
                  "Stored usage point; user-level aggregates; zero or more alert events.",
                  [
                      "Receive r through the ingestion endpoint.",
                      "If deviceId or timestamp is null, or energyConsumed is negative, reject r.",
                      "Map r to an energy-usage event and publish it to Kafka.",
                      "For each consumed event, write deviceId, energyConsumed, and timestamp to InfluxDB.",
                      "Every 10 s, query the preceding one-hour energy_usage values grouped by deviceId.",
                      "Resolve each deviceId to its userId and sum device totals for each user.",
                      "For each user with alerting enabled, compare the one-hour total with the configured threshold.",
                      "If the total exceeds the threshold, publish an energy-alerts event.",
                      "Return control to asynchronous consumers; dashboard queries retrieve stored aggregates separately.",
                  ])

    add_algorithm(doc, "Algorithm 2: Deterministic energy insight generation",
                  "Selected user u; observation length D; per-device energy totals; tariff configuration.",
                  "Bill forecast, appliance ranking, standby estimate, and recommendation text.",
                  [
                      "Retrieve per-device totals for user u over D days.",
                      "Compute E_tot by summing the device totals.",
                      "Compute E_day = E_tot / max(D, 1) and E_30 = 30 x E_day.",
                      "Apply the configured slab function to calculate current and projected costs.",
                      "Sort devices by descending energyConsumed and calculate each contribution percentage.",
                      "Compute P_base = 0.20 x (E_tot / (24D)) x 1000.",
                      "Derive standby daily/monthly kWh, cost, and LOW/MEDIUM/HIGH risk from P_base.",
                      "If Ollama is available, request concise wording using deterministic outputs; otherwise create fallback wording.",
                      "Return numeric outputs and the associated explanation.",
                  ])

    add_heading(doc, "3.5 Workflow of the Proposed System", 2)
    add_body(doc, "Fig. 2 illustrates the complete operational workflow. It starts with a generated telemetry input and ends with functional evaluation evidence and user-facing output. The flow makes clear that event processing, analytics, and alerting are distinct steps, and that no supervised learning stage is present in the current implementation.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_figure(
        p,
        workflow,
        Inches(6.55),
        "Workflow from Wokwi telemetry input through validation, storage, analysis, evaluation, and user output.",
    )
    add_caption(doc, "Fig. 2. Workflow of the proposed Home Energy Tracker.")
    add_body(doc, "After validation, each reading is made durable in the time-series layer through the Kafka-to-Usage-Service path. The stored data are queried over an explicit time range before deterministic analytics are calculated. Threshold evaluation produces an additional branch only when the household total exceeds its configured value. The evaluation stage does not assert performance; it records the evidence required to evaluate the implementation once controlled runs are conducted.")

    # Section 4
    add_heading(doc, "4. IMPLEMENTATION AND EXPERIMENTAL SETUP", 1)
    add_heading(doc, "4.1 Hardware Requirements", 2)
    add_body(doc, "The current implementation is simulation-based. It does not require a physical smart meter, current transformer, voltage sensor, or PZEM module for the Review-2 experiment. The Wokwi ESP32 environment is used to validate the firmware and API contract. Host-computer specifications must be reported from the machine used for final experiments.")
    add_table(doc, "Table 2. Hardware requirements for the current prototype.",
              ["Component", "Specification", "Purpose"],
              [
                  ["Development computer", "[TO BE PROVIDED: processor, RAM, operating system, and storage]", "Runs Docker containers, browser dashboard, VS Code, and local verification tools."],
                  ["Virtual ESP32 DevKit", "ESP32 DevKit simulated in Wokwi", "Executes Arduino firmware and submits test telemetry."],
                  ["Virtual potentiometer and buttons", "One potentiometer and three pushbuttons in Wokwi", "Select controlled Eco, Heavy, and Alert Test telemetry scenarios."],
                  ["Virtual OLED", "SSD1306 OLED simulated through I2C", "Displays mode, simulated kWh, Wi-Fi state, and HTTP response status."],
                  ["Physical energy sensor", "Not used in the current prototype", "[TO BE PROVIDED only if a calibrated meter is added for final evaluation.]"],
              ], widths=[1.5, 2.25, 2.8])

    add_heading(doc, "4.2 Software Requirements", 2)
    add_table(doc, "Table 3. Software platforms and versions evidenced by the current repository.",
              ["Software or platform", "Version", "Purpose"],
              [
                  ["Java", "17", "Builds and runs the Spring Boot services."],
                  ["Spring Boot", "3.5.3", "Implements the Insight Service; [TO BE PROVIDED: confirm uniform version across all services]."],
                  ["React", "19.2.8", "Implements the web dashboard."],
                  ["Vite", "8.2.2", "Frontend development and build tool."],
                  ["TypeScript", "6.0.2", "Frontend language and type checking."],
                  ["Tailwind CSS", "4.3.3", "Dashboard styling."],
                  ["Apache Kafka", "Container image tag latest in compose file", "Asynchronous event broker; exact run-time version [TO BE PROVIDED]."],
                  ["MySQL", "8.3.0", "Stores user, device, and alert records."],
                  ["InfluxDB", "2.7", "Stores timestamped energy usage."],
                  ["Docker and Docker Compose", "[TO BE PROVIDED: installed version]", "Runs the local multi-service environment."],
                  ["PlatformIO and Wokwi", "[TO BE PROVIDED: installed version]", "Build and simulate the ESP32 firmware."],
                  ["Ollama and Spring AI", "Spring AI 1.0.3; Ollama version [TO BE PROVIDED]", "Optional natural-language savings-tip generation with fallback."],
                  ["Mailpit", "Container image tag latest in compose file", "Captures demonstration emails locally."],
              ], widths=[1.75, 1.7, 3.1])

    add_heading(doc, "4.3 Dataset Details", 2)
    add_body(doc, "No external benchmark dataset or labelled training corpus is used by the current prototype. The operative data source is a synthetic telemetry stream produced by the Wokwi ESP32 firmware. Each observation is a JSON event rather than a row from a pre-existing dataset. Consequently, a target variable and train/validation/test split are not applicable unless a future supervised anomaly-detection study is added.")
    add_body(doc, "The firmware submits a record at an 8 s interval and generates controlled energyConsumed values according to the chosen mode. The raw payload is validated by the Ingestion Service and is transformed into a Kafka event. The Usage Service stores the time-series representation in InfluxDB. The total number of generated records, the dates of collection, retained duration, and selected scenarios for the experiment must be reported after execution.")
    add_table(doc, "Table 4. Dataset summary for the current Review-2 prototype.",
              ["Attribute", "Current detail"],
              [
                  ["Dataset name", "Synthetic Wokwi energy-telemetry stream (prototype)."],
                  ["Source", "Wokwi ESP32 firmware in iot-simulator/src/main.cpp."],
                  ["Observation unit", "One JSON telemetry event."],
                  ["Number of records", "[TO BE PROVIDED: total valid, invalid, and retained observations]."],
                  ["Features", "deviceId, energyConsumed (kWh), timestamp."],
                  ["Target variable", "Not applicable; no supervised learning model is implemented."],
                  ["Temporal characteristic", "Generated approximately every 8 s; actual observation duration [TO BE PROVIDED]."],
                  ["Scenario ranges", "Eco approximately 1-6 kWh; Heavy approximately 8-20 kWh; Alert Test approximately 26-40 kWh."],
                  ["Format and storage", "REST JSON -> Kafka event -> InfluxDB energy_usage measurement."],
                  ["Preprocessing", "Schema validation, event mapping, deviceId tagging, and time-window aggregation."],
                  ["Train/test/validation split", "Not applicable in the current deterministic prototype."],
                  ["Ground-truth labels", "Not available; [TO BE PROVIDED] before anomaly-detection accuracy can be evaluated."],
              ], widths=[2.0, 4.55])

    add_heading(doc, "4.4 Tools Technologies and Frameworks", 2)
    add_table(doc, "Table 5. Tools, technologies, and frameworks used in the current implementation.",
              ["Category", "Technology or tool", "Purpose"],
              [
                  ["Frontend", "React, TypeScript, Vite, Tailwind CSS", "Provides household, device, reading-simulator, insight, and alert-history views."],
                  ["API layer", "Spring Cloud Gateway", "Routes /api/v1 requests to backend services."],
                  ["Backend", "Java, Spring Boot, Spring Web, Spring Data", "Implements User, Device, Ingestion, Usage, Alert, and Insight services."],
                  ["Messaging", "Apache Kafka", "Carries energy-usage and energy-alerts events."],
                  ["Relational data", "MySQL 8.3.0", "Stores user, device, and alert metadata."],
                  ["Time-series data", "InfluxDB 2.7 and Flux query language", "Stores and aggregates energy_usage points."],
                  ["IoT simulation", "ESP32 Arduino firmware, PlatformIO, Wokwi", "Generates controllable telemetry and displays request status."],
                  ["Recommendation wording", "Spring AI with optional Ollama model", "Generates optional narrative tips; deterministic fallback preserves core behavior."],
                  ["Notification test surface", "Mailpit", "Captures SMTP alert messages locally."],
                  ["Containerization", "Docker and Docker Compose", "Runs the local service topology."],
              ], widths=[1.35, 2.25, 3.0])

    add_heading(doc, "4.5 Implementation Procedure", 2)
    add_numbered(doc, "Configure the local Docker environment with MySQL, Kafka, InfluxDB, Mailpit, the six application services, API Gateway, and the frontend. Record the exact Docker and host-machine versions in the final experiment log.")
    add_numbered(doc, "Create a household profile and register an appliance through the web dashboard or the corresponding gateway API. The resulting device identifier is used to associate telemetry with a household.")
    add_numbered(doc, "Configure the local Wokwi simulator with the assigned device identifier and the local gateway route host.wokwi.internal:8080/api/v1/ingestion. Build the ESP32 firmware with PlatformIO.")
    add_numbered(doc, "Generate controlled telemetry by selecting Eco, Heavy, or Alert Test mode. Capture the emitted JSON payload and HTTP status. A successful request returns HTTP 201 from the ingestion endpoint.")
    add_numbered(doc, "Validate the pipeline by observing an energy-usage event in Kafka and querying the resulting usage through the Gateway. The Usage Service stores a time-stamped InfluxDB point and aggregates device totals for requested time windows.")
    add_numbered(doc, "Retrieve bill forecast, appliance ranking, and standby-detective outputs through the Insight Service. Record the selected days parameter, tariff configuration, input window, and returned numerical values for later reproducibility.")
    add_numbered(doc, "For an alert scenario, use an isolated household with alerting enabled and a documented threshold. Submit one Alert Test reading, wait for the scheduled one-hour aggregation, then verify the alert record and Mailpit email. Disable alerting after evidence capture to prevent repeated alerts under the current non-deduplicated rule.")
    add_numbered(doc, "Execute the evaluation protocol in Section 4.6 and populate the result tables in Section 5 using retained timestamps, event counts, query results, and screenshots. Do not substitute visual dashboard fallback data for backend evidence.")

    add_heading(doc, "4.6 Experimental Setup", 2)
    add_body(doc, "The experimental setup is a local, containerized integration environment. It is intended to assess whether a controlled telemetry event can traverse the implemented pipeline and produce retrievable analytics and an alert when the threshold rule is satisfied. The setup is not a field deployment, hardware-calibration study, or comparative machine-learning experiment.")
    add_table(doc, "Table 6. Experimental configuration for the current prototype.",
              ["Configuration item", "Current configuration", "Reporting requirement"],
              [
                  ["Computing environment", "Local Docker-based environment", "[TO BE PROVIDED: CPU, RAM, OS, Docker version]."],
                  ["Telemetry source", "Wokwi ESP32 virtual meter", "Report number of runs, duration, modes, and selected potentiometer settings."],
                  ["Ingress interval", "8 s in firmware", "Confirm actual interval observed during each run."],
                  ["Observation window for insights", "API default 3 days unless days parameter is supplied", "Record selected D for each result."],
                  ["Alert aggregation window", "Preceding 1 hour", "Record threshold and exact reading time."],
                  ["Alert scheduler", "Every 10 s", "Record alert latency and duplicate-event behavior."],
                  ["Forecast horizon", "30 days", "Use the configured tariff values and report their source."],
                  ["Standby coefficient", "20% of average load", "State that this is a heuristic, not ground truth."],
                  ["Standby risk thresholds", "LOW <= 50 W; MEDIUM > 50 W and <= 150 W; HIGH > 150 W", "Retain these thresholds during reproducibility runs."],
                  ["Baseline methods", "No implemented empirical baseline", "[TO BE PROVIDED: choose baseline before comparative claims]."],
                  ["Evaluation evidence", "HTTP status, Kafka record, InfluxDB/usage query, insight response, alert history, Mailpit email", "Attach timestamps and screenshots to final result set."],
              ], widths=[1.65, 2.35, 2.6])

    # Section 5
    add_heading(doc, "5. RESULTS AND DISCUSSION", 1)
    add_heading(doc, "5.1 Experimental Results", 2)
    add_body(doc, "Two evidence streams were obtained for the current prototype. First, Maven tests were run locally for the six stateful/analytical services. Second, a controlled local end-to-end run was executed on 30 September 2026 after the Docker Compose stack was healthy. The run used a fresh test user and device, three synthetic inputs (0.25, 0.50, and 0.75 kWh), and a three-day query window. The inputs passed through the same REST ingestion endpoint used by the Wokwi ESP32 simulator; they are synthetic test increments, not calibrated physical-meter measurements.")
    add_body(doc, "The controlled run identifier was review2-20260930-215503-375c1ea3. Its raw machine-readable evidence is retained in research-output/evidence/review2_e2e_results_20260930-215503.json. Results below are functional observations for one isolated test case. They establish that the implemented path executed correctly under this configuration; they do not establish general latency, accuracy, scalability, fault tolerance, or household-energy savings.")
    add_table(doc, "Table 7. Verified automated and controlled end-to-end results.",
              ["Evidence category", "Input or test condition", "Observed result", "Status"],
              [
                  ["Automated test suite", "19 tests discovered across user, device, ingestion, usage, alert, and insight services; 3 marked skipped.", "16 executed tests completed with zero failures and zero errors. The insight suite included 7 bill-calculation and 4 insight-service tests.", "Verified unit test."],
                  ["Gateway health", "GET /actuator/health through the API Gateway.", "Overall gateway status was UP.", "Verified runtime."],
                  ["Telemetry acceptance", "Device 305; 0.25, 0.50, and 0.75 kWh with ISO-8601 timestamps.", "All three POST /ingestion requests returned HTTP 201 (3/3 accepted).", "Verified runtime."],
                  ["Kafka-to-Influx persistence", "Accepted test events for device 305.", "InfluxDB returned three energy_usage points: 0.25, 0.50, and 0.75 kWh. Usage API returned 1.50 kWh for the isolated device.", "Verified runtime."],
                  ["Bill forecast", "User 6; D = 3 days; total 1.50 kWh.", "Current cost $0.18; daily rate 0.50 kWh/day; projected monthly use 15.00 kWh; projected monthly cost $11.80.", "Verified runtime; matches formula."],
                  ["Appliance ranking", "One isolated device with 1.50 kWh.", "Device ranked first with 100.00% of total usage and estimated cost $0.18.", "Verified runtime."],
                  ["Standby heuristic", "User 6; D = 3 days; total 1.50 kWh.", "Baseline 4.17 W; daily 0.10 kWh; monthly 3.00 kWh; monthly cost $0.36; risk LOW.", "Verified runtime; heuristic output."],
                  ["Threshold alert and delivery", "User 7: threshold 0.10 kWh; one 1.00 kWh event.", "Condition 1.00 > 0.10 produced a detailed alert, a delivery-status record, and one Mailpit email. Alerting was then disabled.", "Verified runtime."],
                  ["Availability observation", "One controlled run with 3 s usage polling and 5 s alert polling.", "Usage and insight outputs were observed within 3.72 s; alert records within 5.23 s. These are polling-bounded observations, not benchmark latencies.", "Observed once; not a performance claim."],
                  ["Resilience or load behavior", "No load/fault-injection protocol.", "No throughput, resource, or recovery result recorded.", "[TO BE PROVIDED]."],
              ], widths=[1.35, 1.85, 2.3, 1.1])
    add_body(doc, "The direct InfluxDB query for device 305 returned records at 16:25:03.929Z, 16:25:04.032Z, and 16:25:04.078Z with values 0.25, 0.50, and 0.75, respectively. The Mailpit message had subject “Energy Usage Alert for User 7” and contained the threshold (0.1) and energy-consumed value (1.0). Ingestion-service logs recorded the corresponding correlation identifier and EnergyUsageEvent values. These artefacts establish traceability from HTTP acceptance to stored usage and alert delivery for this one run.")
    add_body(doc, "The current dashboard contains presentation-oriented fallback behavior; therefore backend responses, persistence queries, alert records, and Mailpit evidence should be treated as the primary result artefacts. Additional controlled runs, screenshots, and timestamp logs should be attached before final publication.")

    add_heading(doc, "5.2 Performance Metrics", 2)
    add_body(doc, "The metrics should reflect the implemented system rather than an unimplemented classifier. The following functional and performance metrics are appropriate. Classification accuracy, precision, recall, F1-score, and confusion matrices are only applicable if future work introduces labelled anomaly data and a classifier.")
    add_table(doc, "Table 8. Appropriate metrics and required data.",
              ["Metric", "Definition", "Why relevant", "Current evidence"],
              [
                  ["Ingestion acceptance rate", "100 x N_201 / N_submitted, where N_201 is the number of HTTP 201 responses.", "Measures whether valid telemetry is accepted by the API path.", "100% (3/3) in one controlled run; not a load-test result."],
                  ["Event persistence rate", "100 x N_stored / N_accepted.", "Measures end-to-end transfer from accepted event to queryable time-series point.", "100% (3/3) in one controlled run; evidenced by the direct InfluxDB query."],
                  ["End-to-end availability observation", "t_output_observed - t_first_submitted.", "Records time until the usage and insight endpoints expose a result.", "3.72 s in one run with 3 s polling; not a latency benchmark."],
                  ["Alert-observation time", "t_alert_observed - t_submitted.", "Records time from telemetry submission to a persisted alert record.", "5.23 s in one run with 5 s polling; not a latency benchmark."],
                  ["Throughput", "N_accepted / experiment duration.", "Measures the event rate handled during a defined test run.", "[TO BE PROVIDED]."],
                  ["Forecast MAE", "(1/n) sum |y_i - yhat_i| against actual bill/consumption reference.", "Required only when a real, matched reference series is available.", "Not computable without [TO BE PROVIDED: ground truth]."],
                  ["Standby precision/recall", "Compare predicted standby instances with labelled ground truth.", "Required only for a validated standby-detection claim.", "Not computable without [TO BE PROVIDED: labels]."],
                  ["Resource usage", "CPU, memory, and container utilization during test.", "Supports performance/scalability discussion.", "[TO BE PROVIDED]."],
              ], widths=[1.45, 2.05, 2.05, 1.45])
    add_equation(doc, "Acceptance rate = 100 x N_201 / N_submitted,     Persistence rate = 100 x N_stored / N_accepted")
    add_equation(doc, "MAE = (1 / n) x sum |y_i - yhat_i|   (only with a matched ground-truth reference)")

    add_heading(doc, "5.3 Comparison with Existing Methods", 2)
    add_body(doc, "No controlled baseline implementation or comparable experimental dataset is present. Consequently, a numerical comparison would be fabricated and is intentionally not provided. Before final submission, the authors should choose one or more baselines that use the same telemetry source, observation window, and evaluation protocol. Possible baselines must be defined by the actual future experiment, not inferred from the literature review.")
    add_table(doc, "Table 9. Comparison framework for future controlled evaluation.",
              ["Method", "Dataset or scenario", "Metric 1", "Metric 2", "Metric 3", "Computational cost"],
              [
                  ["Proposed Home Energy Tracker", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]"],
                  ["Baseline method A", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]"],
                  ["Baseline method B", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]", "[TO BE PROVIDED]"],
              ], widths=[1.6, 1.35, 0.85, 0.85, 0.85, 1.1])

    add_heading(doc, "5.4 Discussion", 2)
    add_body(doc, "Observed implementation evidence and experimental inference must be distinguished. The controlled run verified that three accepted synthetic telemetry events were available through the usage service, generated the expected deterministic insights, and triggered the threshold-based email path in the selected configuration. The source code shows that the prototype separates user/device metadata, telemetry ingestion, time-series storage, analytics, and alert notification into distinct services, connected through REST and Kafka. This result is functional evidence; it is not evidence of measured scalability, physical-meter accuracy, resilience, or household energy savings.")
    add_body(doc, "The architecture directly operationalizes the integration-oriented research gap by providing a traceable route from a telemetry event to time-series storage, analytics, and notification. Kafka reduces direct coupling between ingestion and downstream processing, while MySQL and InfluxDB separate relational metadata from time-series observations. These design decisions are technically appropriate for the stated system objective, but their performance benefits require load, fault-injection, and end-to-end latency experiments before they can be claimed empirically.")
    add_body(doc, "The Bill Shock Forecast and Standby Energy Detective provide explainable decision-support outputs. Their strength is transparency: the daily rate, 30-day projection, slab tariff, percentage contribution, and standby coefficient are all explicit. Their limitation is that the virtual telemetry is synthetic and the standby estimate is a fixed 20% average-load heuristic. Therefore, the current prototype can demonstrate system integration and rule execution but cannot substantiate a claim of calibrated energy measurement, exact appliance-level standby attribution, or forecast accuracy.")
    add_body(doc, "The optional Ollama integration should also be interpreted carefully. It provides narrative wording for recommendations when available and falls back to deterministic text when unavailable. It does not calculate bill values, identify anomalies, or serve as a validated predictive model. This fallback behavior improves functional robustness of the user interface but does not establish AI-model performance.")

    # Section 6
    add_heading(doc, "6. CONCLUSION", 1)
    add_heading(doc, "6.1 Summary of Proposed Work", 2)
    add_body(doc, "This work proposes a microservices-based Home Energy Tracker that accepts household device telemetry, transmits it through an event-driven pipeline, stores time-series observations, and returns interpretable energy-management outputs. The current implementation combines a Wokwi ESP32 virtual meter, API Gateway, Ingestion, Usage, Insight, and Alert services, Apache Kafka, InfluxDB, MySQL, a React dashboard, and Mailpit for local notification capture. The system produces deterministic bill forecasts, appliance rankings, standby-risk estimates, and threshold alerts.")

    add_heading(doc, "6.2 Key Findings and Expected Contribution", 2)
    add_body(doc, "At the present Review-2 stage, the contribution is an implemented and traceable system design with one controlled functional verification run, rather than a completed empirical performance study. In the recorded run, three synthetic events were accepted and retrieved as 1.50 kWh, and the forecast, ranking, standby heuristic, and threshold alert matched the specified deterministic behavior. Quantified findings on general latency, throughput, forecast error against a meter or utility bill, standby-detection accuracy, resilience, and comparative performance remain [TO BE PROVIDED] after larger controlled experimentation.")

    add_heading(doc, "6.3 Addressing the Research Gap", 2)
    add_table(doc, "Table 10. Current alignment between identified research gaps and the proposed work.",
              ["Research gap", "Proposed solution", "Evidence or result status", "Current contribution"],
              [
                  ["Limited integration of scalable AI-assisted event-driven services", "Gateway, Spring Boot services, Kafka, InfluxDB, MySQL, deterministic insights, and alert service in one prototype.", "Implementation architecture and source code are available; performance evidence [TO BE PROVIDED].", "Directly addressed at integration level."],
                  ["Lack of standardized datasets and metrics", "Results and metric framework in Sections 4-5.", "No benchmark dataset, labels, or numeric metrics currently available.", "Not resolved; evaluation plan is established."],
                  ["Limited context-aware analysis", "Current appliance metadata and windowed usage only.", "No occupancy, weather, schedule, or environmental features implemented.", "Not resolved; retained as limitation/future work."],
                  ["Insufficient resilience validation", "Service separation and asynchronous messaging provide a design basis.", "No fault-injection or load-test result available.", "Not empirically resolved."],
                  ["Privacy-preserving AI and secure access", "Optional local Ollama wording is present.", "Keycloak/JWT and ownership enforcement are not evidenced in the current codebase.", "Not resolved; security implementation is future work."],
              ], widths=[1.35, 2.1, 1.75, 1.35])

    add_heading(doc, "6.4 Limitations", 2)
    add_bullet(doc, "The current input is a Wokwi-generated synthetic telemetry stream; it is not a calibrated physical smart-meter dataset.")
    add_bullet(doc, "The reported energyConsumed values are scenario inputs sent at an 8 s interval. Their numeric cost outputs must not be interpreted as a real household bill without calibrated interval-energy measurements.")
    add_bullet(doc, "No external benchmark dataset, record count, labelled anomaly set, train/test split, or ground-truth reference is available for empirical evaluation.")
    add_bullet(doc, "The standby estimate is a fixed 20% average-load heuristic and does not identify actual low-activity periods or exact responsible appliances.")
    add_bullet(doc, "The alert scheduler can publish repeated breach events because the current implementation has no alert-deduplication window.")
    add_bullet(doc, "The project currently has no evidenced Keycloak/JWT authentication, device ownership enforcement, physical sensor calibration, cloud TLS verification, fault injection, or multi-node Kafka resilience validation.")
    add_bullet(doc, "The optional Ollama component affects recommendation wording only and has not been evaluated as an anomaly-detection or forecasting model.")
    add_bullet(doc, "The dashboard's illustrative client-side 7-day chart is not evidence of backend historical data and should not be used as a research result figure.")

    add_heading(doc, "6.5 Future Scope", 2)
    add_bullet(doc, "Integrate calibrated physical sensing hardware and convert sampled power into validated interval energy before ingestion.")
    add_bullet(doc, "Collect a sufficiently long, consented household telemetry dataset with record counts, sampling interval, appliance metadata, and ground-truth labels where applicable.")
    add_bullet(doc, "Validate the forecast against matched utility-bill or meter readings and report MAE/MAPE with documented tariff assumptions.")
    add_bullet(doc, "Replace the fixed standby coefficient with time-of-day, occupancy-aware, or labelled appliance-level analysis after appropriate data collection.")
    add_bullet(doc, "Implement secure authentication, authorization, device provisioning, and ownership checks before a multi-user deployment.")
    add_bullet(doc, "Add idempotency, retry/DLQ handling, alert deduplication, load testing, fault injection, and observability before claiming resilience or production readiness.")
    add_bullet(doc, "Evaluate an explicitly selected anomaly-detection model only after obtaining an appropriate labelled dataset and a reproducible baseline.")

    # Figure and table plan
    add_heading(doc, "APPENDIX A. REVIEW-2 FIGURE AND TABLE PLAN", 1)
    add_table(doc, "Table 11. Recommended figures and their placement in the Review-2 paper.",
              ["Figure", "Title", "Purpose and contents", "Placement"],
              [
                  ["Fig. 1", "Proposed system architecture", "Implemented input, event processing, data stores, analytics, alerting, and outputs.", "Section 3.2."],
                  ["Fig. 2", "Workflow of the proposed system", "Input -> validation -> Kafka -> storage -> analytics/alerting -> evaluation -> output.", "Section 3.5."],
                  ["Fig. 3", "Detailed telemetry evidence", "[TO BE PROVIDED: one annotated capture showing Wokwi JSON, HTTP 201, and matching Kafka/usage evidence].", "Section 5.1."],
                  ["Fig. 4", "Functional latency or throughput graph", "[TO BE PROVIDED: plot derived from timestamped controlled runs].", "Section 5.1 or 5.2."],
                  ["Fig. 5", "Forecast comparison graph", "[TO BE PROVIDED: projected versus measured energy/cost after ground-truth data collection].", "Section 5.1."],
                  ["Fig. 6", "Standby evaluation graph", "[TO BE PROVIDED: heuristic estimate versus labelled/validated reference, if available].", "Section 5.1."],
                  ["Fig. 7", "Confusion matrix", "Include only if a labelled anomaly classifier is implemented and evaluated.", "Section 5.1; otherwise omit."],
              ], widths=[0.7, 1.6, 3.2, 1.0])
    add_table(doc, "Table 12. Recommended tables and their placement in the Review-2 paper.",
              ["Table", "Title", "Placement"],
              [
                  ["Table 1", "Detailed proposed approach and component interfaces", "Section 3.3."],
                  ["Table 2", "Hardware requirements", "Section 4.1."],
                  ["Table 3", "Software requirements", "Section 4.2."],
                  ["Table 4", "Dataset summary", "Section 4.3."],
                  ["Table 5", "Tools, technologies, and frameworks", "Section 4.4."],
                  ["Table 6", "Experimental configuration", "Section 4.6."],
                  ["Table 7", "Preliminary functional-results framework", "Section 5.1."],
                  ["Table 8", "Performance metrics", "Section 5.2."],
                  ["Table 9", "Comparison framework", "Section 5.3."],
                  ["Table 10", "Research-gap alignment", "Section 6.3."],
              ], widths=[1.0, 3.8, 1.7])

    add_heading(doc, "APPENDIX B. REVIEW-2 COMPLETION CHECKLIST", 1)
    add_table(doc, "Table 13. Review-2 completion checklist.",
              ["Item", "Status", "Notes"],
              [
                  ["Methodology", "Completed", "Evidence-based Section 3.1."],
                  ["System architecture", "Completed", "Fig. 1 uses implemented components only."],
                  ["Methodology diagram", "Completed", "Fig. 1."],
                  ["Detailed proposed approach", "Completed", "Section 3.3 and Table 1."],
                  ["Algorithms and pseudocode", "Completed", "Section 3.4 and Algorithms 1-2."],
                  ["Workflow diagram", "Completed", "Fig. 2."],
                  ["Hardware requirements", "Completed", "Host specifications still [TO BE PROVIDED]."],
                  ["Software requirements", "Completed", "Run-time versions marked where unavailable."],
                  ["Dataset description", "Needs Information", "Record count, duration, and ground truth are [TO BE PROVIDED]."],
                  ["Tools and technologies", "Completed", "Table 5."],
                  ["Implementation procedure", "Completed", "Section 4.5."],
                  ["Experimental setup", "Needs Information", "Host details and final test protocol required."],
                  ["Results", "Completed (functional scope)", "One controlled end-to-end run and automated test outcomes are reported; empirical performance claims remain pending."],
                  ["Performance metrics", "Completed", "Framework in Section 5.2; values pending."],
                  ["Comparison table", "Needs Information", "Baseline methods and values pending."],
                  ["Graphs and result figures", "Needs Information", "Fig. 3 onward require actual experiment data."],
                  ["Discussion", "Completed", "Separates implementation facts, inference, and expected benefit."],
                  ["Conclusion", "Completed", "Does not claim unmeasured performance."],
                  ["Limitations", "Completed", "Matches the implemented scope."],
                  ["Future scope", "Completed", "Derived from current limitations."],
                  ["IEEE references", "Needs Review", "Existing references retained in research.docx; cross-check citations and formatting before final submission."],
                  ["Figure captions", "Completed", "Figs. 1-2 included; future captions planned."],
                  ["Table captions", "Completed", "Tables 1-13 captioned."],
                  ["Overall consistency", "Completed for Sections 3-6", "The continuation limits claims to the implemented Wokwi/API/Kafka/InfluxDB/alert/deterministic-insight scope."],
              ], widths=[2.05, 1.25, 3.2])

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
