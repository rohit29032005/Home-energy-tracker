"""Build the merged Review-2 paper and its result-provenance companion document.

This generator intentionally distinguishes verified local functional evidence from
unmeasured performance or physical-meter claims.
"""

from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt


ROOT = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker")
TOOLS = ROOT / "tools"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

import build_review2_sections as review


ORIGINAL_PAPER = Path(r"C:\Users\anju2\Desktop\Advance_java_research\research.docx")
OUTPUT_DIR = ROOT / "research-output"
FINAL_PAPER = OUTPUT_DIR / "research_final_merged_review2.docx"
HOW_PROVEN = OUTPUT_DIR / "how_results_were_proven.docx"
RAW_EVIDENCE = "research-output/evidence/review2_e2e_results_20260930-215503.json"


def first_body_child(document):
    for child in document._element.body:
        if child.tag.endswith("}sectPr"):
            continue
        return child
    raise RuntimeError("Document body has no content anchor.")


def prepend_paragraph(document, anchor, text="", style=None, size=11, bold=False,
                      italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      after=5, before=0, first_indent=0):
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing = 1.15
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Cm(first_indent)
    run = paragraph.add_run(text)
    review.set_run_font(run, size=size, bold=bold, italic=italic)
    anchor.addprevious(paragraph._p)
    return paragraph


def prepend_heading(document, anchor, text, level=1):
    p = prepend_paragraph(
        document,
        anchor,
        text,
        style=f"Heading {level}",
        size=12 if level == 1 else 11 if level == 2 else 10.5,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        after=5,
        before=11 if level == 1 else 8,
    )
    p.paragraph_format.keep_with_next = True
    return p


def prepend_number(document, anchor, text):
    p = prepend_paragraph(
        document,
        anchor,
        text,
        style="List Number",
        size=10.5,
        after=3,
        first_indent=0,
    )
    p.paragraph_format.line_spacing = 1.1
    return p


def append_reference(document, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    run = p.add_run(text)
    review.set_run_font(run, size=9)
    return p


def source_texts():
    source = Document(ORIGINAL_PAPER)
    return [p.text.strip() for p in source.paragraphs]


def add_preamble_to_final(document):
    source = source_texts()
    anchor = first_body_child(document)

    # Title page
    title = prepend_paragraph(
        document,
        anchor,
        "A Microservices-Based Home Energy Tracker for Real-Time Monitoring, Intelligent Alerting, and AI-Driven Energy Insights",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
        before=6,
    )
    title.paragraph_format.keep_with_next = True
    prepend_paragraph(document, anchor, "Authors", size=12, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    for author in (
        "Rohit Shrivastava | Registration Number: 23BCE0052",
        "Anshika Rai | Registration Number: 23BCE0080",
        "Harsh Gupta | Registration Number: 23BCE0050",
        "Amisha Raj | Registration Number: 23BCE0001",
    ):
        prepend_paragraph(document, anchor, author, size=10.5,
                          align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    prepend_paragraph(document, anchor,
                      "Department of Computer Science and Engineering, School of Computer Science and Engineering\nVellore Institute of Technology, Vellore, Tamil Nadu, India",
                      size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    prepend_paragraph(document, anchor, "Guide", size=12, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    guide = prepend_paragraph(document, anchor,
                              "Dr. Syamasudha Veeradandham\nAssistant Professor, School of Computer Science and Engineering\nVellore Institute of Technology",
                              size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    guide.add_run().add_break(WD_BREAK.PAGE)

    # Corrected front matter
    prepend_heading(document, anchor, "Abstract", 1)
    prepend_paragraph(document, anchor,
                      "Residential energy-monitoring platforms require a traceable path from device telemetry to actionable feedback, but existing solutions frequently address sensing, analytics, or scalable software deployment separately. This work presents a microservices-based Home Energy Tracker that integrates simulated device telemetry, event-driven processing, time-series storage, deterministic energy analytics, and threshold-based alerts. The implemented prototype consists of Spring Boot User, Device, Ingestion, Usage, Insight, and Alert services behind an API Gateway. A Wokwi ESP32 virtual meter produces timestamped JSON telemetry, which is validated by the Ingestion Service, published through Apache Kafka, and stored by the Usage Service in InfluxDB. MySQL maintains household and device metadata, while the Insight Service calculates a slab-based bill forecast, appliance ranking, and a standby-load heuristic. An optional Ollama component generates narrative wording only; numerical outputs remain deterministic and fall back safely when the AI component is unavailable. A controlled local functional run accepted three synthetic readings totaling 1.50 kWh, returned the expected forecast and standby values, and produced a threshold alert captured by Mailpit. These results demonstrate functional integration, not physical-meter accuracy, predictive-model accuracy, scalability, or resilience. The architecture provides a reproducible foundation for future calibrated sensing and empirical evaluation.")
    prepend_heading(document, anchor, "Keywords", 1)
    prepend_paragraph(document, anchor,
                      "Home energy monitoring; Internet of Things; microservices; Spring Boot; Apache Kafka; InfluxDB; deterministic energy analytics; Wokwi simulation; threshold alerting.")

    # Section 1
    prepend_heading(document, anchor, "1. INTRODUCTION", 1)
    intro = [
        "Residential electricity use is increasingly shaped by connected appliances and digital services. However, monthly bills alone provide limited visibility into when energy is consumed, which device contributes most, or whether a configured household threshold has been exceeded. A practical Home Energy Management System therefore needs an auditable route from telemetry to stored records, analytical output, and user notification.",
        "Existing research provides strong but often separate contributions in low-cost sensing, anomaly detection, time-series analysis, and distributed architectures. A software prototype that links these concerns must avoid treating simulated telemetry as calibrated measurement data or treating rule-based outputs as a trained artificial-intelligence model. This distinction is central to the experimental claims of this work.",
        "The proposed Home Energy Tracker implements independent Spring Boot services for household profiles, devices, telemetry ingestion, time-series usage, insights, and alerts. A Wokwi ESP32 simulator emits deviceId, energyConsumed, and timestamp records. The API Gateway routes the record to the Ingestion Service, Kafka decouples ingestion from downstream work, InfluxDB stores the time-series point, and the Insight Service returns deterministic forecast, ranking, and standby outputs. The Alert Service persists a threshold-breach alert and sends an email to a local Mailpit inbox.",
        "The study is scoped as a system-integration and functional-verification project. It reports a controlled local execution, source-backed formulas, and automated test outcomes. Physical meter calibration, authentication/authorization, anomaly classification, benchmark comparisons, resilience testing, and large-scale performance measurements remain future work rather than completed claims."
    ]
    for paragraph in intro:
        prepend_paragraph(document, anchor, paragraph, first_indent=0.63)

    # Section 2 objectives
    prepend_heading(document, anchor, "2. OBJECTIVES", 1)
    prepend_paragraph(document, anchor, "The objectives of the proposed Home Energy Tracker are:")
    objectives = [
        "Design a modular, event-driven microservices architecture for household energy monitoring using Spring Boot and an API Gateway.",
        "Accept timestamped device telemetry through a validated REST ingestion interface and transfer accepted events through Apache Kafka to a time-series store.",
        "Provide transparent, deterministic energy-management outputs: a slab-based bill forecast, appliance ranking, and standby-load heuristic.",
        "Implement a configurable threshold-alert path that persists alert history and demonstrates local email delivery through Mailpit.",
        "Validate the implemented path through automated tests and controlled end-to-end runs, while documenting the limits of simulation-based evidence.",
        "Maintain an extensible foundation for future calibrated sensing, secure access control, data-driven anomaly detection, and resilience evaluation."
    ]
    for item in objectives:
        prepend_number(document, anchor, item)

    # Literature review retained from the user-provided paper, with project-bridge statements corrected.
    prepend_heading(document, anchor, "2.1 LITERATURE REVIEW", 2)
    literature_heading_indices = {62, 67, 72, 77, 82, 87, 92, 97, 102, 107}
    replacements = {
        66: "This study motivates the use of event-driven telemetry and rich analysis. In the present prototype, Apache Kafka and InfluxDB support the integration path, while trained appliance classification, anomaly detection, and secure authentication are not implemented and remain future work.",
        71: "This hardware-oriented study motivates future physical sensing. The current prototype instead uses Wokwi-generated test input to exercise the backend interface; it does not claim PZEM-level measurement accuracy or calibrated electrical parameters.",
        76: "This work provides an architectural basis for service separation. The present implementation adopts Spring Boot microservices, Kafka, InfluxDB, deterministic insight functions, and a threshold alert path; secure authentication and trained AI analytics are outside the current scope.",
        81: "The layered conceptual model informs the separation of ingestion, operational services, and analytics in the proposed prototype. The implemented additions are API routing, Kafka event transfer, InfluxDB storage, deterministic insights, and local alert delivery; Keycloak authentication is reserved for future development.",
        86: "The study underlines the importance of accurate metering and power-quality sensing. It therefore motivates a future calibrated hardware phase, while the current Wokwi simulation is used only for software integration and functional testing.",
        91: "The anomaly-detection literature motivates future analytical extensions. In the current implementation, the Alert Service applies a configured threshold rule to aggregated usage rather than a hybrid anomaly classifier.",
        96: "This framework demonstrates the potential of AI-assisted anomaly detection at scale. The present work addresses service integration at prototype scale; it does not implement or evaluate a trained anomaly model, accuracy, precision, recall, or F1-score.",
        101: "The study supports the value of historical time-series data. The current prototype stores and queries historical usage but does not infer unusual daily behavior; its active alert condition is a configurable threshold.",
        106: "This review motivates future resilience experiments. The current service separation and asynchronous messaging supply a design basis, but no fault injection, chaos testing, or resilience result is claimed in this paper.",
        111: "This work informs the importance of monitoring distributed services. The present prototype applies microservice separation to residential energy telemetry and deterministic user-facing outputs; advanced anomaly analysis remains future work.",
    }
    for index in range(62, 112):
        text = source[index]
        if not text:
            continue
        if index in replacements:
            text = replacements[index]
        if index in literature_heading_indices:
            prepend_heading(document, anchor, text, 3)
        else:
            prepend_paragraph(document, anchor, text, first_indent=0.63)

    prepend_heading(document, anchor, "2.2 RESEARCH GAP", 2)
    for index in range(113, 160):
        text = source[index]
        if not text:
            continue
        if text.startswith("Gap "):
            prepend_heading(document, anchor, text, 3)
        else:
            prepend_paragraph(document, anchor, text, first_indent=0.63)
    prepend_paragraph(document, anchor,
                      "The present prototype directly addresses Gap 5 at the system-integration level by linking simulated telemetry, event-driven microservices, time-series storage, deterministic decision support, and alert delivery. Gaps concerning standardized datasets, contextual analysis, resilience validation, and privacy-preserving access control are not resolved; they are retained as explicit limitations and future work.",
                      first_indent=0.63)

    prepend_heading(document, anchor, "2.3 NOVELTY OF THE PROPOSED WORK", 2)
    prepend_paragraph(document, anchor,
                      "The novelty of the work is the implementation and traceable verification of a compact residential-energy prototype that combines a Wokwi telemetry interface, independently deployable Spring Boot services, Kafka event transfer, InfluxDB time-series storage, deterministic insight computation, and a threshold-alert path in one workflow. The contribution is not a new anomaly-detection algorithm, a calibrated smart meter, or a validated AI predictor. Instead, it provides an auditable integration baseline in which each numerical dashboard value can be linked to a defined formula and each alert can be linked to an explicit threshold condition.",
                      first_indent=0.63)

    prepend_heading(document, anchor, "2.4 NOVEL CONTRIBUTIONS", 2)
    contributions = [
        "An event-driven microservices architecture with User, Device, Ingestion, Usage, Insight, Alert, and API Gateway responsibilities.",
        "A Wokwi ESP32 simulation interface that sends the same timestamped JSON payload shape used by the ingestion API.",
        "Kafka-based decoupling of ingestion from time-series persistence, with InfluxDB used for device-tagged energy_usage records.",
        "Transparent deterministic calculations for bill forecasting, appliance ranking, and standby-load estimation.",
        "A scheduled threshold-alert flow that persists alert records and demonstrates local email delivery through Mailpit.",
        "A reproducible controlled verification protocol and evidence record for a 1.50 kWh analytics case and a 1.00 kWh threshold-breach case."
    ]
    for item in contributions:
        p = prepend_paragraph(document, anchor, item, style="List Bullet", size=10.5,
                              after=2, first_indent=0)
        p.paragraph_format.line_spacing = 1.05


def append_references(document):
    source = source_texts()
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    review.add_heading(document, "REFERENCES", 1)
    for text in source[185:]:
        if text:
            append_reference(document, text)


def build_final_paper():
    review.build_document()
    document = Document(review.OUTPUT_DOCX)
    add_preamble_to_final(document)
    append_references(document)
    document.save(FINAL_PAPER)


def how_heading(document, text, level=1):
    return review.add_heading(document, text, level)


def how_body(document, text, indent=0):
    return review.add_body(document, text, indent=indent)


def build_how_proven():
    document = review.setup_document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("How the Review-2 Results Were Proven")
    review.set_run_font(r, size=16, bold=True)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Home Energy Tracker | Controlled local functional-verification record")
    review.set_run_font(r, size=10.5, italic=True)
    sub.paragraph_format.space_after = Pt(12)

    how_heading(document, "1. PURPOSE AND EVIDENCE BOUNDARY", 1)
    how_body(document,
             "This document explains, step by step, how each reported Review-2 result was obtained and how it can be repeated. The recorded run used synthetic REST telemetry through the same endpoint used by the Wokwi ESP32 simulator. It proves functional flow through the implemented system; it does not prove physical-meter calibration, household-bill accuracy, AI-model accuracy, scalability, fault tolerance, or resilience.")

    how_heading(document, "2. REPRODUCIBILITY RECORD", 1)
    review.add_table(document, "Table 1. Evidence record for the controlled run.",
                     ["Item", "Recorded value"],
                     [
                         ["Run identifier", "review2-20260930-215503-375c1ea3"],
                         ["Execution date", "30 September 2026 (local Docker Compose run)"],
                         ["API base URL", "http://127.0.0.1:8080/api/v1"],
                         ["Gateway health", "UP"],
                         ["Raw result record", RAW_EVIDENCE],
                         ["Analytics test user/device", "User 6 / Device 305 (isolated synthetic test case)"],
                         ["Alert test user/device", "User 7 / Device 306 (isolated synthetic test case)"],
                         ["Data boundary", "Synthetic test increments; no calibrated physical meter or external benchmark dataset."],
                     ], widths=[1.65, 5.15])

    how_heading(document, "3. ENVIRONMENT PREPARATION", 1)
    steps = [
        "Start Docker Desktop and wait until its engine is available.",
        "From the project directory, run docker compose up -d and wait for MySQL, Kafka, InfluxDB, Mailpit, the API Gateway, and the six Java services to complete startup.",
        "Check GET http://127.0.0.1:8080/actuator/health. The recorded run returned status UP.",
        "Run Maven tests using the project-local Maven cache. The recorded execution discovered 19 tests across six services; 16 executed with zero failures and zero errors, while 3 were marked skipped.",
    ]
    for step in steps:
        review.add_numbered(document, step)
    how_body(document,
             "The implementation services used in the evidence chain were User, Device, Ingestion, Usage, Insight, and Alert. Kafka transferred ingestion events; InfluxDB stored time-series points; MySQL stored metadata and alerts; Mailpit captured local email output.")

    how_heading(document, "4. CONTROLLED ANALYTICS CASE", 1)
    how_body(document,
             "A fresh user was created with alerting disabled and a very high threshold to ensure that the analytics case could not create threshold alerts. A single GENERAL device (Device 305) was registered for this user. Three accepted ingestion requests were then sent.")
    review.add_table(document, "Table 2. Synthetic inputs accepted by the ingestion API.",
                     ["Device", "Timestamp (UTC)", "energyConsumed", "HTTP status"],
                     [
                         ["305", "2026-09-30T16:25:03.929Z", "0.25 kWh", "201 Created"],
                         ["305", "2026-09-30T16:25:04.032Z", "0.50 kWh", "201 Created"],
                         ["305", "2026-09-30T16:25:04.078Z", "0.75 kWh", "201 Created"],
                     ], widths=[0.75, 2.3, 1.6, 1.15])
    how_body(document,
             "The arithmetic input total was 0.25 + 0.50 + 0.75 = 1.50 kWh. The Ingestion Service logged all three EnergyUsageEvent values with the run correlation identifier. A direct InfluxDB query for the energy_usage measurement and deviceId 305 returned the same three persisted points. The Usage API then returned one device with energyConsumed = 1.50 kWh. Thus, the observed persistence total matched the submitted sum.")

    how_heading(document, "5. HOW THE ANALYTICS VALUES WERE CALCULATED", 1)
    how_body(document,
             "The Insight Service uses deterministic formulas. The selected observation window was D = 3 days. The configured tariff was a fixed monthly charge of $10.00, then 0–100 kWh at $0.12/kWh, 100–300 kWh at $0.18/kWh, and usage above 300 kWh at $0.25/kWh. Values are rounded to two decimal places using HALF_UP rounding.")
    review.add_table(document, "Table 3. Formula substitution and observed response for the analytics case.",
                     ["Output", "Formula and substitution", "Observed API value", "Match"],
                     [
                         ["Total usage", "T = 0.25 + 0.50 + 0.75", "1.50 kWh", "Yes"],
                         ["Current cost", "C(T) = 1.50 x $0.12", "$0.18", "Yes"],
                         ["Daily rate", "d = T / 3 = 1.50 / 3", "0.50 kWh/day", "Yes"],
                         ["30-day projection", "P = 30d = 30 x 0.50", "15.00 kWh", "Yes"],
                         ["Projected cost", "C(P) + fixed = (15 x $0.12) + $10", "$11.80", "Yes"],
                         ["Device ranking", "100 x device/T = 100 x 1.50/1.50", "100.00%", "Yes"],
                         ["Standby baseline", "round(0.20 x (T/(3 x 24)) x 1000)", "4.17 W", "Yes"],
                         ["Standby daily energy", "round(4.17 x 24 / 1000)", "0.10 kWh", "Yes"],
                         ["Standby monthly energy", "round(0.10 x 30)", "3.00 kWh", "Yes"],
                         ["Standby cost", "C(3.00) = 3.00 x $0.12", "$0.36", "Yes"],
                         ["Standby risk", "LOW when baseline <= 50 W", "LOW", "Yes"],
                     ], widths=[1.25, 3.15, 1.3, 0.75])
    how_body(document,
             "The Usage, bill-forecast, appliance-ranking, and standby-detective endpoint responses all matched the independent formula calculation. The recorded availability observation was 3.72 s from the first submission until the queried outputs were available. Because the script polled every 3 s, this must be reported as a polling-bounded functional observation, not a latency benchmark.")

    how_heading(document, "6. CONTROLLED THRESHOLD-ALERT CASE", 1)
    how_body(document,
             "A second isolated user (User 7) was created with alerting enabled and a threshold of 0.10 kWh. A second device (Device 306) then sent one 1.00 kWh synthetic event. The rule evaluated by the Usage Service was 1.00 > 0.10, which is true.")
    review.add_table(document, "Table 4. Threshold-alert proof chain.",
                     ["Stage", "Observed evidence"],
                     [
                         ["Input", "POST /ingestion for Device 306 with 1.00 kWh returned HTTP 201."],
                         ["Rule", "Aggregated consumption 1.00 kWh exceeded the configured 0.10 kWh threshold."],
                         ["Alert service", "A detailed alert record was persisted with message, threshold 0.1, energyConsumed 1.0, and sent=true."],
                         ["Delivery status", "A second status-only record was created by the alert-delivery flow; this is expected behavior, not a duplicate failure."],
                         ["Mailpit", "One local email was captured with subject ‘Energy Usage Alert for User 7’ and a snippet containing the threshold and energy value."],
                         ["Safety action", "Alerting was disabled through the User API after observation; the update returned HTTP 200."],
                     ], widths=[1.45, 5.0])
    how_body(document,
             "The alert records were observed 5.23 s after the test submission. The scheduler runs every 10 s and the evidence script polls every 5 s, so this observation must not be presented as a stable alert-latency guarantee. Its purpose is to prove that the threshold rule and delivery chain executed in the configured environment.")

    how_heading(document, "7. AUTOMATED TEST EVIDENCE", 1)
    review.add_table(document, "Table 5. Unit-test outcomes used as supplementary evidence.",
                     ["Suite", "Evidence verified", "Outcome"],
                     [
                         ["BillCalculationServiceTest", "Fixed charge, three tariff slabs, 30-day projection, and baseline-watt helper.", "7 tests passed."],
                         ["InsightServiceTest", "Forecast aggregation, descending appliance ranking, standby response presence, and deterministic fallback when Ollama fails.", "4 tests passed."],
                         ["Application-context tests", "User, Device, Ingestion, Usage, Alert, and Insight Spring contexts.", "8 discovered: 5 passed, 3 skipped, 0 failures/errors."],
                         ["All recorded service tests", "19 discovered across the six services.", "16 executed with 0 failures and 0 errors; 3 skipped."],
                     ], widths=[2.0, 3.25, 1.25])

    how_heading(document, "8. HOW TO REPEAT THE RUN", 1)
    how_body(document,
             "The project includes tools/run_review2_e2e_verification.ps1. From the project root, start the Compose stack and execute the script in PowerShell. It creates fresh test identities and writes a timestamped JSON record under research-output/evidence/. Compare the returned values against Table 3 and confirm that all passed fields are true. The script disables alerting after the alert observation to prevent repeated scheduler-driven notifications.")
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("docker compose up -d\n.\\tools\\run_review2_e2e_verification.ps1")
    review.set_run_font(r, size=9.5, name="Consolas")

    how_heading(document, "9. CLAIMS THAT THIS EVIDENCE DOES NOT SUPPORT", 1)
    limits = [
        "No calibrated sensor or physical utility-meter comparison was used; the values are synthetic test increments.",
        "No machine-learning model, labelled anomaly dataset, train/test split, accuracy, precision, recall, or F1-score was evaluated.",
        "The single-run availability and alert-observation times are not throughput, scalability, or latency benchmarks.",
        "No fault injection, service-recovery, multi-node Kafka, authentication, authorization, cloud TLS, or privacy study was performed.",
        "The standby value is a fixed 20% average-load heuristic, not verified appliance-level standby consumption.",
    ]
    for item in limits:
        review.add_bullet(document, item)

    document.save(HOW_PROVEN)


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_final_paper()
    build_how_proven()
    print(FINAL_PAPER)
    print(HOW_PROVEN)
