"""Create the requested single-column Review-2 manuscript from supplied evidence.

The source DOCX supplies every numerical result, author detail, reference, figure,
and original table.  This script does not touch the source file.
"""

from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def m(tag: str) -> OxmlElement:
    return OxmlElement(f"m:{tag}")


def math_run(text: str, roman: bool = False) -> OxmlElement:
    run = m("r")
    if roman:
        rpr = OxmlElement("m:rPr")
        rpr.append(OxmlElement("m:nor"))
        run.append(rpr)
    text_node = OxmlElement("m:t")
    text_node.text = text
    run.append(text_node)
    return run


def subscript(base: str, sub: str, sub_roman: bool = False) -> OxmlElement:
    node = m("sSub")
    e = m("e")
    e.append(math_run(base))
    sub_node = m("sub")
    sub_node.append(math_run(sub, roman=sub_roman))
    node.append(e)
    node.append(sub_node)
    return node


def fraction(children: list[OxmlElement], denominator: str) -> OxmlElement:
    node = m("f")
    num = m("num")
    for child in children:
        num.append(child)
    den = m("den")
    den.append(math_run(denominator, roman=True))
    node.append(num)
    node.append(den)
    return node


def add_equation(doc: Document, condition: bool = False) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    omath_para = m("oMathPara")
    omath = m("oMath")
    if not condition:
        omath.append(subscript("E", "kWh", sub_roman=True))
        omath.append(math_run(" = "))
        numerator = [subscript("P", "start,W", sub_roman=True), math_run(" × "), subscript("Δt", "s", sub_roman=True)]
        omath.append(fraction(numerator, "3,600,000"))
    else:
        omath.append(math_run("0 < "))
        omath.append(subscript("Δt", "s", sub_roman=True))
        omath.append(math_run(" ≤ 16 "))
        omath.append(math_run("s", roman=True))
    omath_para.append(omath)
    para._p.append(omath_para)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=85, bottom=80, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "insideH"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "8C8C8C")
    for edge in ("left", "right", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def style_table(table, widths: list[float], font_size: float = 8.3) -> None:
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            tr_pr.append(repeat_header)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6F8")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(font_size)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 8.3) -> None:
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(table, widths, font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    para = doc.add_paragraph()
    para.style = "Heading 1" if level == 1 else "Heading 2"
    para.paragraph_format.keep_with_next = True
    para.add_run(text)


def add_body(doc: Document, text: str, citation: str | None = None) -> None:
    para = doc.add_paragraph()
    para.style = "Body Text"
    para.add_run(text)
    if citation:
        para.add_run(" " + citation)


def add_bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="Body Text")
    para.paragraph_format.left_indent = Inches(0.24)
    para.paragraph_format.first_line_indent = Inches(-0.18)
    para.add_run("• ")
    para.add_run(text)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    body = styles["Body Text"]
    body.base_style = normal
    body.font.name = "Times New Roman"
    body.font.size = Pt(12)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Inches(0.25)
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(4)
    body.paragraph_format.line_spacing = 1.12

    h1 = styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(13)
    h1.paragraph_format.space_after = Pt(5)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.italic = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(3)
    h2.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("Home Energy Tracker")
    header_run.font.name = "Times New Roman"
    header_run.font.size = Pt(8.5)
    header_run.font.italic = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_label = footer.add_run("Page ")
    page_label.font.name = "Times New Roman"
    page_label.font.size = Pt(9)
    add_page_field(footer)

    doc.core_properties.title = title


def add_front_matter(doc: Document, source: Document) -> None:
    title = source.paragraphs[0].text.strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)

    author_lines = [
        "Rohit Shrivastava¹, Harsh Gupta², Anshika Rai³, Syamasudha Veeragandham⁴",
        "¹,²,³,⁴School of Computer Science and Engineering, Vellore Institute of Technology, Vellore,",
        "Tamil Nadu, India – 632014",
        "Email: ¹rohit.lalbahadur2023@vitstudent.ac.in, ²harsh.gupta2023@vitstudent.ac.in,",
        "³anshika.rai2023@vitstudent.ac.in, ⁴syamasudha.v@vit.ac.in",
        "⁴Corresponding author",
    ]
    for index, line in enumerate(author_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0 if index < 5 else 8)
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10 if index < 5 else 9)
        if index == 5:
            r.italic = True

    abstract = (
        "Residential energy monitoring requires a traceable path from telemetry admission to storage, analysis, and controlled notification. "
        "This study presents a local Home Energy Tracker implemented as an event-driven microservices composition with gateway-mediated access control, device-token ingestion, Apache Kafka, InfluxDB, MySQL, deterministic insights, alert deduplication, baseline Random Forest inference, and service observability. "
        "The repository's legacy classifier was historically evaluated on a supplied 5,000-record labelled CSV and was reproduced offline with the same outputs; its held-out F1 of 0.902439 is therefore historical-dataset-specific evidence rather than a REFIT or field result. "
        "A separate Wokwi-compatible request verifies model loading and response generation only. "
        "To strengthen data provenance, five official REFIT cleaned files were checksum-verified and H1/H2/H4 were normalised into 468,847 interval-energy records, with 386.398411389 kWh estimated by an actual-delta sample-hold rule. "
        "H3/H11 remain a solar-affected sensitivity cohort. "
        "End-to-end REFIT ingestion, replay, alert-deduplication validation, and REFIT anomaly-accuracy evaluation are not completed. "
        "The contribution is consequently an implementation and reproducibility study with explicit evidence boundaries."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    label = p.add_run("Abstract—")
    label.bold = True
    label.font.name = "Times New Roman"
    body = p.add_run(abstract)
    body.font.name = "Times New Roman"
    body.font.size = Pt(11.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    label = p.add_run("Keywords—")
    label.bold = True
    label.font.name = "Times New Roman"
    body = p.add_run(
        "Home energy management; microservices; Apache Kafka; anomaly detection; REFIT; reproducibility; IoT telemetry; alert deduplication."
    )
    body.font.name = "Times New Roman"
    body.font.size = Pt(11)


def add_introduction(doc: Document) -> None:
    add_heading(doc, "1. Introduction")
    add_body(doc, "Residential energy monitoring is not limited to displaying a consumption value. A usable platform must admit telemetry, retain traceable records, derive information, and issue bounded notifications without confusing software evidence with physical-meter accuracy. Microservices-based residential monitoring architectures motivate the separation of these responsibilities into independently inspectable components.", "[1], [2]")
    add_body(doc, "The scope of this study is a local Home Energy Tracker rather than a calibrated utility-metering product. Its telemetry is simulator-compatible, and the study therefore evaluates software-path evidence, reproducibility, and data preparation rather than physical sensor calibration or billing correctness.", "[13], [15]")
    add_body(doc, "An event-driven design is relevant because telemetry admission need not be coupled to downstream aggregation, storage, inference, and alert delivery. Microservices and hybrid communication designs reported for energy-management contexts provide an architectural rationale for this separation, while their deployment settings are not treated as numerical baselines for the present work.", "[1], [3]")
    add_body(doc, "In the implemented system, user-facing requests are mediated by JWT, role, and ownership controls backed by Keycloak. Simulator-compatible telemetry is admitted through a device-bound token. These mechanisms establish an explicit admission boundary before subsequent event processing.", "[1], [3]")
    add_body(doc, "Accepted telemetry is passed to Kafka, from which the Usage Service writes time-series usage and windowed aggregations. InfluxDB retains time-series usage, while MySQL retains household, device, and alert state. This division supports inspection of the processing path and separates relational state from time-series storage.", "[1], [2]")
    add_body(doc, "The system additionally exposes deterministic bill forecasting, appliance ranking, standby detection, saving tips, and related insights. These functions are useful only when their inputs can be traced to admitted telemetry; the project therefore emphasises record handling and evidence boundaries rather than dashboard presentation alone.", "[11], [12]")
    add_body(doc, "Anomaly detection is a further motivation, but smart-meter anomaly research depends on the relationship among the measurement source, label definition, feature construction, and evaluation protocol. Recent work examines faults, misuse, and real-time anomaly detection in smart electricity contexts, illustrating that a model output is not by itself a validated household result.", "[8]–[10]")
    add_body(doc, "The installed Random Forest is a legacy repository artefact. It can be reproduced under its documented local protocol, but the source, licence, label-generation procedure, and causal provenance of the supplied historical labelled CSV remain incomplete. Consequently, its metrics are retained as reproducibility evidence rather than generalisation evidence.", "[4], [8]")
    add_body(doc, "REFIT offers a better documented recorded-household source for future evaluation. Its published household measurements and cleaned dataset records provide a traceable data origin; nevertheless, aggregate-power preparation and anomaly-label alignment must be verified before they can support end-to-end anomaly claims.", "[5]–[7]")
    add_body(doc, "This paper documents the implemented architecture, the reproducible historical baseline, and a checksum-verified REFIT preprocessing path. It explicitly separates completed work from pending isolated provisioning, replay, alert-deduplication validation, and REFIT anomaly-accuracy evaluation.", "[1]–[3], [5]–[7]")

    add_heading(doc, "1.1 Research Objectives", 2)
    objectives = [
        "To analyse the architectural and data-handling requirements of a residential energy-monitoring platform with traceable telemetry admission and controlled downstream processing.",
        "To implement and document an event-driven microservices path for device-token telemetry admission, Kafka processing, storage, baseline inference, alert handling, and observability.",
        "To reproduce and characterise the existing historical Random Forest baseline without representing historical CSV metrics or endpoint smoke-test evidence as real-household accuracy.",
        "To prepare and validate checksum-verified REFIT preprocessing artefacts for H1/H2/H4 while documenting the isolated provisioning and evaluation requirements that remain before an end-to-end experiment."
    ]
    for i, objective in enumerate(objectives, start=1):
        add_body(doc, f"{i}. {objective}")


def add_related_work(doc: Document) -> None:
    add_heading(doc, "2. Related Work")
    add_body(doc, "Núñez et al. describe a microservices-based architecture for residential energy-efficiency monitoring. Its relevance lies in the deliberate allocation of residential-monitoring responsibilities across services; the present study adapts that architectural motivation to a local evidence-aware prototype rather than attempting to reproduce the reported deployment setting.", "[1]")
    add_body(doc, "Motta et al. present a conceptual cloud architecture for home energy-management systems. That work reinforces the distinction between acquisition, operational processing, and analytical use, which is reflected here in separately described admission, usage, inference, and alert responsibilities.", "[2]")
    add_body(doc, "Veloso et al. discuss a microservices-based energy-management solution with hybrid communication in smart-grid environments. Their work is relevant to the use of decoupled components, but it does not remove the need to establish what individual services have actually been validated in a local prototype.", "[3]")
    add_body(doc, "Hernández et al. study anomaly detection in daily activities using smart-meter data. It illustrates that anomaly-oriented conclusions depend on a documented measurement and evaluation context, a requirement that constrains the interpretation of the legacy labelled CSV used in this project.", "[4]")
    add_body(doc, "Murray et al. provide the published two-year UK household load-measurement study underlying REFIT. Its time-stamped household measurements make it suitable for a reproducible recorded-data preparation workflow, although dataset availability does not automatically establish an aggregate anomaly ground truth.", "[5]")
    add_body(doc, "The University of Strathclyde and Zenodo REFIT cleaned-data records provide a traceable dataset route for the present preprocessing work. The local study uses those records for file-level checksum verification and cohort preparation, while retaining the distinction between source data and a completed replay experiment.", "[6], [7]")
    add_body(doc, "Kulkarni et al. address anomaly detection in smart electric meters for faults and misuse. The topic aligns with the alerting motivation of this study, but the present paper does not claim that its fixed legacy Random Forest has been validated on the same fault or misuse problem.", "[8]")
    add_body(doc, "Abu-Arqoub et al. describe an IoT–AI framework for real-time energy anomaly detection, and Maitra et al. present a convolutional-autoencoder approach with a dynamic threshold. These studies motivate model and threshold transparency; however, their methods are not substituted for the repository's existing Random Forest artefact.", "[9], [10]")
    add_body(doc, "Karthika et al. and Liu and Hsiao address smart energy monitoring and smart-plug home energy management with IoT and anomaly detection. They are relevant to the user-facing monitoring context, while the present work focuses more narrowly on traceable service boundaries and controlled evidence interpretation.", "[11], [12]")
    add_body(doc, "Hidayat and Al-Faris, together with Hossain and Hassan, report IoT-based or real-time electrical power-monitoring systems. Such work highlights the practical value of meter-facing telemetry, but no real hardware calibration or physical meter comparison is claimed by the simulator-compatible implementation reported here.", "[13], [15]")
    add_body(doc, "Sirisumrannukul et al. combine cooling-management control with an ANNs–PSO data-analytics microservice platform. This illustrates the broader role of analytics services in energy management, while the current work does not introduce an optimisation method or make a control-performance claim.", "[14]")
    add_body(doc, "Recent IoT energy-consumption tracking, smart-city energy management, and real-time alert systems emphasise monitoring, energy-saving decisions, and notification capabilities. In contrast, the present paper treats alert-deduplication and recovery testing as evidence-gated requirements rather than assuming that a configured alert path has been empirically validated.", "[16]–[18]")

    add_table(
        doc,
        "Table I. Representative prior works and their relationship to the present study",
        ["References", "Primary topic", "Relationship and boundary"],
        [
            ["[1]–[3]", "Microservices and cloud energy architectures", "Motivate component separation; no deployment metric is reused as a benchmark."],
            ["[4], [8]–[10]", "Smart-meter and real-time anomaly detection", "Motivate evaluation discipline; no external model result is claimed for the legacy Random Forest."],
            ["[5]–[7]", "REFIT measurements and cleaned-data records", "Support provenance and repeatable preparation; do not supply validated aggregate anomaly labels for this experiment."],
            ["[11]–[13], [15]", "IoT/home energy monitoring", "Provide application context; the present system does not claim physical sensor calibration."],
            ["[14], [16]–[18]", "Analytics, tracking, energy management, and alerts", "Inform future analytic and alerting scope; replay, deduplication, and recovery tests remain pending."],
        ],
        [0.9, 1.8, 3.55],
        8.6,
    )

    add_heading(doc, "2.1 Research Gaps Identified", 2)
    add_heading(doc, "Research Gap 1: Traceability across the monitoring pipeline", 2)
    add_body(doc, "Prior work motivates modular energy-monitoring architectures, but a local prototype must still demonstrate how telemetry admission, event processing, persistence, analytics, and alert delivery connect in an inspectable path. This work addresses part of the gap through an implemented gateway-to-Kafka-to-service architecture and explicit component-level evidence boundaries.", "[1]–[3]")
    add_heading(doc, "Research Gap 2: Evidence boundaries in anomaly-detection evaluation", 2)
    add_body(doc, "Historical baseline metrics, endpoint smoke tests, and data-preprocessing checks answer different questions. The historical Random Forest metrics describe a supplied labelled CSV; the smoke test shows interface execution; and REFIT preparation demonstrates source integrity and deterministic transformation. Treating any one of these as a replacement for the others would overstate the evidence.", "[4], [8]–[10]")
    add_heading(doc, "Research Gap 3: Reproducible recorded-household evaluation", 2)
    add_body(doc, "REFIT preprocessing can establish a verified, time-aware input cohort, but it does not by itself establish ingestion acceptance, alert behaviour, anomaly labels, or model accuracy. An isolated identity and device-provisioning environment, pilot replay, event reconciliation, and validated label alignment remain necessary before recorded-household claims can be made.", "[5]–[7]")
    add_heading(doc, "Summary of Research Gaps", 2)
    add_body(doc, "The implemented architecture and exact legacy reproduction address traceability and reproducibility at the software and historical-baseline levels. Checksum-verified REFIT normalisation addresses data provenance. The gap that remains is the controlled, isolated, and evidence-producing transition from prepared REFIT records to a verified end-to-end evaluation.")
    add_heading(doc, "Problem Statement", 2)
    add_body(doc, "The technical problem is to build and document a residential energy-monitoring platform in which simulator-compatible telemetry, event processing, storage, analytics, inference, and alert handling can be traced without confusing an implemented component with a completed scientific evaluation. The project therefore requires an evidence-aware system description, a reproducible baseline, and a recorded-household data-preparation path while preserving the remaining gates for future end-to-end validation.")


def add_methodology(doc: Document, media_dir: Path) -> None:
    add_heading(doc, "3. Materials and Methods")
    add_heading(doc, "3.1 Overall System Methodology", 2)
    add_body(doc, "The system accepts a Wokwi-compatible telemetry payload containing deviceId, energyConsumed, and timestamp. User requests follow JWT, role, and ownership controls, whereas simulator-compatible ingestion follows a device-bound-token path. The Ingestion Service validates the agreed telemetry contract and publishes energy-usage events to Kafka.")
    add_body(doc, "The Usage Service consumes energy-usage events and stores time-stamped usage and windowed aggregates. InfluxDB supports time-series usage, while MySQL holds household, device, and alert state. The baseline inference service independently consumes the relevant event class and emits anomaly results. The Alert Service processes threshold and anomaly paths and is configured with a 15 min deduplication window. Prometheus and Grafana are supported for service observation.")
    add_body(doc, "The study treats architecture implementation, historical-baseline reproduction, live interface smoke testing, and REFIT preprocessing as separate evidence streams. This separation ensures that successful service startup or a single inference response is not represented as an accuracy, reliability, or alert-correctness result.")

    add_heading(doc, "3.2 Proposed System Architecture", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(media_dir / "image1.png"), width=Inches(6.05))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    r = cap.add_run("Fig. 1. ")
    r.bold = True
    cap.add_run("Developed architecture for the Home Energy Tracker. The REFIT branch is prepared but is not connected to the runtime.")
    add_body(doc, "Figure 1 distinguishes solid implemented paths from the dashed REFIT preparation path. The solid path connects simulator-compatible telemetry to gateway admission, ingestion, Kafka, usage aggregation, time-series storage, baseline inference, anomaly results, alert processing, relational state, and observability. The dashed REFIT path stops before isolated provisioning, replay, reconciliation, and labelled evaluation.")

    add_table(
        doc,
        "Table II. System components deployed and evidence limits",
        ["Component", "Implemented behaviour", "Evidence boundary"],
        [
            ["Admission", "Keycloak-backed user routes and device-token telemetry admission at the gateway.", "Functional architecture and some authorization controls; no production security evaluation."],
            ["Brokering and ingestion", "Ingestion validates telemetry and publishes Kafka energy-usage events.", "Historical/local event-path evidence; no REFIT replay has been run."],
            ["Storage and usage", "Usage aggregation writes time-stamped usage to InfluxDB; household, device, and alert state are retained in MySQL.", "Component and historical integration evidence; no utility-bill verification."],
            ["Baseline inference", "A FastAPI service loads the fixed Random Forest and returns anomaly results.", "REFIT accuracy is not established by historical-dataset evaluation or a live smoke test."],
            ["Alerting and monitoring", "Threshold/anomaly paths reach Alert Service; a 15 min deduplication window and Prometheus/Grafana observation are configured.", "Historical local evidence exists; fault and replay tests after isolation remain pending."],
            ["REFIT data path", "Checksum-verified, normalised H1/H2/H4 artefacts are prepared separately.", "Device provisioning, ingestion, alert assessment, and tagged-anomaly evaluation are unfinished."],
        ],
        [1.05, 2.7, 2.5],
        8.2,
    )

    add_heading(doc, "3.3 Identity, Admission, and Event Processing", 2)
    add_body(doc, "The gateway forms the API boundary. JWT-backed user routes apply role and ownership checks, while simulator-compatible device telemetry uses a device-bound token. The design does not equate a REFIT household identifier with a backend device identifier; future recorded-data replay requires separately provisioned households, devices, and credentials.")
    add_body(doc, "Following contract validation, energy-usage events are published to Kafka. Consumers then carry out usage aggregation and baseline inference independently. This decoupling allows downstream services to be examined without representing Kafka connectivity or component availability as proof of end-to-end event reconciliation.")

    add_heading(doc, "3.4 Baseline Model and Alerting Boundary", 2)
    add_body(doc, "The persistent Random Forest artefact is random_forest-csv-20261003T205417Z. It uses Electricity_Consumed, Temperature, Humidity, Wind_Speed, Avg_Past_Consumption, Hour, DayOfWeek, and ConsumptionDeviation. The historical protocol retained a chronological 3,500/750/750 train/validation/test split and selected the decision threshold using validation F1.")
    add_body(doc, "The saved model uses 300 trees, min_samples_leaf = 2, class_weight = balanced_subsample, random_state = 42, and threshold 0.31. Wokwi-compatible payloads omit four contextual inputs, for which the service reports median imputation and derives temporal features from the timestamp. A resulting prediction is therefore an interface smoke test rather than a statistical validation of the historical model.")
    add_body(doc, "Alert Service processes threshold and anomaly paths with a configured 15 min deduplication window. The existence of this design does not establish deduplication correctness, notification delivery, latency, recovery behavior, or fault tolerance, because the requisite isolated replay and fault tests have not been completed.")

    add_heading(doc, "3.5 REFIT Preprocessing and Interval-Energy Estimation", 2)
    add_body(doc, "Five local REFIT cleaned files (H1, H2, H3, H4, and H11) were checked against the official Zenodo MD5 manifest. Whole-house aggregate active power was selected in watts. H1/H2/H4 form the primary cohort, while H3/H11 remain a separate sensitivity cohort because the REFIT descriptor records solar-generation interference and no solar or net-export correction was performed.")
    add_body(doc, "For each eligible consecutive record, interval energy is estimated from the starting active-power sample and the actual elapsed interval. The following expression preserves the source method and units.")
    add_equation(doc, condition=False)
    add_body(doc, "Only intervals satisfying the following condition are integrated.")
    add_equation(doc, condition=True)
    add_body(doc, "Longer cross-gap periods are excluded without interpolation, filling, capping, or synthesis. Negative power is excluded from affected integration but retained in the source audit. IQR flags are diagnostic rather than clipping rules, and REFIT Issues values are treated as context rather than anomaly labels.")

    add_heading(doc, "3.6 Algorithmic Procedure for REFIT Normalisation", 2)
    p = doc.add_paragraph(style="Body Text")
    lead = p.add_run("Algorithm 1. REFIT interval-energy normalisation")
    lead.bold = True
    p.paragraph_format.first_line_indent = Inches(0)
    steps = [
        "Input: an MD5-verified cleaned REFIT household file and a selected 14-day time window.",
        "Order records by the preserved Unix timestamp and retain the corresponding Time value.",
        "For each consecutive pair, calculate the actual elapsed interval Δt.",
        "When 0 < Δt_s ≤ 16 s and the starting aggregate power is non-negative, estimate interval energy using the stated left-rectangle/sample-hold rule.",
        "Otherwise, retain the source record for audit and add the interval to the excluded-record log; do not interpolate, cap, or synthesize data.",
        "Produce the cleaned cohort, preprocessing summary, quality report, timestamp-interval report, and excluded-record log.",
        "Output: a deterministic primary H1/H2/H4 cohort and a separately marked H3/H11 solar-sensitivity cohort."
    ]
    for step in steps:
        add_bullet(doc, step)

    add_heading(doc, "3.7 Evidence and Reproducibility Workflow", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(media_dir / "image2.png"), width=Inches(6.05))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    r = cap.add_run("Fig. 2. ")
    r.bold = True
    cap.add_run("Evidence and reproducibility workflow. The REFIT phase remains pending until isolated devices, credentials, pilot reconciliation, and verified evaluation are completed.")
    add_body(doc, "The upper workflow lane records the legacy 5,000-record historical protocol, validation-threshold selection, offline reproduction, and endpoint smoke-test boundary. The lower lane records official REFIT verification, cohort separation, actual-delta normalisation, and byte-identical repeatability. The dashed final stage is intentionally not represented as completed ingestion or anomaly evaluation.")


def add_experimental_setup(doc: Document) -> None:
    add_heading(doc, "4. Experimental Results and Discussion")
    add_heading(doc, "4.1 Experimental Conditions and Evidence Classes", 2)
    add_body(doc, "Four evidence classes are distinguished in this study: architectural implementation evidence, historical supervised-model evidence, deterministic offline reproduction evidence, and REFIT preprocessing evidence. Each class has a different input, procedure, and permissible conclusion. The separation prevents a service response, a repeated training run, or a data-quality result from being reported as an interchangeable accuracy claim.")
    add_body(doc, "The experimental environment documented by the source comprises an 8 GB Windows laptop with a Docker memory limit of 3.696 GiB. The observed main stack required approximately 2.89 GiB. Consequently, an isolated REFIT provisioning stack must not operate concurrently with the main stack on this machine; deployment status cannot be used as a proxy for scientific validation.")
    add_heading(doc, "4.2 Software Components and Data Sources", 2)
    add_body(doc, "The implemented path uses an API gateway, Keycloak, device-token admission, Kafka, InfluxDB, MySQL, a FastAPI Random Forest inference service, Alert Service, Prometheus, and Grafana. Version numbers not recorded in the supplied research evidence are not inferred here. The legacy supervised CSV was supplied within the repository, whereas REFIT is an official recorded-household source with separate dataset references.", "[5]–[7]")
    add_heading(doc, "4.3 Historical Baseline Protocol", 2)
    add_body(doc, "The historical Random Forest was evaluated under the source protocol with a chronological 70/15/15 split represented by 3,500/750/750 train/validation/test records. The test set contains 750 records. Threshold selection was made using validation F1. The logistic-regression comparator used the same held-out split and threshold 0.28.")
    add_heading(doc, "4.4 Offline Reproduction and Live Smoke Test", 2)
    add_body(doc, "A deterministic trainer was applied to 5,000 copied historical records. Exact agreement of reported metrics and the Random Forest artefact SHA-256 indicates reproduction of the legacy experiment with the locally available inputs. Separately, a Wokwi-compatible request using deviceId 1, energyConsumed 0.55, and timestamp 2026-10-11T00:00:00Z tested the deployed inference interface. It is not a latency, alert, or accuracy test.")
    add_heading(doc, "4.5 REFIT Data Configuration and Repeatability", 2)
    add_body(doc, "The H1, H2, and H4 selected windows are 14-day primary-cohort periods. Their processed intervals are quantified in Table IV. The preprocessing pipeline was executed twice with the same five validated source files, criteria, and configuration. The created artefact sets were byte-identical and included the cleaned primary and solar-sensitivity cohorts, manifest, excluded-record log, quality report, preprocessing summary, and timestamp-interval report.")

    add_table(
        doc,
        "Table III. Historical legacy-baseline appraisal, exact reproduction, and smoke-test boundary",
        ["Evidence item", "Protocol", "Observed result", "Permissible interpretation"],
        [
            ["Historical Random Forest", "Chronological 70/15/15 split; 750 test records; validation-tuned threshold.", "Precision 0.880952; Recall 0.925000; F1 0.902439; ROC-AUC 0.997430; PR-AUC 0.948532.", "Historical CSV-specific evaluation, not REFIT or a new field result."],
            ["Historical logistic regression", "Same held-out split; threshold = 0.28.", "Precision 0.053333; Recall 1.000000; F1 0.101266; ROC-AUC 0.422852; PR-AUC 0.059710.", "Legacy comparator only."],
            ["Offline reproduction", "Current deterministic trainer with 5,000 copied records.", "All reported metric differences = 0; reproduced Random Forest artefact SHA-256 matched the legacy artefact.", "Reproduces the legacy experiment; does not prove external validity."],
            ["Live inference smoke test", "deviceId 1; energyConsumed 0.55; timestamp 2026-10-11T00:00:00Z.", "anomaly = false; probability = 0.0; threshold = 0.31; four context features imputed.", "Confirms endpoint execution and declared imputation only."],
        ],
        [1.15, 1.6, 1.85, 1.65],
        7.8,
    )

    add_table(
        doc,
        "Table IV. Normalisation results for the REFIT primary cohort",
        ["Home", "Time window (UTC)", "Chosen rows", "Intervals retained", "Rows excluded", "Energy estimated (kWh)", "Quality output"],
        [
            ["H1", "2014-10-01 to 2014-10-15", "153,762", "153,578", "184", "104.820956111", "PASS; 95.97% native cadence"],
            ["H2", "2014-10-01 through 2014-10-15", "155,097", "154,985", "112", "170.354373889", "PASS; 95.64% native cadence"],
            ["H4", "22 Apr 2015 to 6 May 2015", "160,324", "160,284", "40", "111.223081389", "PASS; 99.30% native cadence"],
            ["Total primary", "Three 14-day periods", "469,183", "468,847", "336", "386.398411389", "All record reconciliations passed"],
        ],
        [0.55, 1.2, 0.78, 0.92, 0.72, 1.0, 1.1],
        7.2,
    )
    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(5)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = note.add_run("Note. ")
    r.bold = True
    note.add_run("Aggregate is whole-house active power in W. Energy is a true-delta, left-rectangle/sample-hold estimate. H3/H11 sensitivity data are omitted from the primary summary because of the known solar-interference caveat.")
    for run in note.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.italic = True

    add_heading(doc, "4.6 Results and Discussion", 2)
    add_body(doc, "The Random Forest has higher historical held-out F1, ROC-AUC, and PR-AUC than the logistic-regression comparator in Table III. The exact offline reproduction and matching artefact hash show that the legacy repository experiment can be regenerated from the available copied input and configuration. These facts do not cure the missing provenance, licence, label-generation procedure, household identity, or causal design of Avg_Past_Consumption in the historical CSV.")
    add_body(doc, "The live smoke-test response shows that the deployed interface can load the fixed model and process a Wokwi-compatible payload with declared imputation. It does not measure response-time distribution, throughput, alert delivery, alert deduplication, recovery, or statistical accuracy. The Wokwi values also lie outside the 0–1 Electricity_Consumed range of the training CSV for five of six evaluated simulator boundaries, which motivates future range or out-of-distribution controls.")
    add_body(doc, "REFIT processing establishes a traceable H1/H2/H4 primary cohort of 468,847 retained intervals and 386.398411389 kWh estimated energy across three selected 14-day periods. No duplicate or out-of-order Unix timestamps were reported in those windows, and each household passed its defined cadence, coverage, missing-slot, and maximum-gap gates. The two preprocessing outputs were byte-identical; the main output SHA-256 was dfdba435f35a1548661309fd67b267548f419edd489cacf25c7566098f189163.")
    add_body(doc, "This preprocessing evidence is intentionally narrower than a model-performance claim. REFIT appliance-level annotations are not treated as aggregate anomaly ground truth: only H1 overlaps the selected primary cohort and the annotation schema, counts, timestamp semantics, and alignment to aggregate readings were not confirmed. Similarly, the prepared interval-energy mapping has not yet been accepted by the live ingestion API.")

    add_table(
        doc,
        "Table V. Publication border and current result status",
        ["Region", "Proven outcome", "Not configured or not validated"],
        [
            ["System implementation", "Gateway, event processing, storage, deterministic insights, alert path, and observability are implemented in the local composition.", "Production-scale capacity, full recovery behavior, and independent penetration security testing."],
            ["Baseline legacy", "A dated endpoint smoke test passed, and offline reproduction of historical Random Forest metrics was exact with the copied CSV.", "External data provenance, actual household generalisation, REFIT accuracy, and a complete post-restart ML-to-alert outcome."],
            ["REFIT data", "Five local official files were MD5-validated; H1/H2/H4 interval artefacts were normalised and replicated byte-for-byte.", "End-to-end ingestion, device/token provisioning, pilot/full replay, alert deduplication, and aggregate anomaly-label evaluation."],
            ["Deployment viability", "Isolated-experiment design and safety assessments define necessary provisioning and storage boundaries.", "Safe isolated deployment on the 8 GB laptop while the main stack is running; the measured memory constraint prevents concurrent operation."],
        ],
        [1.25, 2.7, 2.35],
        8.1,
    )


def add_limitations_conclusion_backmatter(doc: Document, source: Document) -> None:
    add_heading(doc, "5. Limitations and Future Work")
    add_body(doc, "The first limitation is that Wokwi-compatible telemetry is synthetic and has not been physically calibrated. The second is the incomplete provenance and feature-building context of the legacy supervised CSV, which means that its historical metrics cannot be generalised. The third is that REFIT provides documented aggregate power but not a verified aggregate-anomaly-label protocol for the selected H1/H2/H4 cohort.")
    add_body(doc, "The fourth limitation is local resource capacity. The source evidence records an 8 GB Windows laptop, a Docker memory limit of 3.696 GiB, and an observed main-stack requirement of approximately 2.89 GiB. The isolated REFIT provisioning stack must therefore not run concurrently with the main stack on this machine. This restriction prevents deployment status from being used as a substitute for controlled validation.")
    add_body(doc, "Future work requires an isolated Keycloak realm, separate database volumes, a Kafka broker, three household profiles, and three device-bound credentials. It must then verify the derived_interval_energy_kwh to energyConsumed mapping against the actual ingestion contract, reconcile a small pilot replay, and only then proceed to H1/H2/H4 replay. Validated label alignment, household-held-out or temporal evaluation, input-range controls, repeated latency/throughput measurement, alert-deduplication evidence, and fault/recovery testing remain necessary.")

    add_heading(doc, "6. Conclusion")
    add_body(doc, "This work presents an evidence-aware Home Energy Tracker with gateway-mediated admission, Kafka event processing, time-series and relational storage, deterministic insights, a baseline Random Forest service, duplicate-aware alerting, and monitoring. The legacy classifier was exactly reproduced offline with its historical results, and a separate endpoint smoke test was restricted to interface execution evidence. REFIT work established checksum-verified official inputs, a separated H1/H2/H4 primary cohort, actual-delta interval-energy normalisation, quality controls, and deterministic artefact reproduction.")
    add_body(doc, "The study therefore addresses traceability and reproducibility gaps at the architecture, historical-baseline, and recorded-data-preparation levels. It does not establish real-world anomaly accuracy, completed REFIT ingestion, alert-deduplication correctness, production resilience, or generalisable performance. The next credible step is isolated provisioning followed by pilot replay, event reconciliation, validated label alignment, and a separately reported end-to-end evaluation.")

    add_heading(doc, "CRediT Authorship Contribution Statement")
    add_body(doc, "The supplied materials identify the author list but do not document individual CRediT roles. Author-specific contribution roles require confirmation from the authors before formal submission.")

    add_heading(doc, "Funding Statement")
    add_body(doc, source.paragraphs[63].text)
    add_heading(doc, "Declaration of Competing Interests")
    add_body(doc, source.paragraphs[65].text)
    add_heading(doc, "Data Availability")
    add_body(doc, "Official REFIT cleaned data are identified through the dataset records cited in [6] and [7]. The supplied historical labelled CSV and locally derived preprocessing artefacts were not accompanied by a public repository or release statement in the source materials; their availability is therefore subject to the responsible authors and applicable source terms.")

    add_heading(doc, "References")
    for source_para in source.paragraphs[67:85]:
        text = source_para.text.strip()
        if not text:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.92
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.2)


def extract_media(source_path: Path, media_dir: Path) -> None:
    media_dir.mkdir(parents=True, exist_ok=True)
    with ZipFile(source_path) as zf:
        for name in ("word/media/image1.png", "word/media/image2.png"):
            destination = media_dir / Path(name).name
            with zf.open(name) as src, open(destination, "wb") as dst:
                shutil.copyfileobj(src, dst)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: create_review2_final.py SOURCE.docx OUTPUT.docx")
    source_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    media_dir = output_path.parent / "review2_media"
    extract_media(source_path, media_dir)
    source = Document(source_path)
    document = Document()
    configure_document(document, source.paragraphs[0].text.strip())
    add_front_matter(document, source)
    add_introduction(document)
    add_related_work(document)
    add_methodology(document, media_dir)
    add_experimental_setup(document)
    add_limitations_conclusion_backmatter(document, source)
    document.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
