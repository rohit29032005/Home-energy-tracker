from docx import Document

document = Document(r"research-output/Home_Energy_Tracker_Final_Verified_Research_Paper.docx")
print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} sections={len(document.sections)}")
for index, paragraph in enumerate(document.paragraphs):
    text = " ".join(paragraph.text.split())
    if text:
        print(f"P{index}|{paragraph.style.name}|{text[:280]}")
for table_index, table in enumerate(document.tables):
    print(f"TABLE {table_index} {len(table.rows)}x{len(table.columns)}")
    for row in table.rows:
        print(" || ".join(" ".join(cell.text.split())[:140] for cell in row.cells))
