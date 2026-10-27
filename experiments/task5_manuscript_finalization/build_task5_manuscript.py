"""Build the evidence-bounded Task 5 manuscript from the retained two-column template.

This builder intentionally does not read, start, or alter the application runtime. It only
creates a new DOCX and replacement scientific figures in the Task 5 output directory.
"""

from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\anju2\Desktop\Home-Energy-Tracker")
TEMPLATE = ROOT / "research-output" / "Home_Energy_Tracker_Publication_Updated_Final.docx"
OUT_DIR = ROOT / "experiments" / "task5_manuscript_finalization"
OUTPUT = OUT_DIR / "Home_Energy_Tracker_Manuscript_Task5_Revised.docx"
FIG1 = OUT_DIR / "figure_1_implemented_architecture.png"
FIG2 = OUT_DIR / "figure_2_evidence_and_refit_workflow.png"


def font(size, bold=False, italic=False):
    candidates = [
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def box(draw, xy, label, *, fill=(246, 248, 250), outline=(70, 80, 90), dashed=False,
        title=False, size=28):
    x1, y1, x2, y2 = xy
    if dashed:
        dash = 12
        for x in range(x1, x2, dash * 2):
            draw.line((x, y1, min(x + dash, x2), y1), fill=outline, width=3)
            draw.line((x, y2, min(x + dash, x2), y2), fill=outline, width=3)
        for y in range(y1, y2, dash * 2):
            draw.line((x1, y, x1, min(y + dash, y2)), fill=outline, width=3)
            draw.line((x2, y, x2, min(y + dash, y2)), fill=outline, width=3)
    else:
        draw.rounded_rectangle(xy, radius=10, fill=fill, outline=outline, width=3)
    f = font(size, bold=title)
    lines = label.split("\n")
    heights = [draw.textbbox((0, 0), line, font=f)[3] for line in lines]
    total = sum(heights) + (len(lines) - 1) * 6
    y = y1 + ((y2 - y1) - total) / 2
    for line, h in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=f)
        draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2, y), line,
                  font=f, fill=(20, 25, 30))
        y += h + 6


def arrow(draw, start, end, *, dashed=False, label=None):
    color = (55, 65, 75)
    width = 4
    if dashed:
        x1, y1 = start
        x2, y2 = end
        if y1 == y2:
            step = 16 if x2 >= x1 else -16
            x = x1
            while (x < x2 if step > 0 else x > x2):
                nx = min(x + step // 2, x2) if step > 0 else max(x + step // 2, x2)
                draw.line((x, y1, nx, y2), fill=color, width=width)
                x += step
        else:
            draw.line((start, end), fill=color, width=width)
    else:
        draw.line((start, end), fill=color, width=width)
    # simple triangular arrow head
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        triangle = [(x2, y2), (x2 - 18 * direction, y2 - 10), (x2 - 18 * direction, y2 + 10)]
    else:
        direction = 1 if y2 >= y1 else -1
        triangle = [(x2, y2), (x2 - 10, y2 - 18 * direction), (x2 + 10, y2 - 18 * direction)]
    draw.polygon(triangle, fill=color)
    if label:
        f = font(22, italic=True)
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        draw.text((mx - 45, my - 30), label, font=f, fill=color)


def create_figures():
    # Figure 1: direct representation of implemented and pending branches.
    im = Image.new("RGB", (2200, 1160), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 32), "Implemented Home Energy Tracker architecture", font=font(40, bold=True), fill=(0, 0, 0))
    d.text((60, 88), "Solid lines: implemented data path. Dashed lines: prepared REFIT data path not connected to runtime.",
           font=font(25), fill=(50, 50, 50))
    box(d, (55, 220, 330, 360), "Wokwi-compatible\ntelemetry", title=True)
    box(d, (410, 220, 720, 360), "API Gateway\nJWT or device-token admission", title=True)
    box(d, (800, 220, 1060, 360), "Ingestion\nService", title=True)
    box(d, (1140, 220, 1385, 360), "Kafka\nenergy-usage", title=True)
    arrow(d, (330, 290), (410, 290))
    arrow(d, (720, 290), (800, 290))
    arrow(d, (1060, 290), (1140, 290))
    box(d, (1480, 100, 1750, 240), "Usage Service\nwindow aggregation", title=True)
    box(d, (1870, 100, 2140, 240), "InfluxDB\ntime-series store", title=True)
    box(d, (1480, 340, 1750, 480), "ML inference\nlegacy Random Forest", title=True)
    box(d, (1870, 340, 2140, 480), "Anomaly results\ntopic", title=True)
    box(d, (1480, 590, 1750, 730), "Alert Service\n15 min deduplication", title=True)
    box(d, (1870, 590, 2140, 730), "MySQL and Mailpit\nmetadata and delivery", title=True)
    arrow(d, (1385, 270), (1480, 170))
    arrow(d, (1750, 170), (1870, 170))
    arrow(d, (1385, 310), (1480, 410))
    arrow(d, (1750, 410), (1870, 410))
    arrow(d, (1870, 410), (1750, 660))
    arrow(d, (1750, 660), (1870, 660))
    box(d, (55, 610, 410, 750), "User, Device and Insight\nservices", title=True)
    box(d, (520, 610, 800, 750), "MySQL\nuser/device/alert state", title=True)
    arrow(d, (410, 680), (520, 680))
    box(d, (880, 610, 1210, 750), "Prometheus and Grafana\nservice observation", title=True)
    box(d, (55, 895, 410, 1050), "Checksum-verified REFIT\nH1/H2/H4 artifacts", dashed=True, fill=(255, 255, 255), title=True)
    box(d, (520, 895, 800, 1050), "Compatibility adapter\nand isolated devices", dashed=True, fill=(255, 255, 255), title=True)
    box(d, (900, 895, 1210, 1050), "Replay and REFIT\nanomaly evaluation", dashed=True, fill=(255, 255, 255), title=True)
    arrow(d, (410, 972), (520, 972), dashed=True)
    arrow(d, (800, 972), (900, 972), dashed=True)
    d.text((1270, 940), "Pending: isolated provisioning, pilot replay,\nend-to-end reconciliation, and labelled evaluation.",
           font=font(27, italic=True), fill=(70, 70, 70))
    im.save(FIG1, optimize=True)

    # Figure 2: source-specific evidence and future validation boundary.
    im = Image.new("RGB", (2200, 980), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 30), "Evidence and reproducibility workflow", font=font(40, bold=True), fill=(0, 0, 0))
    d.text((60, 84), "Completed evidence is separated from the pending REFIT integration experiment.",
           font=font(25), fill=(50, 50, 50))
    d.text((75, 170), "Legacy supervised baseline", font=font(30, bold=True), fill=(20, 20, 20))
    box(d, (70, 230, 410, 380), "Repository CSV\n5,000 labelled records", title=True)
    box(d, (520, 230, 850, 380), "Chronological split\n3,500 / 750 / 750", title=True)
    box(d, (960, 230, 1290, 380), "Historical RF evaluation\nthreshold selected on validation", title=True)
    box(d, (1400, 230, 1725, 380), "Exact offline reproduction\nmetric difference = 0", title=True)
    box(d, (1835, 230, 2140, 380), "Endpoint smoke test\nfunctional only", title=True)
    arrow(d, (410, 305), (520, 305))
    arrow(d, (850, 305), (960, 305))
    arrow(d, (1290, 305), (1400, 305))
    arrow(d, (1725, 305), (1835, 305))
    d.text((75, 500), "REFIT recorded-household preprocessing", font=font(30, bold=True), fill=(20, 20, 20))
    box(d, (70, 560, 410, 710), "Official files\nMD5 verified", title=True)
    box(d, (520, 560, 850, 710), "H1/H2/H4 primary cohort\nsolar H3/H11 separated", title=True)
    box(d, (960, 560, 1290, 710), "Actual-delta interval\nenergy normalization", title=True)
    box(d, (1400, 560, 1725, 710), "Byte-identical\npreprocessing rerun", title=True)
    box(d, (1835, 560, 2140, 710), "Isolated replay and\nanomaly evaluation", dashed=True, fill=(255, 255, 255), title=True)
    arrow(d, (410, 635), (520, 635))
    arrow(d, (850, 635), (960, 635))
    arrow(d, (1290, 635), (1400, 635))
    arrow(d, (1725, 635), (1835, 635), dashed=True)
    d.text((1460, 785), "Pending: no ingestion, alert-deduplication,\nor REFIT accuracy is claimed.",
           font=font(24, italic=True), fill=(70, 70, 70))
    im.save(FIG2, optimize=True)


def set_columns(section, count):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), "360")


def set_font(run, size=10, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def paragraph(doc, text="", *, align=None, before=0, after=4, line=1.0, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        pPr = p._p.get_or_add_pPr()
        keep_next = OxmlElement("w:keepNext")
        pPr.append(keep_next)
    if text:
        set_font(p.add_run(text))
    return p


def rich_paragraph(doc, parts, *, align=None, before=0, after=4, line=1.0, keep=False, size=10):
    p = paragraph(doc, align=align, before=before, after=after, line=line, keep=keep)
    for item in parts:
        if isinstance(item, str):
            text, bold, italic, superscript = item, False, False, False
        else:
            text, bold, italic, superscript = item
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, italic=italic)
        r.font.superscript = superscript
    return p


def heading(doc, text, *, level="major"):
    p = paragraph(doc, before=7 if level == "major" else 4, after=3, keep=True)
    r = p.add_run(text)
    set_font(r, size=10, bold=True, italic=(level == "minor"))
    return p


def caption(doc, kind, number, title):
    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=5, keep=True)
    r = p.add_run(f"{kind} {number}. ")
    set_font(r, size=8.5, bold=True)
    r = p.add_run(title)
    set_font(r, size=8.5, italic=True)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=70, bottom=65, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def make_table(doc, headers, rows, widths, *, size=7.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.95
        r = p.add_run(header)
        set_font(r, size=size, bold=True, color=(255, 255, 255))
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    header_pr.append(tbl_header)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or (len(str(value)) < 26 and idx != 1) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.95
            r = p.add_run(str(value))
            set_font(r, size=size)
    # Prevent Word from splitting the short tables in a visually awkward way.
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    return table


def full_width(doc):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(section, 1)
    return section


def two_columns(doc):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(section, 2)
    return section


def add_figure(doc, path, width, alt_text):
    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=1, keep=True)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)
    return p


def add_references(doc):
    heading(doc, "REFERENCES")
    refs = [
        '[1] I. Nunez, C. Rovetto, E. Cruz, A. Smolarz, D. Concepcion, and E. E. Cano, "Design of a microservices-based architecture for residential energy efficiency monitoring," International Journal of Electronics and Telecommunications, vol. 70, no. 4, pp. 1089-1098, 2024, doi: 10.24425/ijet.2024.152511.',
        '[2] L. L. Motta, W. S. Alvarez, H. R. Carvajal, F. B. Neto, E. R. de Lima, and L. G. P. Meloni, "A cloud architecture for home energy management systems: a conceptual model," Energy Informatics, vol. 8, art. 142, 2025, doi: 10.1186/s42162-025-00599-1.',
        '[3] A. F. S. Veloso, J. V. Reis, Jr., and R. A. L. Rabelo, "A microservices-based solution with hybrid communication for energy management in smart grid environments," Sensors, vol. 26, no. 5, art. 1714, 2026, doi: 10.3390/s26051714.',
        '[4] A. Hernandez, R. Nieto, L. de Diego-Oton, M. C. Perez-Rubio, J. M. Villadangos-Carrizo, D. Pizarro, and J. Urena, "Detection of anomalies in daily activities using data from smart meters," Sensors, vol. 24, no. 2, art. 515, 2024, doi: 10.3390/s24020515.',
        '[5] D. Murray, L. Stankovic, and V. Stankovic, "An electrical load measurements dataset of United Kingdom households from a two-year longitudinal study," Scientific Data, vol. 4, art. 160122, 2017, doi: 10.1038/sdata.2016.122.',
        '[6] University of Strathclyde, "REFIT: Electrical Load Measurements (Cleaned)," dataset, 2016, doi: 10.15129/9ab14b0e-19ac-4279-938f-27f643078cec.',
        '[7] D. Murray and L. Stankovic, "REFIT: Electrical Load Measurements (Cleaned)," Zenodo, 2021, doi: 10.5281/zenodo.5063428.',
    ]
    for ref in refs:
        p = paragraph(doc, before=0, after=2, line=0.95)
        p.paragraph_format.left_indent = Inches(0.14)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        set_font(p.add_run(ref), size=8.1)


def create_document():
    create_figures()
    doc = Document(str(TEMPLATE))
    # Preserve page geometry and styles, but intentionally replace every editable body element.
    body = doc._element.body
    for child in list(body)[:-1]:
        body.remove(child)

    # The original body-level section properties become the front-matter section.
    first = doc.sections[0]
    set_columns(first, 1)
    first.top_margin = Inches(0.55)
    first.bottom_margin = Inches(0.55)
    first.left_margin = Inches(0.62)
    first.right_margin = Inches(0.62)

    # Default Normal style remains template-derived but we avoid its blue heading theme.
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10)

    # Title and authors retain the existing author information and ordering.
    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.0, keep=True)
    r = p.add_run("A Microservices-Based Home Energy Tracker for Event-Driven Monitoring, ")
    set_font(r, size=11, bold=True)
    r = p.add_run("Deduplicated Alerting, and ML-Based Anomaly Detection")
    set_font(r, size=11, bold=True)

    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1, line=1.0, keep=True)
    for name, marker in [("Rohit Shrivastava", "1"), ("Harsh Gupta", "2"), ("Anshika Rai", "3"), ("Syamasudha Veeragandham", "4")]:
        if p.runs:
            set_font(p.add_run(", "), size=9.2, bold=True)
        set_font(p.add_run(name), size=9.2, bold=True)
        sr = p.add_run(marker)
        set_font(sr, size=7, bold=True)
        sr.font.superscript = True
    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1, line=0.95, keep=True)
    set_font(p.add_run("1,2,3,4School of Computer Science and Engineering, Vellore Institute of Technology, Vellore, Tamil Nadu, India - 632014"), size=8.2, bold=True)
    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=1, line=0.95, keep=True)
    set_font(p.add_run("Email: 1rohit.lalbahadur2023@vitstudent.ac.in, 2harsh.gupta2023@vitstudent.ac.in, 3anshika.rai2023@vitstudent.ac.in, 4syamasudha.v@vit.ac.in"), size=7.8)
    p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, line=0.95, keep=True)
    set_font(p.add_run("4Corresponding author"), size=7.8, italic=True)

    abstract = (
        "Residential energy platforms need a traceable route from telemetry admission to storage, "
        "analysis, and controlled notification. This paper presents a local Home Energy Tracker "
        "implemented as an event-driven microservices composition with gateway-mediated access "
        "control, device-token ingestion, Apache Kafka, InfluxDB, MySQL, deterministic insights, "
        "alert deduplication, and a baseline Random Forest inference service. The legacy classifier "
        "was historically evaluated on a repository-supplied 5,000-record labelled CSV and was "
        "exactly reproduced offline; its held-out F1 is reported as historical dataset-specific "
        "evidence, not as a REFIT or field result. A separate live endpoint smoke check confirms "
        "baseline loading and response generation only. To strengthen data provenance, five official "
        "REFIT cleaned files were checksum-verified and H1/H2/H4 were normalized into 468,847 "
        "retained interval-energy records (386.398411389 kWh estimated by an actual-delta sample-hold "
        "rule). H3/H11 are retained separately as solar-affected sensitivity data. End-to-end REFIT "
        "ingestion, alert-deduplication validation, and REFIT anomaly-accuracy evaluation remain pending "
        "because isolated provisioning and safe deployment have not been completed. The contribution is "
        "therefore an implementation and reproducibility study with explicit evidence boundaries, not a "
        "claim of calibrated metering, generalizable anomaly accuracy, or production readiness."
    )
    rich_paragraph(doc, [("Abstract - ", True, False, False), (abstract, False, False, False)], after=4, line=1.0, size=9.2)
    keywords = "Home energy management, microservices, Apache Kafka, anomaly detection, REFIT, reproducibility, IoT telemetry, alert deduplication."
    rich_paragraph(doc, [("Keywords - ", True, False, False), (keywords, False, False, False)], after=5, line=1.0, size=9.0)

    two_columns(doc)

    heading(doc, "I. INTRODUCTION")
    rich_paragraph(doc, [
        ("Residential energy monitoring requires more than a dashboard: it needs a controlled data path from a device-like source to persisted measurements, user-visible analytics, and bounded alerts. ", False, False, False),
        ("This project addresses that integration problem through a local, event-driven Home Energy Tracker rather than through a claim of utility-grade measurement accuracy. ", False, False, False),
        ("The implementation separates admission, ingestion, aggregation, inference, alerting, identity, and observation concerns so that each function can be inspected and tested independently.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("The active runtime accepts Wokwi-compatible telemetry containing deviceId, energyConsumed, and timestamp. User-facing requests are mediated by Keycloak-backed JWT, role, and ownership checks; simulated-device ingestion uses a device-bound token. Kafka decouples acceptance from downstream consumption, while InfluxDB retains time-series usage and MySQL stores household, device, and alert state. The system also exposes deterministic bill, ranking, standby, forecast, and saving-tip functions. This design makes implementation claims about software components and data flow, not about physical sensor calibration.", False, False, False),
    ])

    heading(doc, "II. RELATED WORK AND RESEARCH GAP")
    heading(doc, "A. Positioning Against Prior Work", level="minor")
    rich_paragraph(doc, [
        ("Recent work frames residential energy management as a data-engineering and architecture problem. Nunez et al. propose a flexible microservices architecture for residential energy-efficiency monitoring [1], while Motta et al. describe a conceptual cloud architecture organized around ingestion, operational, and analytical environments [2]. Veloso et al. integrate microservices with hybrid communication in a smart-grid setting [3]. These studies support the relevance of modular, event-capable designs, but their deployment conditions and evaluation targets differ from this local prototype; their results are not used as a numerical benchmark for this work.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("Smart-meter anomaly studies also show why data provenance and evaluation protocol matter. Hernandez et al. evaluate activity-related anomalies using a commercial meter in a real household [4]. In contrast, the current Random Forest was historically evaluated on an internally supplied labelled CSV whose source, licence, label-generation process, and causal provenance of Avg_Past_Consumption are not documented. Its metrics can therefore support only a reproducible legacy-baseline claim. The REFIT study provides a better documented recorded-household source, including aggregate electrical power sampled at a nominal 8 s cadence [5]-[7].", False, False, False),
    ])
    heading(doc, "B. Research Gap and Contribution", level="minor")
    rich_paragraph(doc, [
        ("The identified gap is not a new anomaly-detection algorithm. It is the need for a traceable prototype that ties residential telemetry admission to event-mediated persistence, deterministic insights, baseline inference, and duplicate-aware alert delivery while clearly separating implemented evidence from pending experimentation. The present contribution is threefold: (i) an implemented microservices architecture with an event-driven processing path; (ii) exact offline reproduction of the repository's historical supervised baseline; and (iii) checksum-traceable REFIT preprocessing that establishes a scientifically usable primary cohort and interval-energy representation. The study deliberately does not claim completed REFIT replay, REFIT-based anomaly accuracy, or production resilience.", False, False, False),
    ])

    heading(doc, "III. IMPLEMENTED SYSTEM AND METHODS")
    heading(doc, "A. Event-Driven Architecture", level="minor")
    rich_paragraph(doc, [
        ("At the API boundary, the gateway admits user requests through JWT, role, and ownership controls and admits simulator-compatible telemetry through a device-bound token. The Ingestion Service validates the accepted telemetry contract and publishes energy-usage events to Kafka. The Usage Service consumes these events for time-series persistence and selected-window aggregation. The baseline FastAPI inference service independently consumes the same class of event and publishes anomaly results. Threshold and anomaly pathways are handled by the Alert Service, which persists notification state and applies a 15 min configured deduplication window. Prometheus and Grafana support service observation. The architecture is shown in Fig. 1.", False, False, False),
    ])

    full_width(doc)
    add_figure(doc, FIG1, 6.85, "Implemented Home Energy Tracker architecture with solid active data paths and a dashed pending REFIT integration path")
    caption(doc, "Fig.", "1", "Implemented Home Energy Tracker architecture. The REFIT branch is intentionally shown as prepared but not connected to the runtime.")
    caption(doc, "TABLE", "I", "Implemented system components and evidence boundaries")
    make_table(doc,
        ["Component", "Implemented behavior", "Evidence boundary"],
        [
            ["Admission", "Keycloak-backed user routes; device-token telemetry admission at the gateway.", "Functional architecture and selected authorization checks; not a production security assessment."],
            ["Ingestion and broker", "Ingestion validates telemetry and publishes Kafka energy-usage events.", "Historical and dated local event-path evidence; no REFIT replay has been run."],
            ["Usage and storage", "Usage aggregation writes time-stamped usage to InfluxDB; household, device, and alert state use MySQL.", "Component and historical integration evidence; no utility-bill validation."],
            ["Baseline inference", "FastAPI service loads the fixed Random Forest and produces anomaly results.", "Historical dataset evaluation and a live functional smoke test are distinct from REFIT accuracy."],
            ["Alerting and observation", "Threshold/anomaly pathways target Alert Service; 15 min deduplication is configured; Prometheus/Grafana observe services.", "Historical local tests exist; post-isolation fault and replay tests remain pending."],
            ["REFIT data path", "Checksum-verified, normalized H1/H2/H4 artifacts are prepared separately.", "No device provisioning, ingestion, alert, or labelled anomaly evaluation has been completed."],
        ], [1050, 3250, 4300], size=7.3)

    two_columns(doc)
    heading(doc, "B. Baseline ML Boundary", level="minor")
    rich_paragraph(doc, [
        ("The deployed baseline is the persisted Random Forest artifact random_forest-csv-20261003T205417Z. It uses eight features: Electricity_Consumed, Temperature, Humidity, Wind_Speed, Avg_Past_Consumption, Hour, DayOfWeek, and ConsumptionDeviation. The historical protocol stably ordered the supplied data, used a chronological 3,500/750/750 train/validation/test partition, and selected the decision threshold by validation F1. The saved model uses 300 trees, min_samples_leaf = 2, class_weight = balanced_subsample, random_state = 42, and threshold 0.31.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("A Wokwi-compatible runtime payload lacks four contextual fields. The service explicitly median-imputes Temperature, Humidity, Wind_Speed, and Avg_Past_Consumption and derives time-related features from the timestamp. Consequently, its live three-field prediction is an inference-interface smoke test, not a validation of the historical model's statistical accuracy. The documented Wokwi values are also outside the training CSV's 0-1 Electricity_Consumed range for five of six tested simulator boundaries; this is an interface-scale mismatch requiring future range or out-of-distribution controls.", False, False, False),
    ])
    heading(doc, "C. REFIT Preprocessing Method", level="minor")
    rich_paragraph(doc, [
        ("Five local REFIT cleaned files (H1, H2, H3, H4, and H11) were verified against the official Zenodo MD5 manifest before inspection. Aggregate was selected as whole-house active power in watts; Unix was selected for chronology, while Time was retained verbatim. The primary cohort is H1/H2/H4. H3/H11 are excluded from primary summaries because the REFIT descriptor documents solar-generation interference; they remain a separately labelled sensitivity cohort without solar or net-export correction.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("For each consecutive record, the interval-energy estimate is E_kWh = P_start,W x delta_t,s / 3,600,000. Only intervals satisfying 0 < delta_t <= 16 s are integrated. Longer cross-gap intervals are excluded without interpolation, filling, capping, or synthesis. Negative power excludes only the affected integration but is retained in the source audit; IQR flags are diagnostic rather than clipping rules; and REFIT Issues values are context, not anomaly labels. This preserves the distinction between active power samples and estimated interval energy.", False, False, False),
    ])

    full_width(doc)
    caption(doc, "TABLE", "II", "Historical legacy-baseline evaluation, exact reproduction, and smoke-test boundary")
    make_table(doc,
        ["Evidence item", "Protocol", "Observed result", "Interpretation"],
        [
            ["Historical Random Forest", "Chronological 70/15/15 split; 750 held-out test records; threshold selected on validation.", "Precision 0.880952; recall 0.925000; F1 0.902439; ROC-AUC 0.997430; PR-AUC 0.948532.", "Historical CSV-specific evaluation; not a REFIT or newly reproduced field result."],
            ["Historical logistic regression", "Same held-out partition; threshold 0.28.", "Precision 0.053333; recall 1.000000; F1 0.101266; ROC-AUC 0.422852; PR-AUC 0.059710.", "Comparator for the legacy baseline only."],
            ["Offline reproduction", "Copied 5,000-record input and current deterministic trainer.", "All reported metric differences = 0; reproduced Random Forest artifact SHA-256 matched the legacy artifact.", "Confirms reproducibility of the available legacy experiment, not external validity."],
            ["Live inference smoke test", "Wokwi-compatible input: deviceId 1, energyConsumed 0.55, timestamp 2026-10-11T00:00:00Z.", "anomaly = false; probability = 0.0; threshold = 0.31; four context features reported as imputed.", "Confirms endpoint execution and declared imputation only; it is not an accuracy, latency, or end-to-end alert result."],
        ], [1500, 2900, 2850, 3350], size=7.0)
    caption(doc, "TABLE", "III", "REFIT primary cohort normalization results")
    make_table(doc,
        ["House", "UTC window", "Selected rows", "Retained intervals", "Excluded rows", "Estimated energy (kWh)", "Quality result"],
        [
            ["H1", "2014-10-01 to 2014-10-15", "153,762", "153,578", "184", "104.820956111", "PASS; 95.97% native cadence"],
            ["H2", "2014-10-01 to 2014-10-15", "155,097", "154,985", "112", "170.354373889", "PASS; 95.64% native cadence"],
            ["H4", "2015-04-22 to 2015-05-06", "160,324", "160,284", "40", "111.223081389", "PASS; 99.30% native cadence"],
            ["Primary total", "Three 14-day windows", "469,183", "468,847", "336", "386.398411389", "All record reconciliations PASS"],
        ], [650, 1800, 900, 980, 830, 1200, 1900], size=6.85)
    p = paragraph(doc, after=6, before=2, line=0.95)
    set_font(p.add_run("Note: Aggregate is whole-house active power in watts. Energy is an actual-delta, left-rectangle/sample-hold estimate. H3/H11 sensitivity data are intentionally omitted from Table III because they have a documented solar-interference caveat."), size=7.5, italic=True)

    two_columns(doc)
    heading(doc, "IV. EVALUATION METHODOLOGY AND REPRODUCIBILITY")
    heading(doc, "A. Evidence Classification", level="minor")
    rich_paragraph(doc, [
        ("The evaluation is organized by evidence type rather than by a single aggregated performance score. Historical supervised metrics retain their original provenance and are explicitly labelled as such. The exact offline reproduction checks whether current code and the copied input reproduce historical outputs. The live inference check verifies a specific service contract and model load. REFIT preprocessing evaluates source integrity, cohort selection, timestamp handling, and unit transformation. These evidence types answer different questions and are not interchangeable.", False, False, False),
    ])
    heading(doc, "B. REFIT Reproducibility", level="minor")
    rich_paragraph(doc, [
        ("The preprocessing pipeline was rerun twice with the same five verified source files, selection criteria, and configuration. Both generated artifact sets passed validation and were byte-identical, including cleaned primary and solar-sensitivity cohorts, the manifest, excluded-record log, quality report, preprocessing summary, and timestamp-interval report. The primary output SHA-256 is dfdba435f35a1548661309fd67b267548f419edd489cacf25c7566098f189163. This supports deterministic preparation of the local data artifacts, not their delivery through the deployed ingestion API.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("REFIT's official appliance-level annotations are not treated as aggregate anomaly ground truth. The published annotation coverage overlaps the local primary cohort only at H1, and the annotation-file schema, counts, timestamp semantics, and alignment to aggregate readings have not been verified. No synthetic anomaly generation was used to create a reported REFIT accuracy value. Figure 2 separates completed evidence from the future isolated replay and evaluation path.", False, False, False),
    ])

    full_width(doc)
    add_figure(doc, FIG2, 6.85, "Evidence workflow separating historical baseline evaluation, reproducible REFIT preprocessing, and pending REFIT integration")
    caption(doc, "Fig.", "2", "Evidence and reproducibility workflow. The pending REFIT step requires isolated devices, credentials, pilot reconciliation, and a verified evaluation protocol.")
    caption(doc, "TABLE", "IV", "Current result status and publication boundary")
    make_table(doc,
        ["Area", "Verified result", "Not established"],
        [
            ["System implementation", "Microservices, gateway, event processing, storage, deterministic insights, alert path, and observability are implemented in the local composition.", "Production-scale capacity, complete recovery behavior, and independent penetration security testing."],
            ["Legacy baseline", "Historical Random Forest metrics were reproduced exactly offline from the available copied CSV; a dated endpoint smoke test succeeded.", "External data provenance, real-household generalization, REFIT accuracy, and a complete post-restart ML-to-alert result."],
            ["REFIT data", "Five local official files were MD5 verified; primary H1/H2/H4 interval artifacts were normalized and reproduced byte-for-byte.", "End-to-end ingestion, device/token provisioning, pilot/full replay, alert deduplication, and aggregate anomaly-label evaluation."],
            ["Deployment feasibility", "An isolated experiment design and safety assessments identify required provisioning and storage boundaries.", "A safe isolated deployment on the 8 GB laptop; the main and isolated stacks must not run concurrently under the measured memory constraint."],
        ], [1600, 4450, 4450], size=7.2)

    two_columns(doc)
    heading(doc, "V. RESULTS AND DISCUSSION")
    heading(doc, "A. Interpreting the Legacy Baseline", level="minor")
    rich_paragraph(doc, [
        ("On the recorded historical protocol, the Random Forest is materially stronger than the logistic-regression comparator in held-out F1, ROC-AUC, and PR-AUC (Table II). The new offline run reproduced those values exactly and matched the deployed artifact hash. This is strong evidence that the repository's legacy experiment is reproducible under its available inputs and configuration. It is not evidence that the same metrics apply to REFIT, to Wokwi data, or to a physical household. The missing source, licence, label-generation process, household identity, and causal construction of Avg_Past_Consumption prevent that stronger interpretation.", False, False, False),
    ])
    heading(doc, "B. Interpreting REFIT Preparation", level="minor")
    rich_paragraph(doc, [
        ("The REFIT contribution improves the study's data provenance and prepares a real recorded-household experiment without prematurely claiming system performance. The primary cohort retains 468,847 interval records across three exact 14-day windows. All primary selection records reconcile, no duplicate or out-of-order Unix timestamps were observed in the selected windows, and each house passed the predefined cadence, coverage, missing-slot, and maximum-gap gates. H3/H11 remain separate because their documented solar-generation interference would otherwise confound ordinary consumption-only summaries.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("The numerical energy totals are estimated integrations of aggregate active-power samples, not a re-labelling of REFIT as direct kWh records. The use of actual consecutive intervals avoids treating gaps as normal 8 s readings; excluded cross-gap records remain traceable in the generated log. These controls make the dataset suitable for a future reconciliation experiment, but they do not establish that energyConsumed has been accepted by the current ingestion service or that alerts behave correctly under REFIT replay.", False, False, False),
    ])
    heading(doc, "C. System Implications", level="minor")
    rich_paragraph(doc, [
        ("The system's practical value is the explicit separation of event admission, persistence, analytics, inference, and alert delivery. This separation allows a future experiment to reconcile accepted telemetry against usage records, ML outcomes, and deduplicated notifications. It also supports an isolated-storage design in which replayed household data cannot contaminate existing development records. Those safeguards are design objectives and preparation work; they must be treated as planned until isolated provisioning, replay, and failure-mode tests produce logs.", False, False, False),
    ])

    heading(doc, "VI. LIMITATIONS AND FUTURE WORK")
    rich_paragraph(doc, [
        ("The study has four material limits. First, Wokwi telemetry is synthetic and uncalibrated. Second, the legacy supervised CSV has unresolved provenance and feature-construction questions, so its historical metrics cannot be generalized. Third, REFIT provides a documented aggregate-power source but does not yet provide a verified aggregate anomaly-label protocol for the H1/H2/H4 cohort. Fourth, the local 8 GB Windows laptop has a Docker memory limit of 3.696 GiB while the observed main stack requires approximately 2.89 GiB; an isolated REFIT provisioning stack must not run concurrently with the main stack. This resource constraint prevents using deployment status as a substitute for scientific validation.", False, False, False),
    ])
    rich_paragraph(doc, [
        ("Future work should first provision a fully isolated Keycloak realm, database volumes, Kafka broker, three household profiles, and three device-bound credentials. It should then validate the derived_interval_energy_kwh to energyConsumed mapping against the actual ingestion contract, run a small pilot replay with request/event reconciliation, and proceed to H1/H2/H4 only after pilot acceptance. A defensible anomaly study additionally requires independently verified label alignment, household-held-out or temporal evaluation, input-range controls, repeated latency/throughput measurements, alert-deduplication evidence, and fault/recovery testing. None of these future steps is represented as completed here.", False, False, False),
    ])

    heading(doc, "VII. CONCLUSION")
    rich_paragraph(doc, [
        ("This paper presents an evidence-aware Home Energy Tracker whose implemented architecture combines gateway-mediated admission, Kafka event processing, time-series and relational storage, deterministic insights, a baseline Random Forest service, duplicate-aware alerting, and monitoring. The legacy classifier's historical results were reproduced exactly offline, while a separate live check confirms baseline endpoint execution only. The REFIT work adds checksum-verified official data, a scientifically separated H1/H2/H4 primary cohort, actual-delta interval-energy normalization, quality controls, and deterministic artifact reproduction. Together, these results support faculty review of a transparent implementation and reproducibility study. They do not support a journal claim of real-world anomaly accuracy, completed REFIT end-to-end operation, production readiness, or generalizable performance until the outstanding isolated deployment and evaluation gates are satisfied.", False, False, False),
    ])

    add_references(doc)

    doc.core_properties.title = "A Microservices-Based Home Energy Tracker for Event-Driven Monitoring, Deduplicated Alerting, and ML-Based Anomaly Detection"
    doc.core_properties.subject = "Evidence-aware implementation and reproducibility study"
    doc.core_properties.author = "Rohit Shrivastava; Harsh Gupta; Anshika Rai; Syamasudha Veeragandham"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    create_document()
